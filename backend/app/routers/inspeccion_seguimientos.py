# ============================================================
# ROUTER PLANES DE ACCIÓN Y SEGUIMIENTO DE HALLAZGOS SST
# FASE 1.1.8.6 — PLANES DE ACCIÓN Y SEGUIMIENTO ENTERPRISE
# Archivo: backend/app/routers/inspeccion_seguimientos.py
# ============================================================

from datetime import date, datetime
from pathlib import Path
import io
import os
import uuid

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from app.auth.dependencies import require_roles, require_permission
from app.core.default_permissions import PERM_REGISTROS_ELIMINAR, PERM_REPORTES_EXPORTAR
from app.core.file_security import validate_upload
from app.database import get_db
from app.models.archivo_sst import ArchivoSST
from app.models.inspeccion import InspeccionHallazgoSST, InspeccionSST
from app.models.inspeccion_seguimiento import InspeccionHallazgoSeguimientoSST
from app.schemas.inspeccion_seguimiento_schema import (
    CierreHallazgoRequest,
    SeguimientoDashboardResponse,
    SeguimientoHallazgoCreate,
    SeguimientoHallazgoResponse,
    SeguimientoHallazgoUpdate,
)

router = APIRouter(prefix="/inspecciones-seguimientos", tags=["Planes de Acción y Seguimiento SST"])
ROLES_SST = ["SUPER_ADMIN", "ADMIN_EMPRESA", "RESPONSABLE_SST"]
EXPORTAR_REPORTES = require_permission(PERM_REPORTES_EXPORTAR)
ELIMINAR_REGISTROS = require_permission(PERM_REGISTROS_ELIMINAR)

UPLOAD_ROOT = Path(os.getenv("UPLOAD_DIR", "app/uploads")).resolve()
SEGUIMIENTOS_UPLOAD_DIR = UPLOAD_ROOT / "inspecciones" / "seguimientos"
SEGUIMIENTOS_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXT = {"pdf", "jpg", "jpeg", "png", "webp"}
ALLOWED_MIME = {"application/pdf", "image/jpeg", "image/png", "image/webp"}
MAX_UPLOAD_MB = 20


def _public_upload_url(file_path: Path) -> str:
    try:
        rel = file_path.resolve().relative_to(UPLOAD_ROOT)
        return "/uploads/" + rel.as_posix()
    except Exception:
        return "/uploads/inspecciones/seguimientos/" + file_path.name


def _optimizar_imagen_bytes(content: bytes, extension: str) -> tuple[bytes, str, str]:
    try:
        from PIL import Image
        image = Image.open(io.BytesIO(content))
        if image.mode not in ("RGB", "RGBA"):
            image = image.convert("RGB")
        image.thumbnail((1920, 1080))
        out = io.BytesIO()
        image.save(out, format="WEBP", quality=78, method=6, optimize=True)
        return out.getvalue(), "webp", "image/webp"
    except Exception:
        mime = "image/jpeg" if extension.lower() in ["jpg", "jpeg"] else f"image/{extension.lower()}"
        return content, extension.lower(), mime


def _guardar_upload(upload: UploadFile) -> tuple[Path, str, str, str, int]:
    validation = validate_upload(upload, allowed_extensions={f".{item}" for item in ALLOWED_EXT}, max_size_mb=MAX_UPLOAD_MB)
    original = validation.safe_filename or "evidencia_seguimiento"
    extension = validation.extension.lstrip(".")
    content = validation.content
    mime_type = validation.mime_type

    if extension in {"jpg", "jpeg", "png", "webp"}:
        content, extension, mime_type = _optimizar_imagen_bytes(content, extension)
    elif extension == "pdf":
        mime_type = "application/pdf"

    filename = f"{uuid.uuid4().hex}.{extension}"
    path = SEGUIMIENTOS_UPLOAD_DIR / filename
    path.write_bytes(content)
    return path, original, filename, mime_type, len(content)

def _archivo_to_dict(archivo: ArchivoSST):
    return {
        "id": archivo.id,
        "empresa_id": archivo.empresa_id,
        "usuario_id": archivo.usuario_id,
        "tipo": archivo.tipo,
        "nombre_original": archivo.nombre_original,
        "nombre_archivo": archivo.nombre_archivo,
        "ruta": archivo.ruta,
        "url": archivo.url,
        "extension": archivo.extension,
        "mime_type": archivo.mime_type,
        "tamano_bytes": archivo.tamano_bytes,
        "modulo": archivo.modulo,
        "referencia_id": archivo.referencia_id,
        "descripcion": archivo.descripcion,
        "activo": archivo.activo,
        "fecha_creacion": archivo.fecha_creacion,
        "thumb_url": getattr(archivo, "thumb_url", None) or archivo.url,
        "preview_url": getattr(archivo, "preview_url", None) or archivo.url,
    }


def _seguimiento_to_response(db: Session, item: InspeccionHallazgoSeguimientoSST):
    data = SeguimientoHallazgoResponse.model_validate(item)
    data.total_evidencias = db.query(func.count(ArchivoSST.id)).filter(
        ArchivoSST.modulo == "INSPECCIONES_SEGUIMIENTOS",
        ArchivoSST.referencia_id == item.id,
        ArchivoSST.activo.is_(True),
    ).scalar() or 0
    return data


def _validar_hallazgo(db: Session, hallazgo_id: int):
    hallazgo = db.query(InspeccionHallazgoSST).options(joinedload(InspeccionHallazgoSST.inspeccion)).filter(InspeccionHallazgoSST.id == hallazgo_id).first()
    if not hallazgo:
        raise HTTPException(status_code=404, detail="Hallazgo no encontrado")
    return hallazgo


@router.get("/dashboard/resumen", response_model=SeguimientoDashboardResponse)
def dashboard_planes_accion(
    inspeccion_id: int | None = Query(default=None),
    empresa_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    query = db.query(InspeccionHallazgoSST).options(joinedload(InspeccionHallazgoSST.inspeccion)).filter(InspeccionHallazgoSST.activo.is_(True))
    if inspeccion_id:
        query = query.filter(InspeccionHallazgoSST.inspeccion_id == inspeccion_id)
    if empresa_id:
        query = query.filter(InspeccionHallazgoSST.empresa_id == empresa_id)
    hallazgos = query.all()
    ids = [h.id for h in hallazgos]
    hoy = date.today()

    seguimientos = []
    evidencias = 0
    if ids:
        seguimientos = db.query(InspeccionHallazgoSeguimientoSST).filter(
            InspeccionHallazgoSeguimientoSST.hallazgo_id.in_(ids),
            InspeccionHallazgoSeguimientoSST.activo.is_(True),
        ).all()
        seg_ids = [s.id for s in seguimientos]
        if seg_ids:
            evidencias = db.query(func.count(ArchivoSST.id)).filter(
                ArchivoSST.modulo == "INSPECCIONES_SEGUIMIENTOS",
                ArchivoSST.referencia_id.in_(seg_ids),
                ArchivoSST.activo.is_(True),
            ).scalar() or 0

    abiertos = sum(1 for h in hallazgos if h.estado != "CERRADO")
    cerrados = sum(1 for h in hallazgos if h.estado == "CERRADO")
    vencidos = sum(1 for h in hallazgos if h.fecha_compromiso and h.fecha_compromiso < hoy and h.estado != "CERRADO")
    criticos = sum(1 for h in hallazgos if h.nivel_riesgo in ["ALTO", "CRITICO"] and h.estado != "CERRADO")
    avance_promedio = round(sum((s.porcentaje_avance or 0) for s in seguimientos) / len(seguimientos), 1) if seguimientos else 0
    cumplimiento = round((cerrados / len(hallazgos)) * 100, 1) if hallazgos else 0
    semaforo = "ROJO" if vencidos > 5 or criticos > 3 else "AMARILLO" if vencidos or abiertos else "VERDE"

    recomendaciones = []
    if vencidos:
        recomendaciones.append("Reprogramar y escalar acciones correctivas vencidas.")
    if criticos:
        recomendaciones.append("Priorizar hallazgos alto/crítico con responsable y evidencia de cierre.")
    if abiertos and not seguimientos:
        recomendaciones.append("Registrar seguimientos periódicos para los hallazgos abiertos.")
    if not recomendaciones:
        recomendaciones.append("Planes de acción bajo control. Mantén trazabilidad documental.")

    return {
        "kpis": {
            "hallazgos": len(hallazgos),
            "abiertos": abiertos,
            "cerrados": cerrados,
            "vencidos": vencidos,
            "criticos": criticos,
            "seguimientos": len(seguimientos),
            "evidencias": evidencias,
            "avance_promedio": avance_promedio,
            "cumplimiento": cumplimiento,
            "semaforo": semaforo,
        },
        "alertas": {"vencidos": vencidos, "criticos": criticos, "sin_seguimiento": max(len(hallazgos) - len({s.hallazgo_id for s in seguimientos}), 0)},
        "recomendaciones": recomendaciones,
    }


@router.get("/exportaciones/excel")
def exportar_seguimientos_excel(
    inspeccion_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    query = db.query(InspeccionHallazgoSeguimientoSST).join(InspeccionHallazgoSST).filter(InspeccionHallazgoSeguimientoSST.activo.is_(True))
    if inspeccion_id:
        query = query.filter(InspeccionHallazgoSST.inspeccion_id == inspeccion_id)
    items = query.order_by(InspeccionHallazgoSeguimientoSST.fecha_registro.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Seguimientos SST"
    headers = ["ID", "Hallazgo", "Tipo acción", "Responsable", "Estado", "Resultado", "% Avance", "Comentario", "Fecha registro", "Próximo seguimiento"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for item in items:
        ws.append([
            item.id,
            item.hallazgo.descripcion if item.hallazgo else "",
            item.tipo_accion,
            item.responsable,
            item.estado,
            item.resultado,
            item.porcentaje_avance,
            item.comentario,
            item.fecha_registro.strftime("%Y-%m-%d %H:%M") if item.fecha_registro else "",
            item.fecha_proximo_seguimiento.isoformat() if item.fecha_proximo_seguimiento else "",
        ])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 2, 45)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=seguimientos_hallazgos_sst.xlsx"})


@router.get("/exportaciones/pdf")
def exportar_seguimientos_pdf(
    inspeccion_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    query = db.query(InspeccionHallazgoSeguimientoSST).join(InspeccionHallazgoSST).filter(InspeccionHallazgoSeguimientoSST.activo.is_(True))
    if inspeccion_id:
        query = query.filter(InspeccionHallazgoSST.inspeccion_id == inspeccion_id)
    items = query.order_by(InspeccionHallazgoSeguimientoSST.fecha_registro.desc()).all()

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(letter), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    story = [Paragraph("Seguimientos de Hallazgos SST", styles["Title"]), Spacer(1, 12)]
    data = [["Hallazgo", "Tipo", "Responsable", "Estado", "%", "Comentario", "Fecha"]]
    for item in items:
        data.append([
            Paragraph((item.hallazgo.descripcion if item.hallazgo else "")[:80], styles["BodyText"]),
            item.tipo_accion or "",
            item.responsable or "",
            item.estado or "",
            str(item.porcentaje_avance or 0),
            Paragraph((item.comentario or "")[:120], styles["BodyText"]),
            item.fecha_registro.strftime("%Y-%m-%d") if item.fecha_registro else "",
        ])
    table = Table(data, repeatRows=1, colWidths=[180, 75, 100, 75, 35, 230, 75])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    doc.build(story)
    out.seek(0)
    return StreamingResponse(out, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=seguimientos_hallazgos_sst.pdf"})


@router.get("/{hallazgo_id}", response_model=list[SeguimientoHallazgoResponse])
def listar_seguimientos_hallazgo(hallazgo_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    _validar_hallazgo(db, hallazgo_id)
    items = db.query(InspeccionHallazgoSeguimientoSST).filter(
        InspeccionHallazgoSeguimientoSST.hallazgo_id == hallazgo_id,
        InspeccionHallazgoSeguimientoSST.activo.is_(True),
    ).order_by(InspeccionHallazgoSeguimientoSST.fecha_registro.desc()).all()
    return [_seguimiento_to_response(db, item) for item in items]


@router.post("/", response_model=SeguimientoHallazgoResponse)
def crear_seguimiento_hallazgo(data: SeguimientoHallazgoCreate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    hallazgo = _validar_hallazgo(db, data.hallazgo_id)
    item = InspeccionHallazgoSeguimientoSST(
        hallazgo_id=data.hallazgo_id,
        usuario_id=getattr(usuario, "id", None),
        tipo_accion=data.tipo_accion,
        responsable=data.responsable or hallazgo.responsable,
        fecha_seguimiento=data.fecha_seguimiento or date.today(),
        fecha_proximo_seguimiento=data.fecha_proximo_seguimiento,
        estado=data.estado,
        resultado=data.resultado,
        comentario=data.comentario,
        porcentaje_avance=data.porcentaje_avance,
        requiere_evidencia=data.requiere_evidencia,
        observaciones=data.observaciones,
        activo=True,
    )
    db.add(item)
    if hallazgo.estado == "ABIERTO":
        hallazgo.estado = "EN_SEGUIMIENTO"
    if data.porcentaje_avance >= 100:
        hallazgo.estado = "EN_SEGUIMIENTO"
    db.commit()
    db.refresh(item)
    return _seguimiento_to_response(db, item)


@router.put("/{seguimiento_id}", response_model=SeguimientoHallazgoResponse)
def actualizar_seguimiento_hallazgo(seguimiento_id: int, data: SeguimientoHallazgoUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = db.query(InspeccionHallazgoSeguimientoSST).filter(InspeccionHallazgoSeguimientoSST.id == seguimiento_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Seguimiento no encontrado")
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(item, key, value)
    if item.hallazgo and item.hallazgo.estado == "ABIERTO":
        item.hallazgo.estado = "EN_SEGUIMIENTO"
    db.commit()
    db.refresh(item)
    return _seguimiento_to_response(db, item)


@router.delete("/{seguimiento_id}")
def eliminar_seguimiento_hallazgo(seguimiento_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    item = db.query(InspeccionHallazgoSeguimientoSST).filter(InspeccionHallazgoSeguimientoSST.id == seguimiento_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Seguimiento no encontrado")
    item.activo = False
    db.commit()
    return {"ok": True, "message": "Seguimiento desactivado"}


@router.post("/hallazgos/{hallazgo_id}/cerrar")
def cerrar_hallazgo_con_plan_accion(hallazgo_id: int, data: CierreHallazgoRequest, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    hallazgo = _validar_hallazgo(db, hallazgo_id)
    seguimientos = db.query(InspeccionHallazgoSeguimientoSST).filter(
        InspeccionHallazgoSeguimientoSST.hallazgo_id == hallazgo_id,
        InspeccionHallazgoSeguimientoSST.activo.is_(True),
    ).order_by(InspeccionHallazgoSeguimientoSST.fecha_registro.desc()).all()
    if not seguimientos and not data.forzar_cierre:
        raise HTTPException(status_code=400, detail="No se puede cerrar el hallazgo sin seguimientos registrados")
    ultimo = seguimientos[0] if seguimientos else None
    if ultimo and (ultimo.porcentaje_avance or 0) < 100 and not data.forzar_cierre:
        raise HTTPException(status_code=400, detail="No se puede cerrar el hallazgo con avance menor al 100%")
    evidencias = 0
    if ultimo:
        evidencias = db.query(func.count(ArchivoSST.id)).filter(
            ArchivoSST.modulo == "INSPECCIONES_SEGUIMIENTOS",
            ArchivoSST.referencia_id == ultimo.id,
            ArchivoSST.activo.is_(True),
        ).scalar() or 0
    if ultimo and ultimo.requiere_evidencia and evidencias == 0 and not data.forzar_cierre:
        raise HTTPException(status_code=400, detail="No se puede cerrar el hallazgo sin evidencia del seguimiento final")
    hallazgo.estado = "CERRADO"
    hallazgo.fecha_cierre = date.today()
    hallazgo.observaciones = ((hallazgo.observaciones or "") + f"\n[{datetime.utcnow().isoformat()}] Cierre plan acción: {data.observacion or ''}").strip()
    if ultimo:
        ultimo.estado = "CERRADO"
        ultimo.resultado = "CERRADO"
        ultimo.porcentaje_avance = 100
    db.commit()
    return {"ok": True, "message": "Hallazgo cerrado correctamente", "hallazgo_id": hallazgo.id, "estado": hallazgo.estado, "fecha_cierre": str(hallazgo.fecha_cierre)}


@router.get("/{seguimiento_id}/evidencias")
def listar_evidencias_seguimiento(seguimiento_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = db.query(InspeccionHallazgoSeguimientoSST).filter(InspeccionHallazgoSeguimientoSST.id == seguimiento_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Seguimiento no encontrado")
    archivos = db.query(ArchivoSST).filter(
        ArchivoSST.modulo == "INSPECCIONES_SEGUIMIENTOS",
        ArchivoSST.referencia_id == seguimiento_id,
        ArchivoSST.activo.is_(True),
    ).order_by(ArchivoSST.fecha_creacion.desc()).all()
    return [_archivo_to_dict(a) for a in archivos]


@router.post("/{seguimiento_id}/evidencias")
def subir_evidencia_seguimiento(
    seguimiento_id: int,
    tipo_evidencia: str = Form(default="SEGUIMIENTO"),
    descripcion: str = Form(default=""),
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    item = db.query(InspeccionHallazgoSeguimientoSST).options(joinedload(InspeccionHallazgoSeguimientoSST.hallazgo)).filter(InspeccionHallazgoSeguimientoSST.id == seguimiento_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Seguimiento no encontrado")
    path, original, filename, mime_type, size = _guardar_upload(archivo)
    empresa_id = item.hallazgo.empresa_id if item.hallazgo else None
    registro = ArchivoSST(
        empresa_id=empresa_id,
        usuario_id=getattr(usuario, "id", None),
        tipo=(tipo_evidencia or "SEGUIMIENTO").upper().strip(),
        nombre_original=original,
        nombre_archivo=filename,
        ruta=str(path),
        url=_public_upload_url(path),
        extension=filename.rsplit(".", 1)[-1].lower(),
        mime_type=mime_type,
        tamano_bytes=size,
        modulo="INSPECCIONES_SEGUIMIENTOS",
        referencia_id=item.id,
        descripcion=descripcion or "Evidencia de seguimiento de hallazgo SST",
        activo=True,
    )
    db.add(registro)
    db.commit()
    db.refresh(registro)
    return _archivo_to_dict(registro)


@router.delete("/{seguimiento_id}/evidencias/{archivo_id}")
def eliminar_evidencia_seguimiento(seguimiento_id: int, archivo_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    archivo = db.query(ArchivoSST).filter(
        ArchivoSST.id == archivo_id,
        ArchivoSST.modulo == "INSPECCIONES_SEGUIMIENTOS",
        ArchivoSST.referencia_id == seguimiento_id,
    ).first()
    if not archivo:
        raise HTTPException(status_code=404, detail="Evidencia no encontrada")
    archivo.activo = False
    db.commit()
    return {"ok": True, "message": "Evidencia desactivada"}
