
# ============================================================
# ROUTER AUDITORÍA SISTEMA PRO
# ERP SST PRO ENTERPRISE - FASE 35.4.1
# ============================================================
from __future__ import annotations
from datetime import datetime, timedelta
from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Session
from app.auth.dependencies import require_roles, require_permission
from app.core.default_permissions import PERM_REPORTES_EXPORTAR
from app.database import get_db
from app.models.auditoria import Auditoria
from app.models.empresa import Empresa
from app.models.usuario import Usuario
from app.schemas.auditoria_schema import AuditoriaDetalleResponse, AuditoriaKPIResponse, AuditoriaResponse, AuditoriaSistemaResponse

router=APIRouter(prefix='/auditoria', tags=['Auditoría PRO'])
ROLES_AUDITORIA=['SUPER_ADMIN','AUDITOR']
EXPORTAR_REPORTES = require_permission(PERM_REPORTES_EXPORTAR)

def _fecha_col():
    if hasattr(Auditoria,'fecha'): return Auditoria.fecha
    if hasattr(Auditoria,'fecha_creacion'): return Auditoria.fecha_creacion
    return None

def _fecha_val(item):
    return getattr(item,'fecha',None) or getattr(item,'fecha_creacion',None)

def _nombre_usuario(u):
    if not u: return None
    n=f"{getattr(u,'nombres','') or ''} {getattr(u,'apellidos','') or ''}".strip()
    return n or getattr(u,'correo',None)

def _detalle(item,db):
    u=db.query(Usuario).filter(Usuario.id==item.usuario_id).first() if getattr(item,'usuario_id',None) else None
    e=db.query(Empresa).filter(Empresa.id==item.empresa_id).first() if getattr(item,'empresa_id',None) else None
    return AuditoriaDetalleResponse(id=item.id, usuario_id=getattr(item,'usuario_id',None), empresa_id=getattr(item,'empresa_id',None), metodo=getattr(item,'metodo',None), ruta=getattr(item,'ruta',None), accion=getattr(item,'accion',None), ip=getattr(item,'ip',None), user_agent=getattr(item,'user_agent',None), fecha=_fecha_val(item), usuario_nombre=_nombre_usuario(u), usuario_correo=getattr(u,'correo',None) if u else None, empresa_nombre=getattr(e,'nombre',None) if e else None)

def _parse_fecha(v,fin=False):
    if not v: return None
    try:
        if len(v)==10:
            d=datetime.strptime(v,'%Y-%m-%d')
            return d+timedelta(days=1)-timedelta(microseconds=1) if fin else d
        return datetime.fromisoformat(v)
    except Exception:
        raise HTTPException(status_code=422, detail=f'Formato de fecha inválido: {v}')

def _query(db,usuario_id=None,empresa_id=None,metodo=None,ruta=None,ip=None,q=None,desde=None,hasta=None):
    query=db.query(Auditoria)
    if usuario_id: query=query.filter(Auditoria.usuario_id==usuario_id)
    if empresa_id: query=query.filter(Auditoria.empresa_id==empresa_id)
    if metodo: query=query.filter(func.upper(Auditoria.metodo)==metodo.upper().strip())
    if ruta: query=query.filter(Auditoria.ruta.ilike(f'%{ruta.strip()}%'))
    if ip: query=query.filter(Auditoria.ip.ilike(f'%{ip.strip()}%'))
    if q:
        t=f'%{q.strip()}%'
        query=query.filter(or_(Auditoria.ruta.ilike(t),Auditoria.accion.ilike(t),Auditoria.ip.ilike(t),Auditoria.user_agent.ilike(t),Auditoria.metodo.ilike(t)))
    fc=_fecha_col()
    if fc is not None and desde: query=query.filter(fc>=desde)
    if fc is not None and hasta: query=query.filter(fc<=hasta)
    return query

@router.get('/', response_model=list[AuditoriaResponse])
def listar_auditoria(db:Session=Depends(get_db), usuario=Depends(require_roles(ROLES_AUDITORIA))):
    return db.query(Auditoria).order_by(Auditoria.id.desc()).limit(300).all()

@router.get('/sistema', response_model=AuditoriaSistemaResponse)
def auditoria_sistema(usuario_id:int|None=Query(default=None), empresa_id:int|None=Query(default=None), metodo:str|None=Query(default=None), ruta:str|None=Query(default=None), ip:str|None=Query(default=None), q:str|None=Query(default=None), fecha_desde:str|None=Query(default=None), fecha_hasta:str|None=Query(default=None), limit:int=Query(default=300, ge=1, le=2000), db:Session=Depends(get_db), usuario=Depends(require_roles(ROLES_AUDITORIA))):
    desde=_parse_fecha(fecha_desde); hasta=_parse_fecha(fecha_hasta, True)
    base=_query(db,usuario_id,empresa_id,metodo,ruta,ip,q,desde,hasta)
    total=base.count()
    eventos=base.order_by(Auditoria.id.desc()).limit(limit).all()
    def cm(m): return base.filter(func.upper(Auditoria.metodo)==m).count()
    metodos={r[0] or 'SIN_METODO':r[1] for r in base.with_entities(Auditoria.metodo,func.count(Auditoria.id)).group_by(Auditoria.metodo).all()}
    rutas_top=[{'ruta':r[0] or 'SIN_RUTA','total':r[1]} for r in base.with_entities(Auditoria.ruta,func.count(Auditoria.id)).group_by(Auditoria.ruta).order_by(func.count(Auditoria.id).desc()).limit(10).all()]
    usuarios_top=[]
    for uid,cnt in base.with_entities(Auditoria.usuario_id,func.count(Auditoria.id)).group_by(Auditoria.usuario_id).order_by(func.count(Auditoria.id).desc()).limit(10).all():
        u=db.query(Usuario).filter(Usuario.id==uid).first() if uid else None
        usuarios_top.append({'usuario_id':uid,'usuario':_nombre_usuario(u) or 'Sin usuario','correo':getattr(u,'correo',None) if u else None,'total':cnt})
    k=AuditoriaKPIResponse(total_eventos=total,eventos_get=cm('GET'),eventos_post=cm('POST'),eventos_put=cm('PUT'),eventos_delete=cm('DELETE'),eventos_patch=cm('PATCH'),usuarios_unicos=base.with_entities(Auditoria.usuario_id).distinct().count(),ips_unicas=base.with_entities(Auditoria.ip).distinct().count(),rutas_unicas=base.with_entities(Auditoria.ruta).distinct().count(),eventos_sin_usuario=base.filter(Auditoria.usuario_id.is_(None)).count())
    rec=[]
    if k.eventos_sin_usuario: rec.append('Existen eventos sin usuario asociado. Revisar OPTIONS, rutas públicas o middleware de auditoría.')
    if k.eventos_delete: rec.append('Revisar eventos DELETE para confirmar eliminación lógica y trazabilidad.')
    if not rec: rec.append('Auditoría del sistema sin hallazgos críticos en los filtros actuales.')
    return AuditoriaSistemaResponse(kpis=k,eventos=[_detalle(x,db) for x in eventos],metodos=metodos,rutas_top=rutas_top,usuarios_top=usuarios_top,recomendaciones=rec)

@router.get('/usuario/{usuario_id}', response_model=list[AuditoriaResponse])
def auditoria_por_usuario(usuario_id:int, db:Session=Depends(get_db), usuario=Depends(require_roles(ROLES_AUDITORIA))):
    return db.query(Auditoria).filter(Auditoria.usuario_id==usuario_id).order_by(Auditoria.id.desc()).limit(300).all()

@router.get('/{auditoria_id}', response_model=AuditoriaDetalleResponse)
def detalle_auditoria(auditoria_id:int, db:Session=Depends(get_db), usuario=Depends(require_roles(ROLES_AUDITORIA))):
    item=db.query(Auditoria).filter(Auditoria.id==auditoria_id).first()
    if not item: raise HTTPException(status_code=404, detail='Evento de auditoría no encontrado')
    return _detalle(item,db)

@router.get('/export/excel')
def exportar_auditoria_excel(usuario_id:int|None=Query(default=None), empresa_id:int|None=Query(default=None), metodo:str|None=Query(default=None), ruta:str|None=Query(default=None), ip:str|None=Query(default=None), q:str|None=Query(default=None), fecha_desde:str|None=Query(default=None), fecha_hasta:str|None=Query(default=None), limit:int=Query(default=2000, ge=1, le=10000), db:Session=Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    eventos=_query(db,usuario_id,empresa_id,metodo,ruta,ip,q,_parse_fecha(fecha_desde),_parse_fecha(fecha_hasta,True)).order_by(Auditoria.id.desc()).limit(limit).all()
    wb=Workbook(); ws=wb.active; ws.title='Auditoría Sistema'
    headers=['ID','Fecha','Usuario ID','Usuario','Correo','Empresa ID','Empresa','Método','Ruta','Acción','IP','User Agent']
    ws.append(['ERP SST PRO - Auditoría del Sistema']); ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers)); ws['A1'].font=Font(bold=True,color='FFFFFF',size=14); ws['A1'].fill=PatternFill('solid',fgColor='173A8A'); ws['A1'].alignment=Alignment(horizontal='center')
    ws.append([f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"]); ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers)); ws['A2'].alignment=Alignment(horizontal='center')
    ws.append([]); ws.append(headers)
    border=Border(left=Side(style='thin',color='CBD5E1'),right=Side(style='thin',color='CBD5E1'),top=Side(style='thin',color='CBD5E1'),bottom=Side(style='thin',color='CBD5E1'))
    for cell in ws[4]: cell.font=Font(bold=True,color='1E293B'); cell.fill=PatternFill('solid',fgColor='EAF0FF'); cell.alignment=Alignment(horizontal='center'); cell.border=border
    for item in eventos:
        d=_detalle(item,db); ws.append([d.id,d.fecha.strftime('%Y-%m-%d %H:%M:%S') if d.fecha else '',d.usuario_id or '',d.usuario_nombre or '',d.usuario_correo or '',d.empresa_id or '',d.empresa_nombre or '',d.metodo or '',d.ruta or '',d.accion or '',d.ip or '',d.user_agent or ''])
    for row in ws.iter_rows(min_row=5):
        for c in row: c.border=border; c.alignment=Alignment(vertical='top',wrap_text=True)
    widths=[9,22,12,28,34,12,28,12,42,52,18,70]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='A5'; out=BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f"attachment; filename=auditoria_sistema_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"})

@router.get('/export/pdf')
def exportar_auditoria_pdf(usuario_id:int|None=Query(default=None), empresa_id:int|None=Query(default=None), metodo:str|None=Query(default=None), ruta:str|None=Query(default=None), ip:str|None=Query(default=None), q:str|None=Query(default=None), fecha_desde:str|None=Query(default=None), fecha_hasta:str|None=Query(default=None), limit:int=Query(default=300, ge=1, le=2000), db:Session=Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib.styles import ParagraphStyle,getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph,SimpleDocTemplate,Spacer,Table,TableStyle
    eventos=_query(db,usuario_id,empresa_id,metodo,ruta,ip,q,_parse_fecha(fecha_desde),_parse_fecha(fecha_hasta,True)).order_by(Auditoria.id.desc()).limit(limit).all()
    buf=BytesIO(); doc=SimpleDocTemplate(buf,pagesize=landscape(A4),rightMargin=cm,leftMargin=cm,topMargin=cm,bottomMargin=cm); styles=getSampleStyleSheet(); title=ParagraphStyle('TituloAuditoria',parent=styles['Title'],alignment=TA_CENTER,fontSize=16,textColor=colors.HexColor('#173A8A'))
    story=[Paragraph('ERP SST PRO - Auditoría del Sistema',title),Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')} · Total exportado: {len(eventos)}",styles['Normal']),Spacer(1,.25*cm)]
    small=ParagraphStyle('Small',parent=styles['BodyText'],fontSize=7,leading=9)
    rows=[['ID','Fecha','Usuario','Método','Ruta','IP','Acción']]
    for item in eventos:
        d=_detalle(item,db); rows.append([str(d.id),d.fecha.strftime('%Y-%m-%d %H:%M') if d.fecha else '',Paragraph(d.usuario_nombre or 'Sin usuario',small),d.metodo or '',Paragraph(d.ruta or '',small),d.ip or '',Paragraph(d.accion or '',small)])
    if len(rows)==1: rows.append(['Sin registros','','','','','',''])
    table=Table(rows,repeatRows=1,colWidths=[1.5*cm,3.1*cm,4*cm,1.7*cm,7*cm,2.8*cm,7*cm])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#173A8A')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),7),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#CBD5E1')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F8FAFC')])]))
    story.append(table); doc.build(story); buf.seek(0)
    return StreamingResponse(buf,media_type='application/pdf',headers={'Content-Disposition':f"attachment; filename=auditoria_sistema_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"})
