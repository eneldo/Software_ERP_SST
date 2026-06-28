# ============================================================
# EXPORTACIONES INSPECCIONES SST ENTERPRISE
# FASE 1.1.8.3 + Dashboard PDF 1.1.8.4
# Archivo: backend/app/routers/inspecciones_exportaciones.py
# ============================================================

from datetime import date, datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from app.auth.dependencies import require_roles
from app.database import get_db
from app.models.archivo_sst import ArchivoSST
from app.models.inspeccion import InspeccionHallazgoSST, InspeccionSST
from app.models.inspeccion_seguimiento import InspeccionHallazgoSeguimientoSST

router = APIRouter(prefix="/inspecciones/exportaciones", tags=["Exportación Inspecciones SST"])
ROLES_SST = ["SUPER_ADMIN", "ADMIN_EMPRESA", "RESPONSABLE_SST"]


def _safe(v):
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return v.strftime("%Y-%m-%d")
    return str(v)


def _filename(nombre, ext):
    return f"{nombre}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"


def _query(db, empresa_id=None, sede_id=None, area_id=None, estado=None, riesgo=None):
    q = db.query(InspeccionSST).options(
        joinedload(InspeccionSST.empresa), joinedload(InspeccionSST.sede),
        joinedload(InspeccionSST.area), joinedload(InspeccionSST.cargo), joinedload(InspeccionSST.empleado)
    ).filter(InspeccionSST.activo.is_(True))
    if empresa_id:
        q = q.filter(InspeccionSST.empresa_id == empresa_id)
    if sede_id:
        q = q.filter(InspeccionSST.sede_id == sede_id)
    if area_id:
        q = q.filter(InspeccionSST.area_id == area_id)
    if estado:
        q = q.filter(func.upper(InspeccionSST.estado) == estado.upper())
    if riesgo:
        q = q.filter(func.upper(InspeccionSST.nivel_riesgo) == riesgo.upper())
    return q.order_by(InspeccionSST.fecha_inspeccion.desc(), InspeccionSST.id.desc())


def _xlsx_response(wb, filename):
    bio = BytesIO(); wb.save(bio); bio.seek(0)
    return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _pdf_response(buffer, filename):
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _build_wb(title, headers, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = title[:31]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78"); font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for c in ws[1]:
        c.fill = fill; c.font = font; c.alignment = Alignment(horizontal="center"); c.border = Border(bottom=thin)
    for row in rows: ws.append(row)
    for col in ws.columns:
        width = min(max(len(_safe(cell.value)) for cell in col) + 3, 55)
        ws.column_dimensions[get_column_letter(col[0].column)].width = width
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    return wb


def _pdf_base(title):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28, title=title)
    return buffer, doc


def _styles():
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    st = getSampleStyleSheet()
    st.add(ParagraphStyle(name="TitleCenter", parent=st["Title"], alignment=TA_CENTER, fontSize=16, leading=20))
    st.add(ParagraphStyle(name="Small", parent=st["BodyText"], fontSize=8, leading=10))
    return st


def _table(data, widths=None):
    from reportlab.platypus import Table
    from reportlab.lib import colors
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1F4E78")), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,-1), 7),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D9E2F3")), ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F7FAFC")]),
    ])
    return t


@router.get("/excel-general")
def excel_general(empresa_id: int | None = None, sede_id: int | None = None, area_id: int | None = None, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    items = _query(db, empresa_id, sede_id, area_id).all()
    headers = ["Código", "Título", "Empresa", "Sede", "Área", "Fecha", "Estado", "Resultado", "Riesgo", "% Cumplimiento", "Hallazgos", "Responsable"]
    rows = []
    for i in items:
        hall = db.query(func.count(InspeccionHallazgoSST.id)).filter(InspeccionHallazgoSST.inspeccion_id == i.id, InspeccionHallazgoSST.activo.is_(True)).scalar() or 0
        rows.append([i.codigo, i.titulo, i.empresa.nombre if i.empresa else "", i.sede.nombre if i.sede else "", i.area.nombre if i.area else "", _safe(i.fecha_inspeccion), i.estado, i.resultado, i.nivel_riesgo, float(i.cumplimiento or 0), hall, i.responsable or ""])
    return _xlsx_response(_build_wb("Inspecciones", headers, rows), _filename("inspecciones_general", "xlsx"))


@router.get("/hallazgos-excel")
def hallazgos_excel(empresa_id: int | None = None, inspeccion_id: int | None = None, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    q = db.query(InspeccionHallazgoSST).join(InspeccionSST, InspeccionSST.id == InspeccionHallazgoSST.inspeccion_id).filter(InspeccionHallazgoSST.activo.is_(True))
    if empresa_id: q = q.filter(InspeccionHallazgoSST.empresa_id == empresa_id)
    if inspeccion_id: q = q.filter(InspeccionHallazgoSST.inspeccion_id == inspeccion_id)
    rows=[]
    for h in q.order_by(InspeccionHallazgoSST.id.desc()).all():
        rows.append([h.inspeccion.codigo if h.inspeccion else "", h.descripcion, h.tipo_hallazgo, h.nivel_riesgo, h.estado, h.responsable or "", _safe(h.fecha_compromiso), _safe(h.fecha_cierre), h.accion_recomendada or ""])
    headers=["Inspección", "Descripción", "Tipo", "Riesgo", "Estado", "Responsable", "Compromiso", "Cierre", "Acción recomendada"]
    return _xlsx_response(_build_wb("Hallazgos", headers, rows), _filename("hallazgos_inspecciones", "xlsx"))


@router.get("/pdf-general")
def pdf_general(empresa_id: int | None = None, sede_id: int | None = None, area_id: int | None = None, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    from reportlab.platypus import Paragraph, Spacer
    buffer, doc = _pdf_base("Inspecciones SST - General"); st = _styles(); story=[Paragraph("INSPECCIONES SST - REPORTE GENERAL", st["TitleCenter"]), Spacer(1,8)]
    data = [["Código", "Título", "Fecha", "Estado", "Riesgo", "%"]]
    for i in _query(db, empresa_id, sede_id, area_id).limit(120).all():
        data.append([i.codigo, i.titulo[:38], _safe(i.fecha_inspeccion), i.estado, i.nivel_riesgo, _safe(float(i.cumplimiento or 0))])
    story.append(_table(data, [65, 210, 65, 70, 55, 40])); doc.build(story); return _pdf_response(buffer, _filename("inspecciones_general", "pdf"))


def _get_inspeccion(db, inspeccion_id):
    item = _query(db).filter(InspeccionSST.id == inspeccion_id).first()
    if not item: raise HTTPException(status_code=404, detail="Inspección no encontrada")
    return item


@router.get("/{inspeccion_id}/pdf-individual")
def pdf_individual(inspeccion_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    from reportlab.platypus import Paragraph, Spacer
    item=_get_inspeccion(db, inspeccion_id); st=_styles(); buffer, doc=_pdf_base(f"Inspección {item.codigo}")
    story=[Paragraph("REPORTE INDIVIDUAL DE INSPECCIÓN SST", st["TitleCenter"]), Spacer(1,10)]
    story.append(_table([["Campo", "Información"],["Código", item.codigo],["Título", item.titulo],["Empresa", item.empresa.nombre if item.empresa else ""],["Sede / Área", f"{item.sede.nombre if item.sede else ''} / {item.area.nombre if item.area else ''}"],["Fecha", _safe(item.fecha_inspeccion)],["Estado / Resultado", f"{item.estado} / {item.resultado}"],["Riesgo", item.nivel_riesgo],["Cumplimiento", f"{float(item.cumplimiento or 0)}%"],["Responsable", item.responsable or ""],["Observaciones", item.observaciones or ""]], [115,390]))
    story.append(Spacer(1,10)); story.append(Paragraph("Hallazgos", st["Heading2"]))
    hall = db.query(InspeccionHallazgoSST).filter(InspeccionHallazgoSST.inspeccion_id==item.id, InspeccionHallazgoSST.activo.is_(True)).all()
    data=[["Descripción","Riesgo","Estado","Responsable","Compromiso"]]+[[h.descripcion[:80], h.nivel_riesgo, h.estado, h.responsable or "", _safe(h.fecha_compromiso)] for h in hall]
    story.append(_table(data, [210,55,70,90,80])); doc.build(story); return _pdf_response(buffer, _filename(f"inspeccion_{item.codigo}", "pdf"))


@router.get("/{inspeccion_id}/acta-pdf")
def acta_pdf(inspeccion_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    from reportlab.platypus import Paragraph, Spacer
    item=_get_inspeccion(db, inspeccion_id); st=_styles(); buffer, doc=_pdf_base(f"Acta {item.codigo}")
    firmas = [["Inspector", item.firma_inspector_nombre or "Pendiente", _safe(item.firma_inspector_fecha)], ["Responsable Área", item.firma_responsable_area_nombre or "Pendiente", _safe(item.firma_responsable_area_fecha)], ["SST", item.firma_sst_nombre or "Pendiente", _safe(item.firma_sst_fecha)]]
    story=[Paragraph("ACTA OFICIAL DE INSPECCIÓN SST", st["TitleCenter"]), Spacer(1,8), _table([["Código",item.codigo],["Título",item.titulo],["Fecha",_safe(item.fecha_inspeccion)],["Cierre digital", "Sí" if item.cierre_digital else "No"],["Fecha cierre", _safe(item.cierre_digital_fecha)]], [120,385]), Spacer(1,10), Paragraph("Firmas y cierre", st["Heading2"]), _table([["Rol","Firmante","Fecha"]]+firmas, [130,230,145])]
    doc.build(story); return _pdf_response(buffer, _filename(f"acta_inspeccion_{item.codigo}", "pdf"))


@router.get("/hallazgos-pdf")
def hallazgos_pdf(empresa_id: int | None = None, inspeccion_id: int | None = None, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    from reportlab.platypus import Paragraph, Spacer
    q=db.query(InspeccionHallazgoSST).join(InspeccionSST, InspeccionSST.id==InspeccionHallazgoSST.inspeccion_id).filter(InspeccionHallazgoSST.activo.is_(True))
    if empresa_id: q=q.filter(InspeccionHallazgoSST.empresa_id==empresa_id)
    if inspeccion_id: q=q.filter(InspeccionHallazgoSST.inspeccion_id==inspeccion_id)
    buffer, doc=_pdf_base("Hallazgos Inspecciones"); st=_styles(); data=[["Inspección","Hallazgo","Riesgo","Estado","Resp.","Compromiso"]]
    for h in q.order_by(InspeccionHallazgoSST.id.desc()).limit(150).all(): data.append([h.inspeccion.codigo if h.inspeccion else "", h.descripcion[:70], h.nivel_riesgo, h.estado, h.responsable or "", _safe(h.fecha_compromiso)])
    story=[Paragraph("HALLAZGOS DE INSPECCIONES SST", st["TitleCenter"]), Spacer(1,8), _table(data, [65,210,50,65,70,60])]
    doc.build(story); return _pdf_response(buffer, _filename("hallazgos_inspecciones", "pdf"))


@router.get("/seguimientos-pdf")
def seguimientos_pdf(inspeccion_id: int | None = None, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    from reportlab.platypus import Paragraph, Spacer
    q=db.query(InspeccionHallazgoSeguimientoSST).join(InspeccionHallazgoSST, InspeccionHallazgoSST.id==InspeccionHallazgoSeguimientoSST.hallazgo_id).filter(InspeccionHallazgoSeguimientoSST.activo.is_(True))
    if inspeccion_id: q=q.filter(InspeccionHallazgoSST.inspeccion_id==inspeccion_id)
    buffer, doc=_pdf_base("Seguimientos Inspecciones"); st=_styles(); data=[["Fecha","Hallazgo","Avance","Comentario"]]
    for s in q.order_by(InspeccionHallazgoSeguimientoSST.fecha_registro.desc()).limit(160).all(): data.append([_safe(s.fecha_registro), str(s.hallazgo_id), f"{s.porcentaje_avance or 0}%", s.comentario[:120]])
    story=[Paragraph("SEGUIMIENTOS DE HALLAZGOS SST", st["TitleCenter"]), Spacer(1,8), _table(data, [75,60,55,315])]
    doc.build(story); return _pdf_response(buffer, _filename("seguimientos_inspecciones", "pdf"))


@router.get("/dashboard-ejecutivo-pdf")
def dashboard_ejecutivo_pdf(empresa_id: int | None = None, sede_id: int | None = None, area_id: int | None = None, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    from reportlab.platypus import Paragraph, Spacer
    items=_query(db, empresa_id, sede_id, area_id).all(); ids=[i.id for i in items]; total=len(items)
    hall=[]
    if ids: hall=db.query(InspeccionHallazgoSST).filter(InspeccionHallazgoSST.inspeccion_id.in_(ids), InspeccionHallazgoSST.activo.is_(True)).all()
    cumplimiento=round(sum(float(i.cumplimiento or 0) for i in items)/total,1) if total else 0
    criticos=sum(1 for h in hall if h.nivel_riesgo in ["ALTO","CRITICO"])
    cerradas=sum(1 for i in items if i.estado=="CERRADA")
    buffer, doc=_pdf_base("Dashboard Ejecutivo Inspecciones"); st=_styles()
    kpi=[["Indicador","Valor"],["Total inspecciones", total],["Cerradas", cerradas],["Hallazgos", len(hall)],["Hallazgos críticos", criticos],["Cumplimiento promedio", f"{cumplimiento}%"]]
    story=[Paragraph("DASHBOARD EJECUTIVO INSPECCIONES SST", st["TitleCenter"]), Spacer(1,8), _table(kpi,[220,285]), Spacer(1,12), Paragraph("Lectura ejecutiva", st["Heading2"]), Paragraph("Reporte automático para seguimiento gerencial de inspecciones, hallazgos, riesgos, cierre digital y trazabilidad.", st["BodyText"])]
    doc.build(story); return _pdf_response(buffer, _filename("dashboard_ejecutivo_inspecciones", "pdf"))
