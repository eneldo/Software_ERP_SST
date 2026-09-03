# ============================================================
# ROUTER PERFIL SOCIODEMOGRÁFICO EMPLEADO - ERP SST PRO
# Encuesta integral de perfil sociodemográfico, salud y hoja de vida
# ============================================================

import io
import logging
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.auth.dependencies import require_roles
from app.database import get_db
from app.models.empleado import Empleado
from app.models.empleado_perfil_sociodemografico import EmpleadoPerfilSociodemografico
from app.models.empresa import Empresa
from app.schemas.empleado_perfil_schema import (
    EmpleadoPerfilCreate,
    EmpleadoPerfilResponse,
    EmpleadoPerfilUpdate,
)

router = APIRouter(
    prefix="/empleados-perfil",
    tags=["Perfil Sociodemográfico Empleado"],
)
ROLES_SST = ["SUPER_ADMIN", "ADMIN_EMPRESA", "RESPONSABLE_SST"]
logger = logging.getLogger("app.empleados_perfil")


# ── Helpers ─────────────────────────────────────────────────
def _get_perfil_or_404(empleado_id: int, db: Session, empresa_id: int | None = None):
    q = db.query(EmpleadoPerfilSociodemografico).filter(
        EmpleadoPerfilSociodemografico.empleado_id == empleado_id
    )
    if empresa_id:
        q = q.filter(EmpleadoPerfilSociodemografico.empresa_id == empresa_id)
    return q.first()


def _ensure_empleado(exists: bool, empleado_id: int, db: Session):
    if not exists:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")


# ── CRUD ────────────────────────────────────────────────────
@router.get(
    "/empleado/{empleado_id}",
    response_model=EmpleadoPerfilResponse,
    dependencies=[Depends(require_roles(ROLES_SST))],
)
def obtener_perfil(empleado_id: int, empresa_id: int | None = Query(None), db: Session = Depends(get_db)):
    emp = db.query(Empleado).filter(Empleado.id == empleado_id).first()
    _ensure_empleado(emp, empleado_id, db)
    perfil = _get_perfil_or_404(empleado_id, db, empresa_id)
    if not perfil:
        raise HTTPException(status_code=404, detail="Perfil sociodemográfico no encontrado para este empleado")
    return perfil


@router.post(
    "",
    response_model=EmpleadoPerfilResponse,
    status_code=201,
    dependencies=[Depends(require_roles(ROLES_SST))],
)
def crear_perfil(data: EmpleadoPerfilCreate, db: Session = Depends(get_db)):
    emp = db.query(Empleado).filter(Empleado.id == data.empleado_id).first()
    _ensure_empleado(emp, data.empleado_id, db)

    existente = _get_perfil_or_404(data.empleado_id, db, data.empresa_id)
    if existente:
        raise HTTPException(status_code=409, detail="El empleado ya tiene un perfil sociodemográfico. Use PUT para actualizar.")

    payload = data.model_dump(exclude_unset=True)
    perfil = EmpleadoPerfilSociodemografico(**payload)
    db.add(perfil)
    db.commit()
    db.refresh(perfil)
    logger.info("Perfil sociodemográfico creado para empleado %s", data.empleado_id)
    return perfil


@router.put(
    "/empleado/{empleado_id}",
    response_model=EmpleadoPerfilResponse,
    dependencies=[Depends(require_roles(ROLES_SST))],
)
def actualizar_perfil(
    empleado_id: int,
    data: EmpleadoPerfilUpdate,
    empresa_id: int | None = Query(None),
    db: Session = Depends(get_db),
):
    emp = db.query(Empleado).filter(Empleado.id == empleado_id).first()
    _ensure_empleado(emp, empleado_id, db)

    perfil = _get_perfil_or_404(empleado_id, db, empresa_id)
    if not perfil:
        raise HTTPException(status_code=404, detail="Perfil sociodemográfico no encontrado. Use POST para crear.")

    payload = data.model_dump(exclude_unset=True)
    for key, value in payload.items():
        setattr(perfil, key, value)
    db.commit()
    db.refresh(perfil)
    logger.info("Perfil sociodemográfico actualizado para empleado %s", empleado_id)
    return perfil


@router.delete(
    "/empleado/{empleado_id}",
    status_code=204,
    dependencies=[Depends(require_roles(ROLES_SST))],
)
def eliminar_perfil(
    empleado_id: int,
    empresa_id: int | None = Query(None),
    db: Session = Depends(get_db),
):
    perfil = _get_perfil_or_404(empleado_id, db, empresa_id)
    if not perfil:
        raise HTTPException(status_code=404, detail="Perfil no encontrado")
    db.delete(perfil)
    db.commit()
    logger.info("Perfil sociodemográfico eliminado para empleado %s", empleado_id)


# ── LISTA MASIVA ────────────────────────────────────────────
@router.get(
    "",
    response_model=list[EmpleadoPerfilResponse],
    dependencies=[Depends(require_roles(ROLES_SST))],
)
def listar_perfiles(
    empresa_id: int = Query(...),
    completado: bool | None = Query(None),
    db: Session = Depends(get_db),
):
    q = db.query(EmpleadoPerfilSociodemografico).filter(
        EmpleadoPerfilSociodemografico.empresa_id == empresa_id
    )
    if completado is not None:
        q = q.filter(EmpleadoPerfilSociodemografico.completado == completado)
    return q.order_by(EmpleadoPerfilSociodemografico.id.desc()).all()


# ── EXPORTACIÓN EXCEL ───────────────────────────────────────
@router.get(
    "/exportar-excel",
    dependencies=[Depends(require_roles(ROLES_SST))],
)
def exportar_perfiles_excel(
    empresa_id: int = Query(...),
    db: Session = Depends(get_db),
):
    perfiles = (
        db.query(EmpleadoPerfilSociodemografico)
        .filter(EmpleadoPerfilSociodemografico.empresa_id == empresa_id)
        .order_by(EmpleadoPerfilSociodemografico.id.desc())
        .all()
    )
    if not perfiles:
        raise HTTPException(status_code=404, detail="No hay perfiles sociodemográficos para exportar")

    wb = Workbook()
    ws = wb.active
    ws.title = "Perfiles Sociodemográficos"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = [
        "ID", "Empleado ID", "Nombres Completos", "Tipo Doc", "No. Documento",
        "Fecha Nac", "Lugar Nac", "Edad", "Étnia", "Teléfono",
        "Estado Civil", "Cónyuge", "Ocupación Cónyuge", "Dependientes",
        "Dirección", "Barrio", "Ciudad/Municipio", "Estrato", "Tipo Vivienda",
        "Transporte", "Tiempo Desplaz.", "Cargo Actual", "Área/Depto",
        "Tipo Contrato", "Tiempo Laborado", "Escolaridad", "EPS", "Fondo Pensiones",
        "RH", "Diagnóstico", "Actividad Física", "Cigarrillo", "Alcohol",
        "Talla Camisa", "Talla Pantalón", "Talla Calzado",
        "Ref. 1 Nombre", "Ref. 1 Teléfono", "Ref. 2 Nombre", "Ref. 2 Teléfono",
        "Consentimiento", "Completado", "Fuente",
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, p in enumerate(perfiles, 2):
        values = [
            p.id, p.empleado_id, p.nombres_completos, p.tipo_documento, p.numero_documento,
            p.fecha_nacimiento, p.lugar_nacimiento, p.edad, p.raza_pertenencia_etnica, p.telefono_celular,
            p.estado_civil, p.conyuge_nombre, p.conyuge_ocupacion, p.numero_dependientes,
            p.direccion_residencia, p.barrio, p.ciudad_municipio, p.estrato_socioeconomico, p.tipo_vivienda,
            p.medio_transporte, p.tiempo_desplazamiento, p.cargo_actual, p.area_departamento,
            p.tipo_contrato, p.tiempo_laborado, p.nivel_escolaridad, p.eps_actual, p.fondo_pensiones,
            p.tipo_rh, p.diagnostico_detalle if p.diagnostico_previo else "No",
            p.actividad_fisica, p.consumo_cigarrillo, p.consumo_alcohol,
            p.talla_camisa, p.talla_pantalon, p.talla_calzado,
            p.referencia_1_nombre, p.referencia_1_telefono, p.referencia_2_nombre, p.referencia_2_telefono,
            "Sí" if p.consentimiento_informado else "No",
            "Sí" if p.completado else "No", p.fuente,
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    fecha = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"perfiles_sociodemograficos_{fecha}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ── PLANTILLA EXCEL VACÍA ──────────────────────────────────
_PLANTILLA_HEADERS = [
    "Empleado ID",
    "Nombres Completos", "Tipo Doc", "No. Documento", "Libreta Militar",
    "Fecha Nacimiento (AAAA-MM-DD)", "Lugar Nacimiento", "Edad", "Raza / Pertenencia Étnica", "Teléfono / Celular",
    "Estado Civil", "Nombre Cónyuge", "Ocupación Cónyuge", "Edad Cónyuge", "Celular Cónyuge",
    "N° Dependientes", "Hijos (JSON: [{nombre,fecha_nacimiento,edad,escolaridad}])",
    "Dirección Residencia", "Barrio", "Ciudad / Municipio", "Estrato (1-6)", "Tipo Vivienda",
    "Medio Transporte", "Otro Transporte", "Tiempo Desplazamiento",
    "Cargo Actual", "Área / Depto", "Sede / Centro Trabajo", "Tipo Contrato",
    "Tiempo Laborado", "Antigüedad Cargo", "Última Empresa", "Nivel Escolaridad", "Detalle Títulos",
    "EPS Actual", "Fondo Pensiones", "Tipo RH",
    "Diagnóstico Previo (Sí/No)", "Detalle Diagnóstico",
    "Actividad Física (Sí/No)", "Consumo Cigarrillo (Nunca/Ocasional/Frecuente)", "Consumo Alcohol (Nunca/Ocasional/Frecuente)",
    "Talla Camisa", "Talla Pantalón", "Talla Chaqueta", "Talla Overol", "Talla Calzado",
    "Ref 1 Nombre", "Ref 1 Ocupación", "Ref 1 Teléfono",
    "Ref 2 Nombre", "Ref 2 Ocupación", "Ref 2 Teléfono",
    "Consentimiento (Sí/No)", "Fecha Firma (AAAA-MM-DD)", "Completado (Sí/No)",
]


@router.get(
    "/plantilla-excel",
    dependencies=[Depends(require_roles(ROLES_SST))],
)
def descargar_plantilla_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla Perfil Sociodemográfico"

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col, header in enumerate(_PLANTILLA_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = 22

    ws.row_dimensions[1].height = 40

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_perfil_sociodemografico.xlsx"},
    )


# ── IMPORTACIÓN EXCEL ───────────────────────────────────────
_PERFIL_FIELD_MAP = {
    1: "nombres_completos", 2: "tipo_documento", 3: "numero_documento",
    4: "libreta_militar", 5: "fecha_nacimiento", 6: "lugar_nacimiento",
    7: "edad", 8: "raza_pertenencia_etnica", 9: "telefono_celular",
    10: "estado_civil", 11: "conyuge_nombre", 12: "conyuge_ocupacion",
    13: "conyuge_edad", 14: "conyuge_celular", 15: "numero_dependientes",
    17: "direccion_residencia", 18: "barrio", 19: "ciudad_municipio",
    20: "estrato_socioeconomico", 21: "tipo_vivienda", 22: "medio_transporte",
    23: "medio_transporte_otro", 24: "tiempo_desplazamiento",
    25: "cargo_actual", 26: "area_departamento", 27: "sede_centro_trabajo",
    28: "tipo_contrato", 29: "tiempo_laborado", 30: "antiguedad_cargo",
    31: "ultima_empresa", 32: "nivel_escolaridad", 33: "detalle_titulos",
    34: "eps_actual", 35: "fondo_pensiones", 36: "tipo_rh",
    37: "diagnostico_previo", 38: "diagnostico_detalle",
    39: "actividad_fisica", 40: "consumo_cigarrillo", 41: "consumo_alcohol",
    42: "talla_camisa", 43: "talla_pantalon", 44: "talla_chaqueta",
    45: "talla_overol", 46: "talla_calzado",
    47: "referencia_1_nombre", 48: "referencia_1_ocupacion", 49: "referencia_1_telefono",
    50: "referencia_2_nombre", 51: "referencia_2_ocupacion", 52: "referencia_2_telefono",
    53: "consentimiento_informado", 54: "fecha_firma", 55: "completado",
}

_INT_FIELDS = {"edad", "conyuge_edad", "numero_dependientes", "estrato_socioeconomico"}
_BOOL_FIELDS = {"diagnostico_previo", "consentimiento_informado", "completado"}


def _parse_bool(value):
    if value is None:
        return False
    s = str(value).strip().lower()
    return s in ("sí", "si", "true", "1", "yes")


def _parse_int(value):
    if value is None:
        return None
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return None


def _parse_fecha(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if hasattr(value, "date"):
        return value.date()
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


@router.post(
    "/importar-excel",
    dependencies=[Depends(require_roles(ROLES_SST))],
)
async def importar_perfil_excel(
    archivo: UploadFile = File(...),
    empresa_id: int = Query(...),
    db: Session = Depends(get_db),
):
    if not archivo.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="El archivo debe ser un Excel (.xlsx)")

    contents = await archivo.read()
    try:
        wb = load_workbook(io.BytesIO(contents), read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo Excel. Verifique que no esté corrupto.")

    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="El archivo Excel está vacío o no tiene datos después del encabezado")

    resultados = {"importados": 0, "actualizados": 0, "errores": []}

    for row_idx, row in enumerate(rows, 2):
        if not row or not any(cell is not None for cell in row):
            continue

        empleado_id = _parse_int(row[0]) if len(row) > 0 else None
        if not empleado_id:
            resultados["errores"].append(f"Fila {row_idx}: Falta Empleado ID")
            continue

        emp = db.query(Empleado).filter(Empleado.id == empleado_id).first()
        if not emp:
            resultados["errores"].append(f"Fila {row_idx}: Empleado ID {empleado_id} no encontrado")
            continue

        data = {}
        for col_idx, field_name in _PERFIL_FIELD_MAP.items():
            if col_idx < len(row):
                value = row[col_idx]
                if value is not None:
                    if field_name in _INT_FIELDS:
                        data[field_name] = _parse_int(value)
                    elif field_name in _BOOL_FIELDS:
                        data[field_name] = _parse_bool(value)
                    elif field_name in ("fecha_nacimiento", "fecha_firma"):
                        data[field_name] = _parse_fecha(value)
                    else:
                        data[field_name] = str(value).strip() if str(value).strip() else None

        data["empresa_id"] = empresa_id
        data["empleado_id"] = empleado_id
        data["fuente"] = "IMPORTADO"

        perfil_existente = _get_perfil_or_404(empleado_id, db, empresa_id)
        if perfil_existente:
            for key, value in data.items():
                if key not in ("empresa_id", "empleado_id") and value is not None:
                    setattr(perfil_existente, key, value)
            resultados["actualizados"] += 1
        else:
            perfil = EmpleadoPerfilSociodemografico(**data)
            db.add(perfil)
            resultados["importados"] += 1

    db.commit()
    wb.close()
    logger.info(
        "Importación Excel: %d importados, %d actualizados, %d errores",
        resultados["importados"], resultados["actualizados"], len(resultados["errores"]),
    )
    return resultados
