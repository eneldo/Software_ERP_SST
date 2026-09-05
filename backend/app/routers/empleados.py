# ============================================================
# ROUTER EMPLEADOS - ERP SST PRO
# FASE 1.1.5.3 — EXPORTACIÓN PDF / EXCEL
# Archivo: backend/app/routers/empleados.py
# No crea tablas nuevas. Compatible con Empleados Enterprise 360°
# y Empleados Analytics PRO.
# ============================================================

from io import BytesIO
from datetime import datetime, date
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, or_

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.database import get_db
from app.models.empleado import Empleado
from app.models.empresa import Empresa
from app.models.sede import Sede
from app.models.area import Area
from app.models.cargo import Cargo
from app.schemas.empleado_schema import EmpleadoCreate, EmpleadoUpdate, EmpleadoResponse
from app.auth.dependencies import require_roles
from app.routers.empresas import validar_acceso_empresa


router = APIRouter(prefix="/empleados", tags=["Empleados SST Enterprise 360"])
ROLES_SST = ["SUPER_ADMIN", "ADMIN_EMPRESA", "RESPONSABLE_SST"]


# ============================================================
# Helpers
# ============================================================
def _texto(valor, defecto="Sin dato"):
    if valor is None:
        return defecto
    valor = str(valor).strip()
    return valor if valor else defecto


def _fecha(valor):
    if not valor:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")
    return str(valor)


def _bool_text(valor):
    return "Activo" if bool(valor) else "Inactivo"


def _nombre_archivo(prefijo: str, extension: str):
    fecha = datetime.now().strftime("%Y%m%d_%H%M")
    return f"{prefijo}_{fecha}.{extension}"


def _empleado_to_response(empleado: Empleado) -> EmpleadoResponse:
    """Respuesta compatible con Pydantic v2 y campos extra opcionales."""
    data = EmpleadoResponse.model_validate(empleado)
    if hasattr(data, "empresa_nombre"):
        data.empresa_nombre = empleado.empresa.nombre if empleado.empresa else None
    if hasattr(data, "sede_nombre"):
        data.sede_nombre = empleado.sede.nombre if empleado.sede else None
    if hasattr(data, "area_nombre"):
        data.area_nombre = empleado.area.nombre if empleado.area else None
    if hasattr(data, "cargo_nombre"):
        data.cargo_nombre = empleado.cargo.nombre if empleado.cargo else None
    return data


def _validar_relaciones(db: Session, data):
    empresa_id = getattr(data, "empresa_id", None)
    sede_id = getattr(data, "sede_id", None)
    area_id = getattr(data, "area_id", None)
    cargo_id = getattr(data, "cargo_id", None)

    if empresa_id:
        empresa = db.query(Empresa).filter(Empresa.id == empresa_id).first()
        if not empresa:
            raise HTTPException(status_code=404, detail="Empresa no encontrada")
    if sede_id:
        sede = db.query(Sede).filter(Sede.id == sede_id).first()
        if not sede:
            raise HTTPException(status_code=404, detail="Sede no encontrada")
    if area_id:
        area = db.query(Area).filter(Area.id == area_id).first()
        if not area:
            raise HTTPException(status_code=404, detail="Área no encontrada")
    if cargo_id:
        cargo = db.query(Cargo).filter(Cargo.id == cargo_id).first()
        if not cargo:
            raise HTTPException(status_code=404, detail="Cargo no encontrado")


def _payload_limpio(data):
    payload = data.model_dump(exclude_unset=True)
    for key in ["sede_id", "area_id", "cargo_id"]:
        if payload.get(key) in ["", 0, "0"]:
            payload[key] = None
    if payload.get("correo") == "":
        payload["correo"] = None
    return payload


def _empresa_id_autorizada(usuario, empresa_id: int | None) -> int | None:
    if str(getattr(usuario, "rol", "")).strip().upper() == "SUPER_ADMIN":
        return empresa_id
    usuario_empresa_id = getattr(usuario, "empresa_id", None)
    if not usuario_empresa_id:
        raise HTTPException(status_code=403, detail="Usuario sin empresa asignada")
    if empresa_id and int(empresa_id) != int(usuario_empresa_id):
        raise HTTPException(status_code=403, detail="No tiene permisos sobre esta empresa")
    return int(usuario_empresa_id)


def _query_empleados_filtrada(
    db: Session,
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    cargo_id: int | None = None,
    estado: str | None = None,
    q: str | None = None,
):
    query = db.query(Empleado)

    if empresa_id:
        query = query.filter(Empleado.empresa_id == empresa_id)
    if sede_id:
        query = query.filter(Empleado.sede_id == sede_id)
    if area_id:
        query = query.filter(Empleado.area_id == area_id)
    if cargo_id:
        query = query.filter(Empleado.cargo_id == cargo_id)
    if estado:
        estado_up = estado.upper().strip()
        if estado_up in ["ACTIVO", "INACTIVO", "RETIRADO", "SUSPENDIDO"]:
            query = query.filter(func.upper(Empleado.estado_laboral) == estado_up)
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(
                Empleado.nombres.ilike(like),
                Empleado.apellidos.ilike(like),
                Empleado.documento.ilike(like),
                Empleado.correo.ilike(like),
                Empleado.telefono.ilike(like),
            )
        )
    return query.order_by(Empleado.id.desc())


def _fila_empleado(empleado: Empleado):
    return [
        empleado.documento,
        f"{empleado.nombres} {empleado.apellidos}",
        empleado.tipo_documento,
        empleado.correo or "",
        empleado.telefono or "",
        empleado.empresa.nombre if empleado.empresa else "Sin empresa",
        empleado.sede.nombre if empleado.sede else "Sin sede",
        empleado.area.nombre if empleado.area else "Sin área",
        empleado.cargo.nombre if empleado.cargo else "Sin cargo",
        empleado.tipo_contrato or "",
        empleado.estado_laboral or "",
        _bool_text(empleado.activo),
        _fecha(empleado.fecha_ingreso),
        _fecha(empleado.fecha_nacimiento),
    ]


# ============================================================
# CRUD + Analytics
# ============================================================
@router.post("/", response_model=EmpleadoResponse)
def crear_empleado(
    data: EmpleadoCreate,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    validar_acceso_empresa(usuario, data.empresa_id)
    existe = db.query(Empleado).filter(Empleado.documento == data.documento).first()
    if existe:
        raise HTTPException(status_code=400, detail="Ya existe un empleado con este documento")

    _validar_relaciones(db, data)
    empleado = Empleado(**_payload_limpio(data))
    db.add(empleado)
    db.commit()
    db.refresh(empleado)
    return _empleado_to_response(empleado)


@router.get("/", response_model=list[EmpleadoResponse])
def listar_empleados(
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    cargo_id: int | None = None,
    estado: str | None = None,
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    empresa_id = _empresa_id_autorizada(usuario, empresa_id)
    empleados = _query_empleados_filtrada(db, empresa_id, sede_id, area_id, cargo_id, estado, q).all()
    return [_empleado_to_response(e) for e in empleados]


@router.get("/dashboard")
def dashboard_empleados(
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    cargo_id: int | None = None,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    empresa_id = _empresa_id_autorizada(usuario, empresa_id)
    empleados = _query_empleados_filtrada(db, empresa_id, sede_id, area_id, cargo_id).all()
    total = len(empleados)
    activos = sum(1 for e in empleados if (e.estado_laboral or "").upper() == "ACTIVO" and e.activo)
    inactivos = total - activos
    sin_sede = sum(1 for e in empleados if not e.sede_id)
    sin_area = sum(1 for e in empleados if not e.area_id)
    sin_cargo = sum(1 for e in empleados if not e.cargo_id)
    sin_correo = sum(1 for e in empleados if not e.correo)

    def agrupar(nombre_relacion, default="Sin dato"):
        tmp = {}
        for e in empleados:
            obj = getattr(e, nombre_relacion, None)
            nombre = getattr(obj, "nombre", None) or default
            tmp[nombre] = tmp.get(nombre, 0) + 1
        return [{"name": k, "value": v} for k, v in sorted(tmp.items(), key=lambda x: x[1], reverse=True)]

    def agrupar_attr(attr, default="Sin dato"):
        tmp = {}
        for e in empleados:
            nombre = getattr(e, attr, None) or default
            tmp[str(nombre)] = tmp.get(str(nombre), 0) + 1
        return [{"name": k, "value": v} for k, v in sorted(tmp.items(), key=lambda x: x[1], reverse=True)]

    estructura_completa = total - (sin_sede + sin_area + sin_cargo)
    completitud = round((estructura_completa / total) * 100, 1) if total else 100
    activos_pct = round((activos / total) * 100, 1) if total else 0

    return {
        "kpis": {
            "total": total,
            "activos": activos,
            "inactivos": inactivos,
            "empresas": len({e.empresa_id for e in empleados if e.empresa_id}),
            "sin_sede": sin_sede,
            "sin_area": sin_area,
            "sin_cargo": sin_cargo,
            "sin_correo": sin_correo,
            "completitud_organizacional": completitud,
            "activos_pct": activos_pct,
            "estructura_completa": estructura_completa,
            "pendientes_criticos": sin_sede + sin_area + sin_cargo + sin_correo,
        },
        "charts": {
            "por_empresa": agrupar("empresa", "Sin empresa"),
            "por_sede": agrupar("sede", "Sin sede"),
            "por_area": agrupar("area", "Sin área"),
            "por_cargo": agrupar("cargo", "Sin cargo"),
            "por_estado": agrupar_attr("estado_laboral", "Sin estado"),
            "por_contrato": agrupar_attr("tipo_contrato", "Sin contrato"),
        },
        "alertas": {
            "sin_sede": sin_sede,
            "sin_area": sin_area,
            "sin_cargo": sin_cargo,
            "sin_correo": sin_correo,
        },
    }


# ============================================================
# EXPORTACIÓN EXCEL / PDF
# Deben ir antes de /{empleado_id} para evitar conflicto de rutas.
# ============================================================
@router.get("/export/excel")
def exportar_empleados_excel(
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    cargo_id: int | None = None,
    estado: str | None = None,
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    empresa_id = _empresa_id_autorizada(usuario, empresa_id)
    empleados = _query_empleados_filtrada(db, empresa_id, sede_id, area_id, cargo_id, estado, q).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Empleados SST"

    headers = [
        "Documento", "Empleado", "Tipo Doc.", "Correo", "Teléfono", "Empresa", "Sede", "Área",
        "Cargo", "Tipo contrato", "Estado laboral", "Activo", "Fecha ingreso", "Fecha nacimiento",
    ]
    ws.append(["ERP SST PRO - Reporte de Empleados SST"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    ws["A1"].fill = PatternFill("solid", fgColor="173A8A")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.append([f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.append([])
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="EAF0FF")
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for cell in ws[4]:
        cell.font = Font(bold=True, color="1E293B")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for empleado in empleados:
        ws.append(_fila_empleado(empleado))

    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 12
        for cell in column_cells:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 35)

    ws.freeze_panes = "A5"
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = _nombre_archivo("empleados_sst", "xlsx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/pdf")
def exportar_empleados_pdf(
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    cargo_id: int | None = None,
    estado: str | None = None,
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    empresa_id = _empresa_id_autorizada(usuario, empresa_id)
    empleados = _query_empleados_filtrada(db, empresa_id, sede_id, area_id, cargo_id, estado, q).all()

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=1 * cm,
        leftMargin=1 * cm,
        topMargin=1 * cm,
        bottomMargin=1 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TituloEmpleados",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=16,
        textColor=colors.HexColor("#173A8A"),
        spaceAfter=8,
    )
    normal = ParagraphStyle("NormalSmall", parent=styles["BodyText"], fontSize=7, leading=9)

    story = [
        Paragraph("ERP SST PRO - Reporte General de Empleados", title_style),
        Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')} · Total registros: {len(empleados)}", styles["Normal"]),
        Spacer(1, 0.25 * cm),
    ]

    headers = ["Documento", "Empleado", "Empresa", "Sede", "Área", "Cargo", "Estado", "Ingreso"]
    data = [headers]
    for e in empleados:
        data.append([
            Paragraph(_texto(e.documento, ""), normal),
            Paragraph(_texto(f"{e.nombres} {e.apellidos}", ""), normal),
            Paragraph(_texto(e.empresa.nombre if e.empresa else None), normal),
            Paragraph(_texto(e.sede.nombre if e.sede else None), normal),
            Paragraph(_texto(e.area.nombre if e.area else None), normal),
            Paragraph(_texto(e.cargo.nombre if e.cargo else None), normal),
            Paragraph(_texto(e.estado_laboral, ""), normal),
            Paragraph(_fecha(e.fecha_ingreso), normal),
        ])

    table = Table(data, colWidths=[2.3 * cm, 4.0 * cm, 4.0 * cm, 3.0 * cm, 3.0 * cm, 3.5 * cm, 2.2 * cm, 2.2 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(table)
    doc.build(story)
    buffer.seek(0)

    filename = _nombre_archivo("empleados_sst", "pdf")
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/{empleado_id}/pdf")
def exportar_ficha_empleado_pdf(
    empleado_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    empleado = db.query(Empleado).filter(Empleado.id == empleado_id).first()
    if not empleado:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    validar_acceso_empresa(usuario, empleado.empresa_id)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1.5 * cm, leftMargin=1.5 * cm, topMargin=1.3 * cm, bottomMargin=1.3 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TituloFichaEmpleado", parent=styles["Title"], alignment=TA_CENTER, fontSize=17, textColor=colors.HexColor("#173A8A"))
    section_style = ParagraphStyle("Seccion", parent=styles["Heading2"], fontSize=11, textColor=colors.HexColor("#173A8A"), spaceBefore=10, spaceAfter=6)

    def tabla_pares(rows):
        tabla = Table(rows, colWidths=[5 * cm, 11 * cm])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF4FF")),
            ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#1E293B")),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ]))
        return tabla

    story = [
        Paragraph("Ficha Individual de Empleado SST", title_style),
        Paragraph(f"ERP SST PRO · Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]),
        Spacer(1, 0.25 * cm),
        Paragraph("Información personal", section_style),
        tabla_pares([
            ["Nombre completo", _texto(f"{empleado.nombres} {empleado.apellidos}", "")],
            ["Tipo documento", _texto(empleado.tipo_documento, "")],
            ["Documento", _texto(empleado.documento, "")],
            ["Correo", _texto(empleado.correo)],
            ["Teléfono", _texto(empleado.telefono)],
            ["Fecha nacimiento", _fecha(empleado.fecha_nacimiento) or "Sin dato"],
        ]),
        Paragraph("Información laboral", section_style),
        tabla_pares([
            ["Empresa", _texto(empleado.empresa.nombre if empleado.empresa else None)],
            ["Sede", _texto(empleado.sede.nombre if empleado.sede else None)],
            ["Área", _texto(empleado.area.nombre if empleado.area else None)],
            ["Cargo", _texto(empleado.cargo.nombre if empleado.cargo else None)],
            ["Fecha ingreso", _fecha(empleado.fecha_ingreso) or "Sin dato"],
            ["Tipo contrato", _texto(empleado.tipo_contrato)],
            ["Estado laboral", _texto(empleado.estado_laboral)],
            ["Activo", _bool_text(empleado.activo)],
        ]),
        Paragraph("Trazabilidad", section_style),
        tabla_pares([
            ["Fecha creación", _fecha(getattr(empleado, "fecha_creacion", None)) or "Sin dato"],
            ["Última actualización", _fecha(getattr(empleado, "fecha_actualizacion", None)) or "Sin dato"],
        ]),
    ]

    doc.build(story)
    buffer.seek(0)
    filename = f"ficha_empleado_{empleado.documento}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})


# ============================================================
# Rutas dinámicas al final
# ============================================================
@router.get("/{empleado_id}", response_model=EmpleadoResponse)
def obtener_empleado(
    empleado_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    empleado = db.query(Empleado).filter(Empleado.id == empleado_id).first()
    if not empleado:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    validar_acceso_empresa(usuario, empleado.empresa_id)
    return _empleado_to_response(empleado)


@router.put("/{empleado_id}", response_model=EmpleadoResponse)
def actualizar_empleado(
    empleado_id: int,
    data: EmpleadoUpdate,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    empleado = db.query(Empleado).filter(Empleado.id == empleado_id).first()
    if not empleado:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    validar_acceso_empresa(usuario, empleado.empresa_id)

    payload = _payload_limpio(data)
    if payload.get("empresa_id"):
        validar_acceso_empresa(usuario, payload["empresa_id"])
    temporal = type("Temporal", (), payload)()
    _validar_relaciones(db, temporal)

    if "documento" in payload:
        existe = db.query(Empleado).filter(Empleado.documento == payload["documento"], Empleado.id != empleado_id).first()
        if existe:
            raise HTTPException(status_code=400, detail="Ya existe otro empleado con este documento")

    for key, value in payload.items():
        setattr(empleado, key, value)

    db.commit()
    db.refresh(empleado)
    return _empleado_to_response(empleado)


@router.delete("/{empleado_id}")
def eliminar_empleado(
    empleado_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(["SUPER_ADMIN", "ADMIN_EMPRESA"])),
):
    empleado = db.query(Empleado).filter(Empleado.id == empleado_id).first()
    if not empleado:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    validar_acceso_empresa(usuario, empleado.empresa_id)
    empleado.activo = False
    empleado.estado_laboral = "INACTIVO"
    db.commit()
    return {"mensaje": "Empleado desactivado correctamente"}
