# ============================================================
# ROUTER INCIDENTES Y ACCIDENTES SST ENTERPRISE
# FASE 1.1.8.8.5 — DASHBOARD Y EXPORTACIONES
# Archivo: backend/app/routers/incidentes.py
# ============================================================

from datetime import date, datetime
from pathlib import Path
import io
import logging
import os
import uuid

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, joinedload

from app.auth.dependencies import require_roles, require_permission
from app.core.default_permissions import PERM_REGISTROS_ELIMINAR, PERM_REPORTES_EXPORTAR
from app.core.file_security import validate_upload
from app.database import get_db
from app.models.archivo_sst import ArchivoSST
from app.models.area import Area
from app.models.cargo import Cargo
from app.models.empleado import Empleado
from app.models.empresa import Empresa
from app.models.capa import CapaSST
from app.models.incidente import IncidenteAccidenteSST, IncidenteLesionadoSST, IncidenteTestigoSST
from app.models.sede import Sede
from app.schemas.incidente_schema import (
    IncidenteCreate,
    IncidenteDashboardResponse,
    IncidenteResponse,
    IncidenteUpdate,
    IncidenteInvestigacionUpdate,
    IncidenteArbolCausasUpdate,
    IncidenteCierreInvestigacionRequest,
    LesionadoCreate,
    LesionadoResponse,
    LesionadoUpdate,
    TestigoCreate,
    TestigoResponse,
    TestigoUpdate,
)

router = APIRouter(prefix="/incidentes", tags=["Incidentes y Accidentes SST Enterprise"])
ROLES_SST = ["SUPER_ADMIN", "ADMIN_EMPRESA", "RESPONSABLE_SST"]
EXPORTAR_REPORTES = require_permission(PERM_REPORTES_EXPORTAR)
ELIMINAR_REGISTROS = require_permission(PERM_REGISTROS_ELIMINAR)
logger = logging.getLogger("app.incidentes")

UPLOAD_ROOT = Path(os.getenv("UPLOAD_DIR", "app/uploads")).resolve()
INCIDENTES_UPLOAD_DIR = UPLOAD_ROOT / "incidentes"
INCIDENTES_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXT = {"pdf", "jpg", "jpeg", "png", "webp", "doc", "docx", "xls", "xlsx", "csv"}
IMAGE_EXT = {"jpg", "jpeg", "png", "webp"}
MAX_UPLOAD_MB = 25


def _upper(value, default=None):
    if value is None:
        return default
    value = str(value).strip().upper()
    return value if value else default


def _es_super_admin(usuario) -> bool:
    return str(getattr(usuario, "rol", "") or "").upper() == "SUPER_ADMIN"


def _empresa_usuario_id(usuario) -> int | None:
    return getattr(usuario, "empresa_id", None)


def _validar_empresa_usuario(usuario, empresa_id: int | None) -> None:
    if not empresa_id or _es_super_admin(usuario):
        return
    usuario_empresa_id = _empresa_usuario_id(usuario)
    if usuario_empresa_id is not None and int(usuario_empresa_id) != int(empresa_id):
        raise HTTPException(status_code=403, detail="No tiene permisos sobre esta empresa")


def _filtrar_empresa_usuario(query, usuario, model):
    if usuario is None or _es_super_admin(usuario):
        return query
    usuario_empresa_id = _empresa_usuario_id(usuario)
    if usuario_empresa_id is None:
        raise HTTPException(status_code=403, detail="Usuario sin empresa asignada para operación multiempresa")
    return query.filter(model.empresa_id == usuario_empresa_id)


def _public_upload_url(file_path: Path) -> str:
    try:
        rel = file_path.resolve().relative_to(UPLOAD_ROOT)
        return "/uploads/" + rel.as_posix()
    except Exception:
        return "/uploads/incidentes/" + file_path.name


def _thumb_preview_urls(url: str, mime_type: str | None = None):
    if not url or not (mime_type or "").startswith("image/"):
        return None, None
    if "." not in url:
        return None, None
    stem, ext = url.rsplit(".", 1)
    return f"{stem}_thumb.{ext}", f"{stem}_preview.{ext}"


def _optimizar_imagenes(content: bytes, extension: str):
    try:
        from PIL import Image
        image = Image.open(io.BytesIO(content))
        if image.mode not in ("RGB", "RGBA"):
            image = image.convert("RGB")

        def guardar(max_size, quality):
            img = image.copy()
            img.thumbnail(max_size)
            out = io.BytesIO()
            img.save(out, format="WEBP", quality=quality, method=6, optimize=True)
            return out.getvalue()

        return {
            "main": guardar((1920, 1920), 82),
            "preview": guardar((1100, 1100), 76),
            "thumb": guardar((160, 160), 70),
            "extension": "webp",
            "mime_type": "image/webp",
        }
    except Exception:
        mime = "image/jpeg" if extension.lower() in ["jpg", "jpeg"] else f"image/{extension.lower()}"
        return {"main": content, "preview": None, "thumb": None, "extension": extension.lower(), "mime_type": mime}


def _optimizar_pdf_bytes(content: bytes) -> bytes:
    try:
        import pikepdf
        src = io.BytesIO(content)
        out = io.BytesIO()
        with pikepdf.Pdf.open(src) as pdf:
            pdf.save(out, compress_streams=True, object_stream_mode=pikepdf.ObjectStreamMode.generate, linearize=True)
        optimized = out.getvalue()
        return optimized if len(optimized) < len(content) else content
    except Exception:
        return content


def _guardar_upload(upload: UploadFile) -> tuple[Path, str, str, str, int]:
    validation = validate_upload(upload, allowed_extensions={f".{item}" for item in ALLOWED_EXT}, max_size_mb=MAX_UPLOAD_MB)
    original = validation.safe_filename or "evidencia_incidente"
    extension = validation.extension.lstrip(".")
    content = validation.content

    if extension in IMAGE_EXT:
        data = _optimizar_imagenes(content, extension)
        filename = f"{uuid.uuid4().hex}.{data['extension']}"
        path = INCIDENTES_UPLOAD_DIR / filename
        path.write_bytes(data["main"])
        if data.get("preview"):
            (INCIDENTES_UPLOAD_DIR / filename.replace(f".{data['extension']}", f"_preview.{data['extension']}")).write_bytes(data["preview"])
        if data.get("thumb"):
            (INCIDENTES_UPLOAD_DIR / filename.replace(f".{data['extension']}", f"_thumb.{data['extension']}")).write_bytes(data["thumb"])
        return path, original, filename, data["mime_type"], len(data["main"])

    filename = f"{uuid.uuid4().hex}.{extension}"
    path = INCIDENTES_UPLOAD_DIR / filename
    if extension == "pdf":
        content = _optimizar_pdf_bytes(content)
    path.write_bytes(content)
    return path, original, filename, validation.mime_type, len(content)

def _archivo_to_dict(archivo: ArchivoSST):
    thumb, preview = _thumb_preview_urls(archivo.url, archivo.mime_type)
    return {
        "id": archivo.id,
        "empresa_id": archivo.empresa_id,
        "usuario_id": archivo.usuario_id,
        "tipo": archivo.tipo,
        "nombre_original": archivo.nombre_original,
        "nombre_archivo": archivo.nombre_archivo,
        "ruta": archivo.ruta,
        "url": archivo.url,
        "thumb_url": thumb,
        "preview_url": preview,
        "extension": archivo.extension,
        "mime_type": archivo.mime_type,
        "tamano_bytes": archivo.tamano_bytes,
        "modulo": archivo.modulo,
        "referencia_id": archivo.referencia_id,
        "descripcion": archivo.descripcion,
        "activo": archivo.activo,
        "fecha_creacion": archivo.fecha_creacion,
    }


def _validar_empresa(db: Session, empresa_id: int):
    empresa = db.query(Empresa).filter(Empresa.id == empresa_id).first()
    if not empresa:
        raise HTTPException(status_code=404, detail="Empresa no encontrada")
    return empresa


def _validar_opcional(db: Session, model, item_id: int | None, label: str):
    if not item_id:
        return None
    item = db.query(model).filter(model.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail=f"{label} no encontrado")
    return item


def _obtener_incidente_db(db: Session, incidente_id: int, usuario=None) -> IncidenteAccidenteSST:
    item = db.query(IncidenteAccidenteSST).filter(IncidenteAccidenteSST.id == incidente_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Incidente/accidente no encontrado")
    _validar_empresa_usuario(usuario, item.empresa_id)
    return item


INVESTIGACION_FIELDS = [
    "equipo_investigador", "investigador_lider", "fecha_investigacion",
    "metodologia_investigacion", "estado_investigacion", "descripcion_hechos",
    "agente_material", "mecanismo_evento", "tipo_contacto", "acto_inseguro",
    "condicion_insegura", "causa_inmediata", "causa_basica", "causa_raiz",
    "porque_1", "porque_2", "porque_3", "porque_4", "porque_5",
    "factores_personales", "factores_trabajo", "factores_organizacionales",
    "causas_directas", "causas_indirectas", "arbol_causas",
    "controles_existentes", "controles_recomendados", "plan_investigacion",
    "conclusion_investigacion", "recomendaciones_investigacion",
    "investigacion_cerrada", "fecha_cierre_investigacion",
]


def _investigacion_to_dict(item: IncidenteAccidenteSST) -> dict:
    data = {field: getattr(item, field, None) for field in INVESTIGACION_FIELDS}
    data.update({
        "id": item.id,
        "codigo": item.codigo,
        "titulo": item.titulo,
        "tipo_evento": item.tipo_evento,
        "clasificacion": item.clasificacion,
        "estado": item.estado,
        "severidad": item.severidad,
        "requiere_capa": item.requiere_capa,
        "capa_id": item.capa_id,
        "trazabilidad": item.trazabilidad,
    })
    return data


def _query_incidentes(db: Session, empresa_id=None, sede_id=None, area_id=None, cargo_id=None, empleado_id=None, tipo_evento=None, clasificacion=None, estado=None, severidad=None, q=None, usuario=None):
    query = db.query(IncidenteAccidenteSST).options(
        joinedload(IncidenteAccidenteSST.empresa),
        joinedload(IncidenteAccidenteSST.sede),
        joinedload(IncidenteAccidenteSST.area),
        joinedload(IncidenteAccidenteSST.cargo),
        joinedload(IncidenteAccidenteSST.empleado),
    )
    query = _filtrar_empresa_usuario(query, usuario, IncidenteAccidenteSST)
    if empresa_id:
        _validar_empresa_usuario(usuario, empresa_id)
        query = query.filter(IncidenteAccidenteSST.empresa_id == empresa_id)
    if sede_id:
        query = query.filter(IncidenteAccidenteSST.sede_id == sede_id)
    if area_id:
        query = query.filter(IncidenteAccidenteSST.area_id == area_id)
    if cargo_id:
        query = query.filter(IncidenteAccidenteSST.cargo_id == cargo_id)
    if empleado_id:
        query = query.filter(IncidenteAccidenteSST.empleado_id == empleado_id)
    if tipo_evento:
        query = query.filter(func.upper(IncidenteAccidenteSST.tipo_evento) == tipo_evento.upper().strip())
    if clasificacion:
        query = query.filter(func.upper(IncidenteAccidenteSST.clasificacion) == clasificacion.upper().strip())
    if estado:
        query = query.filter(func.upper(IncidenteAccidenteSST.estado) == estado.upper().strip())
    if severidad:
        query = query.filter(func.upper(IncidenteAccidenteSST.severidad) == severidad.upper().strip())
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(or_(IncidenteAccidenteSST.codigo.ilike(like), IncidenteAccidenteSST.titulo.ilike(like), IncidenteAccidenteSST.descripcion.ilike(like), IncidenteAccidenteSST.lugar.ilike(like)))
    return query.order_by(IncidenteAccidenteSST.id.desc())


def _incidente_to_response(db: Session, item: IncidenteAccidenteSST):
    data = IncidenteResponse.model_validate(item)
    data.empresa_nombre = item.empresa.nombre if item.empresa else None
    data.sede_nombre = item.sede.nombre if item.sede else None
    data.area_nombre = item.area.nombre if item.area else None
    data.cargo_nombre = item.cargo.nombre if item.cargo else None
    if item.empleado:
        nombres = getattr(item.empleado, "nombres", "") or ""
        apellidos = getattr(item.empleado, "apellidos", "") or ""
        data.empleado_nombre = f"{nombres} {apellidos}".strip()
        data.empleado_documento = getattr(item.empleado, "documento", None)
    data.total_lesionados = db.query(func.count(IncidenteLesionadoSST.id)).filter(IncidenteLesionadoSST.incidente_id == item.id, IncidenteLesionadoSST.activo.is_(True)).scalar() or 0
    data.total_testigos = db.query(func.count(IncidenteTestigoSST.id)).filter(IncidenteTestigoSST.incidente_id == item.id, IncidenteTestigoSST.activo.is_(True)).scalar() or 0
    data.total_evidencias = db.query(func.count(ArchivoSST.id)).filter(ArchivoSST.modulo == "INCIDENTES", ArchivoSST.referencia_id == item.id, ArchivoSST.activo.is_(True)).scalar() or 0
    if item.fecha_evento:
        data.dias_desde_evento = (date.today() - item.fecha_evento).days
        data.vencido = data.dias_desde_evento > 2 and item.estado not in ["CERRADO", "ANULADO"]
    return data


@router.get("/", response_model=list[IncidenteResponse])
def listar_incidentes(empresa_id: int | None = Query(default=None), sede_id: int | None = Query(default=None), area_id: int | None = Query(default=None), cargo_id: int | None = Query(default=None), empleado_id: int | None = Query(default=None), tipo_evento: str | None = Query(default=None), clasificacion: str | None = Query(default=None), estado: str | None = Query(default=None), severidad: str | None = Query(default=None), q: str | None = Query(default=None), db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    items = _query_incidentes(db, empresa_id, sede_id, area_id, cargo_id, empleado_id, tipo_evento, clasificacion, estado, severidad, q, usuario=usuario).filter(IncidenteAccidenteSST.activo.is_(True)).all()
    return [_incidente_to_response(db, item) for item in items]


@router.get("/dashboard/resumen", response_model=IncidenteDashboardResponse)
def dashboard_incidentes(empresa_id: int | None = Query(default=None), sede_id: int | None = Query(default=None), area_id: int | None = Query(default=None), db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    items = _query_incidentes(db, empresa_id, sede_id, area_id, usuario=usuario).filter(IncidenteAccidenteSST.activo.is_(True)).all()
    ids = [i.id for i in items]
    total = len(items)
    incidentes = sum(1 for i in items if i.tipo_evento == "INCIDENTE")
    accidentes = sum(1 for i in items if i.tipo_evento == "ACCIDENTE")
    graves = sum(1 for i in items if i.clasificacion == "ACCIDENTE_GRAVE")
    mortales = sum(1 for i in items if i.clasificacion == "ACCIDENTE_MORTAL")
    abiertos = sum(1 for i in items if i.estado not in ["CERRADO", "ANULADO"])
    cerrados = sum(1 for i in items if i.estado == "CERRADO")
    vencidos = sum(1 for i in items if i.fecha_evento and (date.today() - i.fecha_evento).days > 2 and i.estado not in ["CERRADO", "ANULADO"])
    cumplimiento = round((cerrados / total) * 100, 1) if total else 0
    total_lesionados = db.query(func.count(IncidenteLesionadoSST.id)).filter(IncidenteLesionadoSST.incidente_id.in_(ids), IncidenteLesionadoSST.activo.is_(True)).scalar() if ids else 0
    total_testigos = db.query(func.count(IncidenteTestigoSST.id)).filter(IncidenteTestigoSST.incidente_id.in_(ids), IncidenteTestigoSST.activo.is_(True)).scalar() if ids else 0
    total_evidencias = db.query(func.count(ArchivoSST.id)).filter(ArchivoSST.modulo == "INCIDENTES", ArchivoSST.referencia_id.in_(ids), ArchivoSST.activo.is_(True)).scalar() if ids else 0

    def conteo(attr):
        data = {}
        for item in items:
            key = attr(item) or "Sin dato"
            data[key] = data.get(key, 0) + 1
        return [{"name": k, "value": v} for k, v in sorted(data.items(), key=lambda x: x[1], reverse=True)[:8]]

    score = mortales * 60 + graves * 35 + vencidos * 20 + abiertos * 8
    semaforo = "ROJO" if score >= 60 else "AMARILLO" if score >= 20 else "VERDE"
    recomendaciones = []
    if mortales or graves:
        recomendaciones.append("Priorizar investigación inmediata de accidentes graves o mortales.")
    if vencidos:
        recomendaciones.append("Cerrar investigaciones con más de 48 horas abiertas o justificar seguimiento.")
    if abiertos:
        recomendaciones.append("Registrar lesionados, testigos y evidencias de soporte.")
    if not recomendaciones:
        recomendaciones.append("Gestión estable. Mantener investigación oportuna y seguimiento preventivo.")

    return {
        "kpis": {"total": total, "incidentes": incidentes, "accidentes": accidentes, "graves": graves, "mortales": mortales, "abiertos": abiertos, "cerrados": cerrados, "vencidos": vencidos, "cumplimiento": cumplimiento, "lesionados": total_lesionados or 0, "testigos": total_testigos or 0, "evidencias": total_evidencias or 0, "score_riesgo": score, "semaforo": semaforo},
        "charts": {"por_tipo": conteo(lambda x: x.tipo_evento), "por_clasificacion": conteo(lambda x: x.clasificacion), "por_estado": conteo(lambda x: x.estado), "por_severidad": conteo(lambda x: x.severidad), "por_area": conteo(lambda x: x.area.nombre if x.area else "Sin área"), "por_consecuencia": conteo(lambda x: x.consecuencia or "Sin consecuencia")},
        "alertas": {"abiertos": abiertos, "graves": graves, "mortales": mortales, "vencidos": vencidos},
        "recomendaciones": recomendaciones,
    }


# ============================================================
# FASE 1.1.8.8.5 — DASHBOARD Y EXPORTACIONES INCIDENTES SST
# ============================================================

def _fmt(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d %H:%M" if isinstance(value, datetime) else "%Y-%m-%d")
    return str(value)


def _filename(name: str) -> str:
    return name.replace(" ", "_").replace("/", "-")


def _excel_response(workbook, filename: str):
    out = io.BytesIO()
    workbook.save(out)
    out.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="{_filename(filename)}"'}
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


def _pdf_response(elements, filename: str, title: str = "Reporte SST"):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except Exception as exc:
        logger.exception("Dependencia PDF no disponible para exportacion de incidentes")
        raise HTTPException(status_code=500, detail="No fue posible generar el PDF.") from exc

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=letter, rightMargin=1.2*cm, leftMargin=1.2*cm, topMargin=1.2*cm, bottomMargin=1.2*cm)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]), Spacer(1, 10)]
    story.extend(elements)
    doc.build(story)
    out.seek(0)
    headers = {"Content-Disposition": f'attachment; filename="{_filename(filename)}"'}
    return StreamingResponse(out, media_type="application/pdf", headers=headers)


def _table(data, col_widths=None):
    try:
        from reportlab.lib import colors
        from reportlab.platypus import Table, TableStyle
    except Exception as exc:
        logger.exception("Dependencia PDF no disponible para tabla de incidentes")
        raise HTTPException(status_code=500, detail="No fue posible generar el PDF.") from exc
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def _p(text, style="Normal"):
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph
    return Paragraph(str(text or ""), getSampleStyleSheet()[style])


def _incidentes_filtrados(db, empresa_id=None, sede_id=None, area_id=None, tipo_evento=None, clasificacion=None, estado=None, severidad=None, q=None, usuario=None):
    return _query_incidentes(db, empresa_id=empresa_id, sede_id=sede_id, area_id=area_id, tipo_evento=tipo_evento, clasificacion=clasificacion, estado=estado, severidad=severidad, q=q, usuario=usuario).filter(IncidenteAccidenteSST.activo.is_(True)).all()


@router.get("/exportaciones/excel-general")
def exportar_incidentes_excel_general(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    tipo_evento: str | None = Query(default=None),
    clasificacion: str | None = Query(default=None),
    estado: str | None = Query(default=None),
    severidad: str | None = Query(default=None),
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception as exc:
        logger.exception("Dependencia Excel no disponible para exportacion de incidentes")
        raise HTTPException(status_code=500, detail="No fue posible generar el Excel.") from exc
    items = _incidentes_filtrados(db, empresa_id, sede_id, area_id, tipo_evento, clasificacion, estado, severidad, q, usuario=usuario)
    wb = Workbook()
    ws = wb.active
    ws.title = "Incidentes SST"
    headers = ["Código", "Título", "Empresa", "Sede", "Área", "Fecha", "Tipo", "Clasificación", "Severidad", "Estado", "Lesionados", "Testigos", "Evidencias", "CAPA", "Causa raíz"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center")
    for item in items:
        les = db.query(func.count(IncidenteLesionadoSST.id)).filter(IncidenteLesionadoSST.incidente_id == item.id, IncidenteLesionadoSST.activo.is_(True)).scalar() or 0
        tes = db.query(func.count(IncidenteTestigoSST.id)).filter(IncidenteTestigoSST.incidente_id == item.id, IncidenteTestigoSST.activo.is_(True)).scalar() or 0
        evi = db.query(func.count(ArchivoSST.id)).filter(ArchivoSST.modulo == "INCIDENTES", ArchivoSST.referencia_id == item.id, ArchivoSST.activo.is_(True)).scalar() or 0
        ws.append([item.codigo, item.titulo, item.empresa.nombre if item.empresa else "", item.sede.nombre if item.sede else "", item.area.nombre if item.area else "", _fmt(item.fecha_evento), item.tipo_evento, item.clasificacion, item.severidad, item.estado, les, tes, evi, item.capa_id or "", item.causa_raiz or ""])
    widths = [18, 36, 30, 24, 24, 14, 16, 20, 14, 18, 12, 12, 12, 10, 50]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width if i <= 26 else 20
    return _excel_response(wb, "incidentes_accidentes_sst_general.xlsx")


@router.get("/exportaciones/pdf-general")
def exportar_incidentes_pdf_general(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    tipo_evento: str | None = Query(default=None),
    clasificacion: str | None = Query(default=None),
    estado: str | None = Query(default=None),
    severidad: str | None = Query(default=None),
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    items = _incidentes_filtrados(db, empresa_id, sede_id, area_id, tipo_evento, clasificacion, estado, severidad, q, usuario=usuario)
    data = [["Código", "Evento", "Fecha", "Tipo", "Clasificación", "Estado", "Severidad"]]
    for item in items[:80]:
        data.append([item.codigo, item.titulo, _fmt(item.fecha_evento), item.tipo_evento, item.clasificacion, item.estado, item.severidad])
    elements = [_table(data)]
    return _pdf_response(elements, "incidentes_accidentes_sst_general.pdf", "Reporte General de Incidentes y Accidentes SST")


@router.get("/exportaciones/dashboard-ejecutivo-pdf")
def exportar_dashboard_incidentes_pdf(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    dash = dashboard_incidentes(empresa_id=empresa_id, sede_id=sede_id, area_id=area_id, db=db, usuario=usuario)
    k = dash.get("kpis", {})
    data = [["Indicador", "Valor"], ["Total eventos", k.get("total", 0)], ["Incidentes", k.get("incidentes", 0)], ["Accidentes", k.get("accidentes", 0)], ["Graves", k.get("graves", 0)], ["Mortales", k.get("mortales", 0)], ["Abiertos", k.get("abiertos", 0)], ["Cerrados", k.get("cerrados", 0)], ["Lesionados", k.get("lesionados", 0)], ["Testigos", k.get("testigos", 0)], ["Evidencias", k.get("evidencias", 0)], ["Cumplimiento", f"{k.get('cumplimiento', 0)}%"], ["Semáforo", k.get("semaforo", "VERDE")]]
    elements = [_table(data), _p("Recomendaciones PRO", "Heading2")]
    for rec in dash.get("recomendaciones", []):
        elements.append(_p(f"• {rec}"))
    return _pdf_response(elements, "dashboard_ejecutivo_incidentes_sst.pdf", "Dashboard Ejecutivo de Incidentes y Accidentes SST")


def _elementos_detalle_incidente(db: Session, item: IncidenteAccidenteSST, titulo: str):
    from reportlab.platypus import Spacer
    les = db.query(IncidenteLesionadoSST).filter(IncidenteLesionadoSST.incidente_id == item.id, IncidenteLesionadoSST.activo.is_(True)).all()
    tes = db.query(IncidenteTestigoSST).filter(IncidenteTestigoSST.incidente_id == item.id, IncidenteTestigoSST.activo.is_(True)).all()
    elementos = []
    base = [["Campo", "Valor"], ["Código", item.codigo], ["Título", item.titulo], ["Empresa", item.empresa.nombre if item.empresa else ""], ["Fecha evento", _fmt(item.fecha_evento)], ["Hora", item.hora_evento or ""], ["Lugar", item.lugar or ""], ["Tipo", item.tipo_evento], ["Clasificación", item.clasificacion], ["Severidad", item.severidad], ["Estado", item.estado], ["Consecuencia", item.consecuencia or ""], ["Descripción", item.descripcion or ""]]
    elementos.append(_table(base))
    elementos.append(Spacer(1, 10))
    inv = [["Investigación", "Detalle"], ["Investigador líder", item.investigador_lider or ""], ["Metodología", item.metodologia_investigacion or ""], ["Descripción hechos", item.descripcion_hechos or ""], ["Acto inseguro", item.acto_inseguro or ""], ["Condición insegura", item.condicion_insegura or ""], ["Causa inmediata", item.causa_inmediata or ""], ["Causa básica", item.causa_basica or ""], ["Causa raíz", item.causa_raiz or ""], ["Controles recomendados", item.controles_recomendados or ""], ["Conclusión", item.conclusion_investigacion or ""]]
    elementos.append(_p("Investigación y árbol de causas", "Heading2"))
    elementos.append(_table(inv))
    elementos.append(Spacer(1, 10))
    if les:
        elementos.append(_p("Lesionados", "Heading2"))
        elementos.append(_table([["Nombre", "Documento", "Parte afectada", "Gravedad", "Días incapacidad"]] + [[l.nombre, l.documento or "", l.parte_cuerpo_afectada or "", l.gravedad or "", l.dias_incapacidad or 0] for l in les]))
        elementos.append(Spacer(1, 10))
    if tes:
        elementos.append(_p("Testigos", "Heading2"))
        elementos.append(_table([["Nombre", "Documento", "Cargo", "Declaración"]] + [[t.nombre, t.documento or "", t.cargo or "", t.declaracion or ""] for t in tes]))
        elementos.append(Spacer(1, 10))
    if item.trazabilidad:
        elementos.append(_p("Trazabilidad", "Heading2"))
        elementos.append(_p(item.trazabilidad.replace("\n", "<br/>")))
    return elementos


@router.get("/{incidente_id}/pdf-individual")
def exportar_incidente_pdf_individual(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    return _pdf_response(_elementos_detalle_incidente(db, item, "Incidente"), f"incidente_accidente_{item.codigo}.pdf", f"Incidente / Accidente SST {item.codigo}")


@router.get("/{incidente_id}/acta-investigacion-pdf")
def exportar_acta_investigacion_pdf(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    return _pdf_response(_elementos_detalle_incidente(db, item, "Acta"), f"acta_investigacion_{item.codigo}.pdf", f"Acta Oficial de Investigación SST {item.codigo}")


@router.get("/{incidente_id}/informe-incidente-pdf")
def exportar_informe_incidente_pdf(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    if item.tipo_evento != "INCIDENTE":
        raise HTTPException(status_code=400, detail="El evento no está clasificado como INCIDENTE")
    return _pdf_response(_elementos_detalle_incidente(db, item, "Informe Incidente"), f"informe_incidente_{item.codigo}.pdf", f"Informe de Incidente SST {item.codigo}")


@router.get("/{incidente_id}/informe-accidente-pdf")
def exportar_informe_accidente_pdf(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    if item.tipo_evento != "ACCIDENTE":
        raise HTTPException(status_code=400, detail="El evento no está clasificado como ACCIDENTE")
    return _pdf_response(_elementos_detalle_incidente(db, item, "Informe Accidente"), f"informe_accidente_{item.codigo}.pdf", f"Informe de Accidente SST {item.codigo}")


@router.get("/{incidente_id}", response_model=IncidenteResponse)
def obtener_incidente(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    return _incidente_to_response(db, _obtener_incidente_db(db, incidente_id, usuario))


@router.post("/", response_model=IncidenteResponse)
def crear_incidente(data: IncidenteCreate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    _validar_empresa_usuario(usuario, data.empresa_id)
    _validar_empresa(db, data.empresa_id)
    _validar_opcional(db, Sede, data.sede_id, "Sede")
    _validar_opcional(db, Area, data.area_id, "Área")
    _validar_opcional(db, Cargo, data.cargo_id, "Cargo")
    _validar_opcional(db, Empleado, data.empleado_id, "Empleado")
    existe = db.query(IncidenteAccidenteSST).filter(IncidenteAccidenteSST.empresa_id == data.empresa_id, func.upper(IncidenteAccidenteSST.codigo) == data.codigo.upper()).first()
    if existe:
        raise HTTPException(status_code=400, detail="Ya existe un incidente/accidente con ese código para la empresa")
    payload = data.model_dump()
    payload["usuario_id"] = payload.get("usuario_id") or getattr(usuario, "id", None)
    ahora = datetime.utcnow()
    payload["trazabilidad"] = (payload.get("trazabilidad") + "\n" if payload.get("trazabilidad") else "") + f"[{ahora.isoformat()}] Evento creado por usuario {getattr(usuario, 'id', '')}. Estado {payload.get('estado', 'REPORTADO')}."
    item = IncidenteAccidenteSST(**payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return obtener_incidente(item.id, db, usuario)


@router.put("/{incidente_id}", response_model=IncidenteResponse)
def actualizar_incidente(incidente_id: int, data: IncidenteUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(item, key, value)
    ahora = datetime.utcnow()
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + f"[{ahora.isoformat()}] Evento actualizado por usuario {getattr(usuario, 'id', '')}."
    db.commit()
    db.refresh(item)
    return obtener_incidente(item.id, db, usuario)


@router.delete("/{incidente_id}")
def eliminar_incidente(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    item.activo = False
    item.estado = "ANULADO"
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + f"[{datetime.utcnow().isoformat()}] Evento anulado por usuario {getattr(usuario, 'id', '')}."
    db.commit()
    return {"ok": True, "message": "Incidente/accidente anulado"}


@router.get("/{incidente_id}/investigacion")
def obtener_investigacion_incidente(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    return _investigacion_to_dict(item)


@router.put("/{incidente_id}/investigacion", response_model=IncidenteResponse)
def actualizar_investigacion_incidente(incidente_id: int, data: IncidenteInvestigacionUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    payload = data.model_dump(exclude_unset=True)
    for key, value in payload.items():
        setattr(item, key, value)
    if item.estado == "REPORTADO":
        item.estado = "EN_INVESTIGACION"
    if not item.estado_investigacion or item.estado_investigacion == "PENDIENTE":
        item.estado_investigacion = "EN_PROCESO"
    item.requiere_investigacion = True
    ahora = datetime.utcnow()
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + f"[{ahora.isoformat()}] Investigación actualizada por usuario {getattr(usuario, 'id', '')}."
    db.commit()
    db.refresh(item)
    return obtener_incidente(item.id, db, usuario)


@router.get("/{incidente_id}/arbol-causas")
def obtener_arbol_causas_incidente(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    return {
        "incidente_id": item.id,
        "codigo": item.codigo,
        "descripcion_hechos": item.descripcion_hechos,
        "arbol_causas": item.arbol_causas,
        "causa_inmediata": item.causa_inmediata,
        "causa_basica": item.causa_basica,
        "causa_raiz": item.causa_raiz,
        "porque_1": item.porque_1,
        "porque_2": item.porque_2,
        "porque_3": item.porque_3,
        "porque_4": item.porque_4,
        "porque_5": item.porque_5,
        "factores_personales": item.factores_personales,
        "factores_trabajo": item.factores_trabajo,
        "factores_organizacionales": item.factores_organizacionales,
        "controles_recomendados": item.controles_recomendados,
    }


@router.put("/{incidente_id}/arbol-causas", response_model=IncidenteResponse)
def actualizar_arbol_causas_incidente(incidente_id: int, data: IncidenteArbolCausasUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(item, key, value)
    if item.estado == "REPORTADO":
        item.estado = "EN_INVESTIGACION"
    item.estado_investigacion = "ANALISIS_CAUSAL"
    ahora = datetime.utcnow()
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + f"[{ahora.isoformat()}] Árbol de causas actualizado por usuario {getattr(usuario, 'id', '')}."
    db.commit()
    db.refresh(item)
    return obtener_incidente(item.id, db, usuario)


def _prioridad_capa_por_incidente(item: IncidenteAccidenteSST) -> str:
    if item.clasificacion == "ACCIDENTE_MORTAL" or item.severidad == "CRITICA":
        return "CRITICA"
    if item.clasificacion == "ACCIDENTE_GRAVE" or item.severidad == "ALTA":
        return "ALTA"
    if item.tipo_evento == "ACCIDENTE" or item.severidad == "MEDIA":
        return "MEDIA"
    return "BAJA"


def _crear_capa_automatica_desde_incidente(db: Session, item: IncidenteAccidenteSST, usuario) -> CapaSST:
    if item.capa_id:
        existente = db.query(CapaSST).filter(CapaSST.id == item.capa_id, CapaSST.activo.is_(True)).first()
        if existente:
            return existente

    existente = db.query(CapaSST).filter(
        CapaSST.origen.in_(["INCIDENTE", "ACCIDENTE"]),
        CapaSST.observaciones.ilike(f"%Incidente/Accidente ID: {item.id}%"),
        CapaSST.activo.is_(True),
    ).first()
    if existente:
        item.capa_id = existente.id
        item.requiere_capa = True
        item.estado = "CON_CAPA"
        return existente

    consecutivo = (db.query(func.count(CapaSST.id)).filter(CapaSST.empresa_id == item.empresa_id).scalar() or 0) + 1
    prefijo = "ACC" if item.tipo_evento == "ACCIDENTE" else "INC"
    codigo = f"CAPA-{prefijo}-{item.id:04d}-{consecutivo:03d}"
    causa_raiz = item.causa_raiz or item.causa_basica or item.causa_inmediata or "Pendiente análisis de causa raíz del evento SST."
    accion_correctiva = item.controles_recomendados or item.accion_inmediata or "Definir e implementar controles correctivos derivados de la investigación."
    descripcion = (
        f"CAPA generada automáticamente desde {item.tipo_evento.lower()} SST {item.codigo}. "
        f"Clasificación: {item.clasificacion}. Descripción del evento: {item.descripcion}"
    )
    capa = CapaSST(
        empresa_id=item.empresa_id,
        sede_id=item.sede_id,
        area_id=item.area_id,
        cargo_id=item.cargo_id,
        empleado_id=item.empleado_id,
        usuario_id=getattr(usuario, "id", None),
        codigo=codigo,
        titulo=f"CAPA por {item.tipo_evento.lower()} SST {item.codigo}",
        descripcion=descripcion,
        tipo_accion="CORRECTIVA" if item.tipo_evento == "ACCIDENTE" else "PREVENTIVA",
        origen=item.tipo_evento,
        prioridad=_prioridad_capa_por_incidente(item),
        estado="ABIERTA",
        responsable=item.investigador_lider or "Responsable SST",
        fecha_apertura=date.today(),
        fecha_compromiso=None,
        avance=0,
        causa_raiz=causa_raiz,
        porque_1=item.porque_1,
        porque_2=item.porque_2,
        porque_3=item.porque_3,
        porque_4=item.porque_4,
        porque_5=item.porque_5,
        accion_inmediata=item.accion_inmediata,
        accion_correctiva=accion_correctiva,
        accion_preventiva=item.recomendaciones_investigacion or "Implementar seguimiento preventivo para evitar recurrencia.",
        observaciones=f"CAPA generada desde Incidente/Accidente ID: {item.id}. Código evento: {item.codigo}.",
        trazabilidad=f"[{datetime.utcnow().isoformat()}] CAPA generada automáticamente desde {item.tipo_evento} {item.codigo} por usuario {getattr(usuario, 'id', '')}.",
        activo=True,
    )
    db.add(capa)
    db.flush()
    item.capa_id = capa.id
    item.requiere_capa = True
    item.estado = "CON_CAPA"
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + f"[{datetime.utcnow().isoformat()}] CAPA automática generada: {capa.codigo}."
    return capa


@router.post("/{incidente_id}/cerrar-investigacion", response_model=IncidenteResponse)
def cerrar_investigacion_incidente(incidente_id: int, data: IncidenteCierreInvestigacionRequest, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    faltantes = []
    if not item.descripcion_hechos:
        faltantes.append("descripción de hechos")
    if not item.causa_raiz:
        faltantes.append("causa raíz")
    if not item.accion_inmediata and not item.controles_recomendados:
        faltantes.append("controles o acción inmediata")
    if faltantes:
        raise HTTPException(status_code=400, detail=f"No se puede cerrar la investigación. Faltan: {', '.join(faltantes)}")
    ahora = datetime.utcnow()
    item.conclusion_investigacion = data.conclusion_investigacion
    item.recomendaciones_investigacion = data.recomendaciones_investigacion
    item.requiere_capa = data.requiere_capa
    item.estado_investigacion = "CERRADA"
    item.investigacion_cerrada = True
    item.fecha_cierre_investigacion = ahora
    item.estado = "CON_CAPA" if data.requiere_capa else "CERRADO"
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + f"[{ahora.isoformat()}] Investigación cerrada por usuario {getattr(usuario, 'id', '')}. {data.observacion or ''}"
    db.commit()
    db.refresh(item)
    return obtener_incidente(item.id, db, usuario)


@router.post("/{incidente_id}/generar-capa", response_model=IncidenteResponse)
def generar_capa_desde_incidente(
    incidente_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    item = _obtener_incidente_db(db, incidente_id, usuario)
    if item.estado == "ANULADO" or item.activo is False:
        raise HTTPException(status_code=400, detail="No se puede generar CAPA desde un incidente anulado")
    if not item.causa_raiz and not item.causa_basica and not item.causa_inmediata:
        raise HTTPException(status_code=400, detail="Registra al menos causa inmediata, básica o raíz antes de generar la CAPA.")
    capa = _crear_capa_automatica_desde_incidente(db, item, usuario)
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + f"[{datetime.utcnow().isoformat()}] Integración CAPA ejecutada. CAPA vinculada: {capa.codigo}."
    db.commit()
    db.refresh(item)
    return obtener_incidente(item.id, db, usuario)


@router.get("/{incidente_id}/lesionados", response_model=list[LesionadoResponse])
def listar_lesionados(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    return db.query(IncidenteLesionadoSST).filter(IncidenteLesionadoSST.incidente_id == incidente_id, IncidenteLesionadoSST.activo.is_(True)).order_by(IncidenteLesionadoSST.id.desc()).all()


@router.post("/{incidente_id}/lesionados", response_model=LesionadoResponse)
def crear_lesionado(incidente_id: int, data: LesionadoCreate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    incidente = _obtener_incidente_db(db, incidente_id, usuario)
    payload = data.model_dump()
    payload["incidente_id"] = incidente.id
    payload["empresa_id"] = incidente.empresa_id
    item = IncidenteLesionadoSST(**payload)
    db.add(item)
    incidente.trazabilidad = (incidente.trazabilidad + "\n" if incidente.trazabilidad else "") + f"[{datetime.utcnow().isoformat()}] Lesionado registrado: {item.nombre}."
    db.commit()
    db.refresh(item)
    return item


@router.put("/lesionados/{lesionado_id}", response_model=LesionadoResponse)
def actualizar_lesionado(lesionado_id: int, data: LesionadoUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = db.query(IncidenteLesionadoSST).filter(IncidenteLesionadoSST.id == lesionado_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Lesionado no encontrado")
    _validar_empresa_usuario(usuario, item.empresa_id)
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(item, key, value)
    db.commit()
    db.refresh(item)
    return item


@router.delete("/lesionados/{lesionado_id}")
def eliminar_lesionado(lesionado_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    item = db.query(IncidenteLesionadoSST).filter(IncidenteLesionadoSST.id == lesionado_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Lesionado no encontrado")
    _validar_empresa_usuario(usuario, item.empresa_id)
    item.activo = False
    db.commit()
    return {"ok": True, "message": "Lesionado desactivado"}


@router.get("/{incidente_id}/testigos", response_model=list[TestigoResponse])
def listar_testigos(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    return db.query(IncidenteTestigoSST).filter(IncidenteTestigoSST.incidente_id == incidente_id, IncidenteTestigoSST.activo.is_(True)).order_by(IncidenteTestigoSST.id.desc()).all()


@router.post("/{incidente_id}/testigos", response_model=TestigoResponse)
def crear_testigo(incidente_id: int, data: TestigoCreate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    incidente = _obtener_incidente_db(db, incidente_id, usuario)
    payload = data.model_dump()
    payload["incidente_id"] = incidente.id
    payload["empresa_id"] = incidente.empresa_id
    if payload.get("firma"):
        payload["firma_fecha"] = datetime.utcnow()
    item = IncidenteTestigoSST(**payload)
    db.add(item)
    incidente.trazabilidad = (incidente.trazabilidad + "\n" if incidente.trazabilidad else "") + f"[{datetime.utcnow().isoformat()}] Testigo registrado: {item.nombre}."
    db.commit()
    db.refresh(item)
    return item


@router.put("/testigos/{testigo_id}", response_model=TestigoResponse)
def actualizar_testigo(testigo_id: int, data: TestigoUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = db.query(IncidenteTestigoSST).filter(IncidenteTestigoSST.id == testigo_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Testigo no encontrado")
    _validar_empresa_usuario(usuario, item.empresa_id)
    payload = data.model_dump(exclude_unset=True)
    if payload.get("firma"):
        payload["firma_fecha"] = datetime.utcnow()
    for key, value in payload.items():
        setattr(item, key, value)
    db.commit()
    db.refresh(item)
    return item


@router.delete("/testigos/{testigo_id}")
def eliminar_testigo(testigo_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    item = db.query(IncidenteTestigoSST).filter(IncidenteTestigoSST.id == testigo_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Testigo no encontrado")
    _validar_empresa_usuario(usuario, item.empresa_id)
    item.activo = False
    db.commit()
    return {"ok": True, "message": "Testigo desactivado"}


@router.get("/{incidente_id}/evidencias")
def listar_evidencias(incidente_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    archivos = db.query(ArchivoSST).filter(ArchivoSST.modulo == "INCIDENTES", ArchivoSST.referencia_id == incidente_id, ArchivoSST.activo.is_(True)).order_by(ArchivoSST.fecha_creacion.desc()).all()
    return [_archivo_to_dict(a) for a in archivos]


@router.post("/{incidente_id}/evidencias")
def subir_evidencia(incidente_id: int, tipo_evidencia: str = Form(default="EVIDENCIA_EVENTO"), descripcion: str = Form(default=""), archivo: UploadFile = File(...), db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    incidente = _obtener_incidente_db(db, incidente_id, usuario)
    path, original, filename, mime_type, size = _guardar_upload(archivo)
    registro = ArchivoSST(empresa_id=incidente.empresa_id, usuario_id=getattr(usuario, "id", None), tipo=(tipo_evidencia or "EVIDENCIA_EVENTO").upper().strip(), nombre_original=original, nombre_archivo=filename, ruta=str(path), url=_public_upload_url(path), extension=filename.rsplit(".", 1)[-1].lower(), mime_type=mime_type, tamano_bytes=size, modulo="INCIDENTES", referencia_id=incidente.id, descripcion=descripcion or "Evidencia de incidente/accidente SST", activo=True)
    db.add(registro)
    incidente.trazabilidad = (incidente.trazabilidad + "\n" if incidente.trazabilidad else "") + f"[{datetime.utcnow().isoformat()}] Evidencia cargada: {original}."
    db.commit()
    db.refresh(registro)
    return _archivo_to_dict(registro)


@router.delete("/{incidente_id}/evidencias/{archivo_id}")
def eliminar_evidencia(incidente_id: int, archivo_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    archivo = db.query(ArchivoSST).filter(ArchivoSST.id == archivo_id, ArchivoSST.modulo == "INCIDENTES", ArchivoSST.referencia_id == incidente_id).first()
    if not archivo:
        raise HTTPException(status_code=404, detail="Evidencia no encontrada")
    archivo.activo = False
    db.commit()
    return {"ok": True, "message": "Evidencia desactivada"}
