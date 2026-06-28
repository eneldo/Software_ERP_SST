# ============================================================
# ROUTER INDICADORES SST BI EXECUTIVE - ERP SST PRO
# FASE 1.1.18.1 — NÚCLEO INDICADORES SST
# Archivo: backend/app/routers/indicadores_sst.py
# ============================================================

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.auth.dependencies import require_roles
from app.database import get_db
from app.models.area import Area
from app.models.auditoria_sst import AuditoriaSST, AuditoriaHallazgoSST
from app.models.capa import CapaSST
from app.models.capacitacion import CapacitacionSST, CapacitacionAsistenteSST
from app.models.cargo import Cargo
from app.models.empleado import Empleado
from app.models.empresa import Empresa
from app.models.epp import EPPEntrega
from app.models.examen_medico import ExamenMedico
from app.models.incidente import IncidenteAccidenteSST
from app.models.indicador_sst import IndicadorSST
from app.models.inspeccion import InspeccionSST, InspeccionHallazgoSST
from app.models.sede import Sede
from app.schemas.indicador_sst_schema import (
    IndicadorSSTCreate,
    IndicadorSSTResponse,
    IndicadorSSTUpdate,
    IndicadoresDashboardResponse,
    KPIAutomaticoResponse,
)

router = APIRouter(prefix="/indicadores", tags=["Indicadores SST BI Executive"])

ROLES_SST = ["SUPER_ADMIN", "ADMIN_EMPRESA", "RESPONSABLE_SST"]
ROLES_ADMIN = ["SUPER_ADMIN", "ADMIN_EMPRESA"]


# ============================================================
# HELPERS
# ============================================================

def normalizar_texto(valor: Any, upper: bool = False, defecto: str | None = None) -> str | None:
    if valor is None:
        return defecto
    if isinstance(valor, str):
        limpio = valor.strip()
    else:
        limpio = str(valor).strip()
    if not limpio:
        return defecto
    return limpio.upper() if upper else limpio


def to_float(valor: Any) -> float:
    try:
        if valor is None:
            return 0.0
        return float(valor)
    except Exception:
        return 0.0


def safe_pct(numerador: float, denominador: float) -> float:
    if denominador <= 0:
        return 0.0
    return round((numerador / denominador) * 100, 2)


def semaforo_por_cumplimiento(valor: float) -> str:
    if valor >= 90:
        return "VERDE"
    if valor >= 70:
        return "AMARILLO"
    return "ROJO"


def semaforo_por_riesgo(valor: float, bajo_es_mejor: bool = True) -> str:
    if bajo_es_mejor:
        if valor <= 0:
            return "VERDE"
        if valor <= 5:
            return "AMARILLO"
        return "ROJO"
    return semaforo_por_cumplimiento(valor)


def limpiar_query_indicadores(db: Session, empresa_id=None, sede_id=None, area_id=None, activo=True):
    query = db.query(IndicadorSST)
    if empresa_id:
        query = query.filter(IndicadorSST.empresa_id == empresa_id)
    if sede_id:
        query = query.filter(IndicadorSST.sede_id == sede_id)
    if area_id:
        query = query.filter(IndicadorSST.area_id == area_id)
    if activo is not None:
        query = query.filter(IndicadorSST.activo == activo)
    return query


def validar_empresa(db: Session, empresa_id: int) -> Empresa:
    empresa = db.query(Empresa).filter(Empresa.id == empresa_id).first()
    if not empresa:
        raise HTTPException(status_code=404, detail="Empresa no encontrada")
    return empresa


def validar_relaciones(db: Session, data: dict):
    empresa_id = data.get("empresa_id")
    if empresa_id:
        validar_empresa(db, empresa_id)

    sede_id = data.get("sede_id")
    if sede_id:
        sede = db.query(Sede).filter(Sede.id == sede_id).first()
        if not sede:
            raise HTTPException(status_code=404, detail="Sede no encontrada")
        if empresa_id and sede.empresa_id != empresa_id:
            raise HTTPException(status_code=400, detail="La sede no pertenece a la empresa seleccionada")

    area_id = data.get("area_id")
    if area_id:
        area = db.query(Area).filter(Area.id == area_id).first()
        if not area:
            raise HTTPException(status_code=404, detail="Área no encontrada")
        if empresa_id and area.empresa_id != empresa_id:
            raise HTTPException(status_code=400, detail="El área no pertenece a la empresa seleccionada")

    cargo_id = data.get("cargo_id")
    if cargo_id:
        cargo = db.query(Cargo).filter(Cargo.id == cargo_id).first()
        if not cargo:
            raise HTTPException(status_code=404, detail="Cargo no encontrado")
        if empresa_id and cargo.empresa_id != empresa_id:
            raise HTTPException(status_code=400, detail="El cargo no pertenece a la empresa seleccionada")


def indicador_to_response(ind: IndicadorSST) -> IndicadorSSTResponse:
    return IndicadorSSTResponse(
        id=ind.id,
        empresa_id=ind.empresa_id,
        sede_id=ind.sede_id,
        area_id=ind.area_id,
        cargo_id=ind.cargo_id,
        usuario_id=ind.usuario_id,
        empresa_nombre=ind.empresa.nombre if ind.empresa else None,
        sede_nombre=ind.sede.nombre if ind.sede else None,
        area_nombre=ind.area.nombre if ind.area else None,
        cargo_nombre=ind.cargo.nombre if ind.cargo else None,
        codigo=ind.codigo,
        nombre=ind.nombre,
        descripcion=ind.descripcion,
        categoria=ind.categoria,
        tipo_indicador=ind.tipo_indicador,
        origen_dato=ind.origen_dato,
        frecuencia=ind.frecuencia,
        formula=ind.formula,
        unidad=ind.unidad,
        meta=ind.meta,
        valor_actual=ind.valor_actual,
        resultado=ind.resultado,
        semaforo=ind.semaforo,
        tendencia=ind.tendencia,
        periodo_inicio=ind.periodo_inicio,
        periodo_fin=ind.periodo_fin,
        responsable=ind.responsable,
        fuente=ind.fuente,
        observaciones=ind.observaciones,
        activo=ind.activo,
        fecha_creacion=ind.fecha_creacion,
        fecha_actualizacion=ind.fecha_actualizacion,
    )


def contar(query) -> int:
    try:
        return int(query.count() or 0)
    except Exception:
        return 0


def calcular_kpis_automaticos(db: Session, empresa_id=None, sede_id=None, area_id=None) -> list[KPIAutomaticoResponse]:
    # Filtros base por empresa/sede/área cuando el modelo lo permite.
    empleados_q = db.query(Empleado).filter(Empleado.activo == True)
    inspecciones_q = db.query(InspeccionSST).filter(InspeccionSST.activo == True)
    hallazgos_q = db.query(InspeccionHallazgoSST).filter(InspeccionHallazgoSST.activo == True)
    capas_q = db.query(CapaSST).filter(CapaSST.activo == True)
    incidentes_q = db.query(IncidenteAccidenteSST).filter(IncidenteAccidenteSST.activo == True)
    capacitaciones_q = db.query(CapacitacionSST).filter(CapacitacionSST.activo == True)
    epp_q = db.query(EPPEntrega).filter(EPPEntrega.activo == True)
    examenes_q = db.query(ExamenMedico).filter(ExamenMedico.activo == True)
    auditorias_q = db.query(AuditoriaSST).filter(AuditoriaSST.activo == True)
    auditoria_hallazgos_q = db.query(AuditoriaHallazgoSST).filter(AuditoriaHallazgoSST.activo == True)

    if empresa_id:
        empleados_q = empleados_q.filter(Empleado.empresa_id == empresa_id)
        inspecciones_q = inspecciones_q.filter(InspeccionSST.empresa_id == empresa_id)
        hallazgos_q = hallazgos_q.filter(InspeccionHallazgoSST.empresa_id == empresa_id)
        capas_q = capas_q.filter(CapaSST.empresa_id == empresa_id)
        incidentes_q = incidentes_q.filter(IncidenteAccidenteSST.empresa_id == empresa_id)
        capacitaciones_q = capacitaciones_q.filter(CapacitacionSST.empresa_id == empresa_id)
        epp_q = epp_q.filter(EPPEntrega.empresa_id == empresa_id)
        auditorias_q = auditorias_q.filter(AuditoriaSST.empresa_id == empresa_id)
        auditoria_hallazgos_q = auditoria_hallazgos_q.filter(AuditoriaHallazgoSST.empresa_id == empresa_id)

    if sede_id:
        empleados_q = empleados_q.filter(Empleado.sede_id == sede_id)
        inspecciones_q = inspecciones_q.filter(InspeccionSST.sede_id == sede_id)
        capas_q = capas_q.filter(CapaSST.sede_id == sede_id)
        incidentes_q = incidentes_q.filter(IncidenteAccidenteSST.sede_id == sede_id)

    if area_id:
        empleados_q = empleados_q.filter(Empleado.area_id == area_id)
        inspecciones_q = inspecciones_q.filter(InspeccionSST.area_id == area_id)
        capas_q = capas_q.filter(CapaSST.area_id == area_id)
        incidentes_q = incidentes_q.filter(IncidenteAccidenteSST.area_id == area_id)

    empleados_total = contar(empleados_q)

    inspecciones_total = contar(inspecciones_q)
    inspecciones_cerradas = contar(inspecciones_q.filter(func.upper(InspeccionSST.estado).in_(["CERRADA", "EJECUTADA", "FINALIZADA"])))
    cumplimiento_inspecciones = safe_pct(inspecciones_cerradas, inspecciones_total)

    hallazgos_total = contar(hallazgos_q)
    hallazgos_cerrados = contar(hallazgos_q.filter(func.upper(InspeccionHallazgoSST.estado).in_(["CERRADO", "CERRADA", "FINALIZADO", "RESUELTO"])))
    hallazgos_abiertos = max(hallazgos_total - hallazgos_cerrados, 0)
    cierre_hallazgos = safe_pct(hallazgos_cerrados, hallazgos_total)

    capas_total = contar(capas_q)
    capas_cerradas = contar(capas_q.filter(func.upper(CapaSST.estado).in_(["CERRADA", "CERRADO", "FINALIZADA"])))
    capas_efectivas = contar(capas_q.filter(CapaSST.efectiva == True))
    cumplimiento_capa = safe_pct(capas_cerradas, capas_total)
    efectividad_capa = safe_pct(capas_efectivas, capas_cerradas if capas_cerradas else capas_total)

    incidentes_total = contar(incidentes_q)
    accidentes_total = contar(incidentes_q.filter(func.upper(IncidenteAccidenteSST.tipo_evento).like("%ACCIDENTE%")))
    graves = contar(incidentes_q.filter(func.upper(IncidenteAccidenteSST.severidad).in_(["ALTA", "GRAVE", "CRITICA", "CRÍTICA"])))
    mortales = contar(incidentes_q.filter(func.upper(IncidenteAccidenteSST.clasificacion).like("%MORTAL%")))
    abiertos_eventos = contar(incidentes_q.filter(func.upper(IncidenteAccidenteSST.estado).notin_(["CERRADO", "CERRADA", "FINALIZADO"])))

    capacitaciones_total = contar(capacitaciones_q)
    capacitaciones_ejecutadas = contar(capacitaciones_q.filter(func.upper(CapacitacionSST.estado).in_(["EJECUTADA", "FINALIZADA", "CERRADA"])))
    cumplimiento_capacitaciones = safe_pct(capacitaciones_ejecutadas, capacitaciones_total)

    asistentes_total = contar(db.query(CapacitacionAsistenteSST).filter(CapacitacionAsistenteSST.activo == True))
    cobertura_capacitacion = min(safe_pct(asistentes_total, empleados_total), 100.0)

    epp_entregas = contar(epp_q)
    cobertura_epp = min(safe_pct(epp_entregas, empleados_total), 100.0)

    examenes_total = contar(examenes_q)
    cobertura_examenes = min(safe_pct(examenes_total, empleados_total), 100.0)

    auditorias_total = contar(auditorias_q)
    auditorias_cerradas = contar(auditorias_q.filter(func.upper(AuditoriaSST.estado).in_(["CERRADA", "FINALIZADA", "EJECUTADA"])))
    cumplimiento_auditorias = safe_pct(auditorias_cerradas, auditorias_total)

    auditoria_hallazgos_total = contar(auditoria_hallazgos_q)
    auditoria_hallazgos_cerrados = contar(auditoria_hallazgos_q.filter(func.upper(AuditoriaHallazgoSST.estado).in_(["CERRADO", "CERRADA", "FINALIZADO"])))
    cierre_hallazgos_auditoria = safe_pct(auditoria_hallazgos_cerrados, auditoria_hallazgos_total)

    tasa_eventos = round((incidentes_total / empleados_total) * 100, 2) if empleados_total else 0.0

    datos = [
        ("KPI-INS-001", "Cumplimiento inspecciones SST", "INSPECCIONES", cumplimiento_inspecciones, 90, "%", "inspecciones_sst", "Inspecciones cerradas / inspecciones totales"),
        ("KPI-HAL-001", "Cierre de hallazgos SST", "HALLAZGOS", cierre_hallazgos, 90, "%", "inspecciones_hallazgos_sst", "Hallazgos cerrados / hallazgos totales"),
        ("KPI-HAL-002", "Hallazgos abiertos", "HALLAZGOS", float(hallazgos_abiertos), 0, "N°", "inspecciones_hallazgos_sst", "Cantidad de hallazgos abiertos"),
        ("KPI-CAPA-001", "Cumplimiento CAPA", "CAPA", cumplimiento_capa, 90, "%", "capas_sst", "CAPA cerradas / CAPA totales"),
        ("KPI-CAPA-002", "Efectividad CAPA", "CAPA", efectividad_capa, 85, "%", "capas_sst", "CAPA efectivas / CAPA cerradas"),
        ("KPI-INC-001", "Eventos SST registrados", "INCIDENTES", float(incidentes_total), 0, "N°", "incidentes_accidentes_sst", "Incidentes y accidentes reportados"),
        ("KPI-INC-002", "Eventos abiertos", "INCIDENTES", float(abiertos_eventos), 0, "N°", "incidentes_accidentes_sst", "Eventos pendientes de cierre"),
        ("KPI-ACC-001", "Accidentes registrados", "ACCIDENTES", float(accidentes_total), 0, "N°", "incidentes_accidentes_sst", "Accidentes registrados"),
        ("KPI-ACC-002", "Eventos graves/mortales", "ACCIDENTES", float(graves + mortales), 0, "N°", "incidentes_accidentes_sst", "Eventos de alta severidad o mortales"),
        ("KPI-CAP-001", "Cumplimiento capacitaciones", "CAPACITACION", cumplimiento_capacitaciones, 90, "%", "capacitaciones_sst", "Capacitaciones ejecutadas / capacitaciones totales"),
        ("KPI-CAP-002", "Cobertura capacitación", "CAPACITACION", cobertura_capacitacion, 90, "%", "capacitaciones_sst_asistentes", "Asistentes registrados / empleados activos"),
        ("KPI-EPP-001", "Cobertura EPP", "EPP", cobertura_epp, 90, "%", "epp_entregas", "Entregas EPP / empleados activos"),
        ("KPI-EXA-001", "Cobertura exámenes médicos", "EXAMENES", cobertura_examenes, 90, "%", "examenes_medicos", "Exámenes médicos / empleados activos"),
        ("KPI-AUD-001", "Cumplimiento auditorías SST", "AUDITORIAS", cumplimiento_auditorias, 90, "%", "auditorias_sst", "Auditorías cerradas / auditorías totales"),
        ("KPI-AUD-002", "Cierre hallazgos auditoría", "AUDITORIAS", cierre_hallazgos_auditoria, 90, "%", "auditorias_hallazgos_sst", "Hallazgos de auditoría cerrados / total"),
        ("KPI-SST-001", "Tasa de eventos por 100 empleados", "BI SST", tasa_eventos, 0, "Tasa", "incidentes_accidentes_sst + empleados", "Eventos SST por cada 100 empleados activos"),
    ]

    kpis: list[KPIAutomaticoResponse] = []
    for codigo, nombre, categoria, valor, meta, unidad, fuente, descripcion in datos:
        bajo_es_mejor = codigo in {"KPI-HAL-002", "KPI-INC-001", "KPI-INC-002", "KPI-ACC-001", "KPI-ACC-002", "KPI-SST-001"}
        cumplimiento = 100.0 if meta == 0 and valor == 0 else (max(0.0, min(100.0, safe_pct(meta, valor))) if bajo_es_mejor and valor > 0 else min(100.0, safe_pct(valor, meta) if meta else 100.0))
        semaforo = semaforo_por_riesgo(valor, True) if bajo_es_mejor else semaforo_por_cumplimiento(cumplimiento)
        kpis.append(KPIAutomaticoResponse(
            codigo=codigo,
            nombre=nombre,
            categoria=categoria,
            valor=round(float(valor), 2),
            meta=float(meta),
            unidad=unidad,
            cumplimiento=round(float(cumplimiento), 2),
            semaforo=semaforo,
            tendencia="ESTABLE",
            fuente=fuente,
            descripcion=descripcion,
        ))
    return kpis


def crear_excel_indicadores(indicadores: list[IndicadorSSTResponse], kpis: list[KPIAutomaticoResponse]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Indicadores SST"

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(["Código", "Indicador", "Categoría", "Valor", "Meta", "Unidad", "Cumplimiento", "Semáforo", "Fuente"])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for kpi in kpis:
        ws.append([kpi.codigo, kpi.nombre, kpi.categoria, kpi.valor, kpi.meta, kpi.unidad, kpi.cumplimiento, kpi.semaforo, kpi.fuente])

    for ind in indicadores:
        ws.append([ind.codigo, ind.nombre, ind.categoria, float(ind.valor_actual), float(ind.meta), ind.unidad, float(ind.resultado), ind.semaforo, ind.fuente or "Manual"])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    widths = [18, 42, 18, 12, 12, 10, 16, 14, 28]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def crear_pdf_indicadores(indicadores: list[IndicadorSSTResponse], kpis: list[KPIAutomaticoResponse], titulo: str = "Reporte Indicadores SST") -> BytesIO:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=32, leftMargin=32, topMargin=32, bottomMargin=32)
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph(f"<b>{titulo}</b>", styles["Title"]))
    story.append(Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]))
    story.append(Spacer(1, 14))

    data = [["Código", "Indicador", "Categoría", "Valor", "Meta", "Cumpl.", "Semáforo"]]
    for kpi in kpis[:18]:
        data.append([kpi.codigo, kpi.nombre, kpi.categoria, f"{kpi.valor} {kpi.unidad}", f"{kpi.meta} {kpi.unidad}", f"{kpi.cumplimiento}%", kpi.semaforo])
    for ind in indicadores[:30]:
        data.append([ind.codigo, ind.nombre, ind.categoria, f"{float(ind.valor_actual)} {ind.unidad}", f"{float(ind.meta)} {ind.unidad}", f"{float(ind.resultado)}%", ind.semaforo])

    table = Table(data, repeatRows=1, colWidths=[82, 220, 92, 76, 76, 70, 72])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    story.append(Spacer(1, 12))
    story.append(Paragraph("Recomendación: revisar indicadores en rojo y amarillo dentro del comité SST y revisión por la dirección.", styles["Normal"]))
    doc.build(story)
    buffer.seek(0)
    return buffer


# ============================================================
# CRUD
# ============================================================

@router.get("/", response_model=list[IndicadorSSTResponse])
def listar_indicadores(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    categoria: str | None = Query(default=None),
    semaforo: str | None = Query(default=None),
    buscar: str | None = Query(default=None),
    activo: bool | None = Query(default=True),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    query = limpiar_query_indicadores(db, empresa_id, sede_id, area_id, activo)

    if categoria:
        query = query.filter(func.upper(IndicadorSST.categoria) == categoria.upper())
    if semaforo:
        query = query.filter(func.upper(IndicadorSST.semaforo) == semaforo.upper())
    if buscar:
        q = f"%{buscar.lower()}%"
        query = query.filter(or_(
            func.lower(IndicadorSST.codigo).like(q),
            func.lower(IndicadorSST.nombre).like(q),
            func.lower(IndicadorSST.descripcion).like(q),
            func.lower(IndicadorSST.responsable).like(q),
            func.lower(IndicadorSST.fuente).like(q),
        ))

    indicadores = query.order_by(IndicadorSST.id.desc()).all()
    return [indicador_to_response(i) for i in indicadores]


@router.post("/", response_model=IndicadorSSTResponse)
def crear_indicador(
    data: IndicadorSSTCreate,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    payload = data.model_dump()
    validar_relaciones(db, payload)

    existe = db.query(IndicadorSST).filter(
        IndicadorSST.empresa_id == data.empresa_id,
        func.lower(IndicadorSST.codigo) == data.codigo.lower(),
    ).first()
    if existe:
        raise HTTPException(status_code=400, detail="Ya existe un indicador con este código para la empresa")

    valor = to_float(data.valor_actual)
    meta = to_float(data.meta)
    resultado = to_float(data.resultado) or (safe_pct(valor, meta) if meta else 0)
    semaforo = normalizar_texto(data.semaforo, upper=True) or semaforo_por_cumplimiento(resultado)

    payload["codigo"] = normalizar_texto(data.codigo, upper=True)
    payload["categoria"] = normalizar_texto(data.categoria, upper=True, defecto="GESTION")
    payload["tipo_indicador"] = normalizar_texto(data.tipo_indicador, upper=True, defecto="RESULTADO")
    payload["origen_dato"] = normalizar_texto(data.origen_dato, upper=True, defecto="MANUAL")
    payload["frecuencia"] = normalizar_texto(data.frecuencia, upper=True, defecto="MENSUAL")
    payload["semaforo"] = semaforo
    payload["resultado"] = Decimal(str(round(resultado, 2)))

    indicador = IndicadorSST(
        **payload,
        usuario_id=getattr(usuario, "id", None) or getattr(usuario, "user_id", None),
    )
    db.add(indicador)
    db.commit()
    db.refresh(indicador)
    return indicador_to_response(indicador)


@router.get("/dashboard/resumen", response_model=IndicadoresDashboardResponse)
def dashboard_indicadores(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    manuales_db = limpiar_query_indicadores(db, empresa_id, sede_id, area_id, True).all()
    manuales = [indicador_to_response(i) for i in manuales_db]
    kpis = calcular_kpis_automaticos(db, empresa_id, sede_id, area_id)

    total = len(kpis) + len(manuales)
    verdes = sum(1 for k in kpis if k.semaforo == "VERDE") + sum(1 for i in manuales if (i.semaforo or "").upper() == "VERDE")
    amarillos = sum(1 for k in kpis if k.semaforo == "AMARILLO") + sum(1 for i in manuales if (i.semaforo or "").upper() == "AMARILLO")
    rojos = sum(1 for k in kpis if k.semaforo == "ROJO") + sum(1 for i in manuales if (i.semaforo or "").upper() == "ROJO")

    cumplimientos = [k.cumplimiento for k in kpis]
    cumplimientos.extend([to_float(i.resultado) for i in manuales])
    cumplimiento_global = round(sum(cumplimientos) / len(cumplimientos), 2) if cumplimientos else 0.0
    semaforo_global = semaforo_por_cumplimiento(cumplimiento_global)

    distribucion_categoria: dict[str, int] = {}
    distribucion_semaforo = {"VERDE": verdes, "AMARILLO": amarillos, "ROJO": rojos}
    for k in kpis:
        distribucion_categoria[k.categoria] = distribucion_categoria.get(k.categoria, 0) + 1
    for i in manuales:
        categoria = i.categoria or "MANUAL"
        distribucion_categoria[categoria] = distribucion_categoria.get(categoria, 0) + 1

    recomendaciones = []
    if rojos:
        recomendaciones.append("Priorizar cierre de indicadores en rojo y definir responsables por acción.")
    if amarillos:
        recomendaciones.append("Revisar indicadores en amarillo durante el comité SST del periodo.")
    if cumplimiento_global < 80:
        recomendaciones.append("Generar plan de mejora transversal del SG-SST por bajo cumplimiento global.")
    if not recomendaciones:
        recomendaciones.append("Gestión SST estable. Mantener seguimiento preventivo y revisión mensual.")

    return IndicadoresDashboardResponse(
        total_indicadores=total,
        automaticos=len(kpis),
        manuales=len(manuales),
        verdes=verdes,
        amarillos=amarillos,
        rojos=rojos,
        cumplimiento_global=cumplimiento_global,
        score_sst=cumplimiento_global,
        semaforo_global=semaforo_global,
        kpis_automaticos=kpis,
        distribucion_categoria=distribucion_categoria,
        distribucion_semaforo=distribucion_semaforo,
        alertas={"rojos": rojos, "amarillos": amarillos, "verdes": verdes},
        recomendaciones=recomendaciones,
    )


@router.get("/{indicador_id}", response_model=IndicadorSSTResponse)
def obtener_indicador(
    indicador_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    indicador = db.query(IndicadorSST).filter(IndicadorSST.id == indicador_id).first()
    if not indicador:
        raise HTTPException(status_code=404, detail="Indicador no encontrado")
    return indicador_to_response(indicador)


@router.put("/{indicador_id}", response_model=IndicadorSSTResponse)
def actualizar_indicador(
    indicador_id: int,
    data: IndicadorSSTUpdate,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    indicador = db.query(IndicadorSST).filter(IndicadorSST.id == indicador_id).first()
    if not indicador:
        raise HTTPException(status_code=404, detail="Indicador no encontrado")

    update = data.model_dump(exclude_unset=True)
    validar_relaciones(db, {**{ "empresa_id": indicador.empresa_id }, **update})

    if "codigo" in update and update["codigo"]:
        existe = db.query(IndicadorSST).filter(
            IndicadorSST.id != indicador_id,
            IndicadorSST.empresa_id == update.get("empresa_id", indicador.empresa_id),
            func.lower(IndicadorSST.codigo) == update["codigo"].lower(),
        ).first()
        if existe:
            raise HTTPException(status_code=400, detail="Ya existe otro indicador con este código")

    for key, value in update.items():
        if key in {"codigo", "categoria", "tipo_indicador", "origen_dato", "frecuencia", "semaforo", "tendencia"} and value:
            value = normalizar_texto(value, upper=True)
        elif isinstance(value, str):
            value = normalizar_texto(value)
        setattr(indicador, key, value)

    if "valor_actual" in update or "meta" in update or "resultado" not in update:
        valor = to_float(indicador.valor_actual)
        meta = to_float(indicador.meta)
        indicador.resultado = Decimal(str(round(safe_pct(valor, meta), 2))) if meta else Decimal("0")
        indicador.semaforo = semaforo_por_cumplimiento(float(indicador.resultado))

    db.commit()
    db.refresh(indicador)
    return indicador_to_response(indicador)


@router.delete("/{indicador_id}")
def eliminar_indicador(
    indicador_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_ADMIN)),
):
    indicador = db.query(IndicadorSST).filter(IndicadorSST.id == indicador_id).first()
    if not indicador:
        raise HTTPException(status_code=404, detail="Indicador no encontrado")
    indicador.activo = False
    db.commit()
    return {"mensaje": "Indicador desactivado correctamente", "indicador_id": indicador_id}


# ============================================================
# EXPORTACIONES
# ============================================================

@router.get("/export/excel")
def exportar_excel_indicadores(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    manuales = [indicador_to_response(i) for i in limpiar_query_indicadores(db, empresa_id, sede_id, area_id, True).all()]
    kpis = calcular_kpis_automaticos(db, empresa_id, sede_id, area_id)
    output = crear_excel_indicadores(manuales, kpis)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=indicadores_sst_bi_executive.xlsx"},
    )


@router.get("/export/pdf")
def exportar_pdf_indicadores(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    manuales = [indicador_to_response(i) for i in limpiar_query_indicadores(db, empresa_id, sede_id, area_id, True).all()]
    kpis = calcular_kpis_automaticos(db, empresa_id, sede_id, area_id)
    output = crear_pdf_indicadores(manuales, kpis, "Reporte Ejecutivo Indicadores SST")
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=indicadores_sst_bi_executive.pdf"},
    )


@router.get("/export/dashboard-pdf")
def exportar_dashboard_pdf_indicadores(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    manuales = [indicador_to_response(i) for i in limpiar_query_indicadores(db, empresa_id, sede_id, area_id, True).all()]
    kpis = calcular_kpis_automaticos(db, empresa_id, sede_id, area_id)
    output = crear_pdf_indicadores(manuales, kpis, "Dashboard Ejecutivo Indicadores SST")
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=dashboard_indicadores_sst.pdf"},
    )
