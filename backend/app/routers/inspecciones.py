# ============================================================
# ROUTER INSPECCIONES SST ENTERPRISE - ERP SST PRO
# FASE 1.1.8 — INSPECCIONES SST ENTERPRISE
# Archivo: backend/app/routers/inspecciones.py
# ============================================================

from datetime import date, datetime, timedelta
from pathlib import Path
import io
import os
import uuid

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, joinedload
from starlette.responses import StreamingResponse

from app.auth.dependencies import require_roles, require_permission
from app.core.default_permissions import PERM_REGISTROS_ELIMINAR, PERM_REPORTES_EXPORTAR
from app.core.file_security import validate_upload
from app.database import get_db
from app.models.archivo_sst import ArchivoSST
from app.models.area import Area
from app.models.cargo import Cargo
from app.models.empleado import Empleado
from app.models.empresa import Empresa
from app.models.inspeccion import InspeccionHallazgoSST, InspeccionSST
from app.models.inspeccion_seguimiento import InspeccionHallazgoSeguimientoSST
from app.models.sede import Sede
from app.schemas.inspeccion_schema import (
    HallazgoCreate,
    HallazgoResponse,
    HallazgoUpdate,
    InspeccionCreate,
    InspeccionDashboardResponse,
    InspeccionResponse,
    InspeccionUpdate,
    InspeccionFirmaRequest,
    InspeccionCierreDigitalRequest,
)

router = APIRouter(prefix="/inspecciones", tags=["Inspecciones SST Enterprise"])
ROLES_SST = ["SUPER_ADMIN", "ADMIN_EMPRESA", "RESPONSABLE_SST"]
EXPORTAR_REPORTES = require_permission(PERM_REPORTES_EXPORTAR)
ELIMINAR_REGISTROS = require_permission(PERM_REGISTROS_ELIMINAR)

UPLOAD_ROOT = Path(os.getenv("UPLOAD_DIR", "app/uploads")).resolve()
INSPECCIONES_UPLOAD_DIR = UPLOAD_ROOT / "inspecciones"
INSPECCIONES_PREVIEW_DIR = INSPECCIONES_UPLOAD_DIR / "previews"
INSPECCIONES_THUMB_DIR = INSPECCIONES_UPLOAD_DIR / "thumbs"

INSPECCIONES_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
INSPECCIONES_PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
INSPECCIONES_THUMB_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXT = {"pdf", "jpg", "jpeg", "png", "webp"}
ALLOWED_MIME = {"application/pdf", "image/jpeg", "image/png", "image/webp"}
MAX_UPLOAD_MB = 20


def _public_upload_url(file_path: Path) -> str:
    try:
        rel = file_path.resolve().relative_to(UPLOAD_ROOT)
        return "/uploads/" + rel.as_posix()
    except Exception:
        return "/uploads/inspecciones/" + file_path.name


def _webp_bytes_from_image(image, max_size: tuple[int, int], quality: int) -> bytes:
    """
    FASE 1.1.8.5 — Optimización Enterprise de evidencias.
    Genera versiones WEBP livianas para almacenamiento, preview y miniatura.
    """
    from PIL import Image

    img = image.copy()
    img.thumbnail(max_size)

    # WEBP no maneja transparencia de forma uniforme en todos los visores.
    # Convertimos a RGB con fondo blanco para evidencias SST imprimibles y PDF.
    if img.mode in ("RGBA", "LA", "P"):
        fondo = Image.new("RGB", img.size, "white")
        if img.mode == "P":
            img = img.convert("RGBA")
        fondo.paste(img, mask=img.split()[-1] if img.mode in ("RGBA", "LA") else None)
        img = fondo
    elif img.mode != "RGB":
        img = img.convert("RGB")

    out = io.BytesIO()
    img.save(out, format="WEBP", quality=quality, method=6, optimize=True)
    return out.getvalue()


def _optimizar_imagen_enterprise(content: bytes) -> dict:
    """
    Convierte JPG/PNG/WEBP a:
    - original optimizado: máximo 1920x1920, calidad 82
    - preview: máximo 1280x1280, calidad 76
    - thumb: máximo 360x360, calidad 68

    Esto reduce almacenamiento y acelera el frontend sin perder calidad operativa.
    """
    try:
        from PIL import Image, ImageOps

        image = Image.open(io.BytesIO(content))
        image = ImageOps.exif_transpose(image)

        original = _webp_bytes_from_image(image, (1920, 1920), 82)
        preview = _webp_bytes_from_image(image, (1280, 1280), 76)
        thumb = _webp_bytes_from_image(image, (360, 360), 68)

        return {
            "ok": True,
            "extension": "webp",
            "mime_type": "image/webp",
            "original": original,
            "preview": preview,
            "thumb": thumb,
        }
    except Exception:
        return {
            "ok": False,
            "extension": "bin",
            "mime_type": "application/octet-stream",
            "original": content,
            "preview": content,
            "thumb": content,
        }


def _optimizar_pdf_bytes(content: bytes) -> bytes:
    try:
        import pikepdf
        src = io.BytesIO(content)
        out = io.BytesIO()
        with pikepdf.Pdf.open(src) as pdf:
            pdf.save(out, compress_streams=True, object_stream_mode=pikepdf.ObjectStreamMode.generate, linearize=True)
        optimized = out.getvalue()
        return optimized if len(optimized) < len(content) else content
    except Exception:
        return content


def _guardar_upload(upload: UploadFile) -> tuple[Path, str, str, str, int]:
    validation = validate_upload(upload, allowed_extensions={f".{item}" for item in ALLOWED_EXT}, max_size_mb=MAX_UPLOAD_MB)
    original = validation.safe_filename or "evidencia_inspeccion"
    extension = validation.extension.lstrip(".")
    content = validation.content

    if extension in {"jpg", "jpeg", "png", "webp"}:
        optimized = _optimizar_imagen_enterprise(content)
        if not optimized["ok"]:
            raise HTTPException(status_code=400, detail="No fue posible procesar la imagen")

        filename = f"{uuid.uuid4().hex}.webp"
        path = INSPECCIONES_UPLOAD_DIR / filename
        preview_path = INSPECCIONES_PREVIEW_DIR / filename
        thumb_path = INSPECCIONES_THUMB_DIR / filename

        path.write_bytes(optimized["original"])
        preview_path.write_bytes(optimized["preview"])
        thumb_path.write_bytes(optimized["thumb"])

        return path, original, filename, optimized["mime_type"], len(optimized["original"])

    if extension == "pdf":
        content = _optimizar_pdf_bytes(content)
        mime_type = "application/pdf"
    else:
        mime_type = validation.mime_type

    filename = f"{uuid.uuid4().hex}.{extension}"
    path = INSPECCIONES_UPLOAD_DIR / filename
    path.write_bytes(content)
    return path, original, filename, mime_type, len(content)

def _archivo_to_dict(archivo: ArchivoSST):
    preview_url = _archivo_variant_url(archivo, "preview")
    thumbnail_url = _archivo_variant_url(archivo, "thumb")

    return {
        "id": archivo.id,
        "empresa_id": archivo.empresa_id,
        "usuario_id": archivo.usuario_id,
        "tipo": archivo.tipo,
        "nombre_original": archivo.nombre_original,
        "nombre_archivo": archivo.nombre_archivo,
        "ruta": archivo.ruta,
        "url": archivo.url,
        "preview_url": preview_url or archivo.url,
        "thumbnail_url": thumbnail_url or preview_url or archivo.url,
        "extension": archivo.extension,
        "mime_type": archivo.mime_type,
        "tamano_bytes": archivo.tamano_bytes,
        "modulo": archivo.modulo,
        "referencia_id": archivo.referencia_id,
        "descripcion": archivo.descripcion,
        "activo": archivo.activo,
        "fecha_creacion": archivo.fecha_creacion,
        "optimizacion": {
            "webp": (archivo.extension or "").lower() == "webp",
            "preview": bool(preview_url),
            "thumbnail": bool(thumbnail_url),
            "max_original_px": 1920,
            "max_preview_px": 1280,
            "max_thumbnail_px": 360,
        },
    }


def _upper(value, default=None):
    if value is None:
        return default
    value = str(value).strip().upper()
    return value if value else default


def _validar_empresa(db: Session, empresa_id: int):
    empresa = db.query(Empresa).filter(Empresa.id == empresa_id).first()
    if not empresa:
        raise HTTPException(status_code=404, detail="Empresa no encontrada")
    return empresa


def _validar_opcional(db: Session, model, item_id: int | None, label: str):
    if not item_id:
        return None
    item = db.query(model).filter(model.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail=f"{label} no encontrado")
    return item


def _query_inspecciones(db: Session, empresa_id=None, sede_id=None, area_id=None, cargo_id=None, empleado_id=None, estado=None, riesgo=None, q=None):
    query = db.query(InspeccionSST).options(
        joinedload(InspeccionSST.empresa),
        joinedload(InspeccionSST.sede),
        joinedload(InspeccionSST.area),
        joinedload(InspeccionSST.cargo),
        joinedload(InspeccionSST.empleado),
    )
    if empresa_id:
        query = query.filter(InspeccionSST.empresa_id == empresa_id)
    if sede_id:
        query = query.filter(InspeccionSST.sede_id == sede_id)
    if area_id:
        query = query.filter(InspeccionSST.area_id == area_id)
    if cargo_id:
        query = query.filter(InspeccionSST.cargo_id == cargo_id)
    if empleado_id:
        query = query.filter(InspeccionSST.empleado_id == empleado_id)
    if estado:
        query = query.filter(func.upper(InspeccionSST.estado) == estado.upper().strip())
    if riesgo:
        query = query.filter(func.upper(InspeccionSST.nivel_riesgo) == riesgo.upper().strip())
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(or_(InspeccionSST.codigo.ilike(like), InspeccionSST.titulo.ilike(like), InspeccionSST.responsable.ilike(like), InspeccionSST.lugar.ilike(like)))
    return query.order_by(InspeccionSST.id.desc())


def _inspeccion_to_response(db: Session, item: InspeccionSST):
    data = InspeccionResponse.model_validate(item)
    data.empresa_nombre = item.empresa.nombre if item.empresa else None
    data.sede_nombre = item.sede.nombre if item.sede else None
    data.area_nombre = item.area.nombre if item.area else None
    data.cargo_nombre = item.cargo.nombre if item.cargo else None
    if item.empleado:
        data.empleado_nombre = f"{item.empleado.nombres} {item.empleado.apellidos}".strip()
        data.empleado_documento = item.empleado.documento
    data.total_hallazgos = db.query(func.count(InspeccionHallazgoSST.id)).filter(InspeccionHallazgoSST.inspeccion_id == item.id, InspeccionHallazgoSST.activo.is_(True)).scalar() or 0
    data.hallazgos_abiertos = db.query(func.count(InspeccionHallazgoSST.id)).filter(InspeccionHallazgoSST.inspeccion_id == item.id, InspeccionHallazgoSST.activo.is_(True), InspeccionHallazgoSST.estado != "CERRADO").scalar() or 0
    data.total_evidencias = db.query(func.count(ArchivoSST.id)).filter(ArchivoSST.modulo == "INSPECCIONES", ArchivoSST.referencia_id == item.id, ArchivoSST.activo.is_(True)).scalar() or 0
    return data


@router.get("/", response_model=list[InspeccionResponse])
def listar_inspecciones(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    cargo_id: int | None = Query(default=None),
    empleado_id: int | None = Query(default=None),
    estado: str | None = Query(default=None),
    riesgo: str | None = Query(default=None),
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    items = _query_inspecciones(db, empresa_id, sede_id, area_id, cargo_id, empleado_id, estado, riesgo, q).all()
    return [_inspeccion_to_response(db, item) for item in items]


@router.get("/dashboard/resumen", response_model=InspeccionDashboardResponse)
def dashboard_inspecciones(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    cargo_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    inspecciones = _query_inspecciones(db, empresa_id, sede_id, area_id, cargo_id).filter(InspeccionSST.activo.is_(True)).all()
    ids = [i.id for i in inspecciones]
    total = len(inspecciones)
    ejecutadas = sum(1 for i in inspecciones if i.estado in ["EJECUTADA", "CERRADA"])
    programadas = sum(1 for i in inspecciones if i.estado == "PROGRAMADA")
    cerradas = sum(1 for i in inspecciones if i.estado == "CERRADA")
    alto_critico = sum(1 for i in inspecciones if i.nivel_riesgo in ["ALTO", "CRITICO"])
    cumplimiento = round(sum(float(i.cumplimiento or 0) for i in inspecciones) / total, 1) if total else 0
    hoy = date.today()
    vencidas = sum(1 for i in inspecciones if i.fecha_programada and i.fecha_programada < hoy and i.estado not in ["EJECUTADA", "CERRADA", "ANULADA"])

    hallazgos = []
    evidencias = 0
    if ids:
        hallazgos = db.query(InspeccionHallazgoSST).filter(InspeccionHallazgoSST.inspeccion_id.in_(ids), InspeccionHallazgoSST.activo.is_(True)).all()
        evidencias = db.query(func.count(ArchivoSST.id)).filter(ArchivoSST.modulo == "INSPECCIONES", ArchivoSST.referencia_id.in_(ids), ArchivoSST.activo.is_(True)).scalar() or 0
    hallazgos_abiertos = sum(1 for h in hallazgos if h.estado != "CERRADO")
    hallazgos_criticos = sum(1 for h in hallazgos if h.nivel_riesgo in ["ALTO", "CRITICO"])

    def conteo(attr):
        data = {}
        for item in inspecciones:
            key = attr(item) or "Sin dato"
            data[key] = data.get(key, 0) + 1
        return [{"name": k, "value": v} for k, v in sorted(data.items(), key=lambda x: x[1], reverse=True)[:8]]

    riesgo_score = 0
    riesgo_score += 35 if hallazgos_criticos else 0
    riesgo_score += 25 if alto_critico else 0
    riesgo_score += 20 if vencidas else 0
    riesgo_score += 10 if hallazgos_abiertos else 0
    riesgo_score += 10 if total and not evidencias else 0
    semaforo = "ROJO" if riesgo_score >= 60 else "AMARILLO" if riesgo_score >= 25 else "VERDE"

    recomendaciones = []
    if vencidas:
        recomendaciones.append("Ejecutar inspecciones programadas vencidas.")
    if hallazgos_criticos:
        recomendaciones.append("Priorizar cierre de hallazgos alto/crítico.")
    if hallazgos_abiertos:
        recomendaciones.append("Asignar responsables y fechas compromiso a hallazgos abiertos.")
    if total and not evidencias:
        recomendaciones.append("Adjuntar evidencias fotográficas o soportes PDF a las inspecciones.")
    if not recomendaciones:
        recomendaciones.append("Gestión de inspecciones estable. Mantén seguimiento periódico.")

    return {
        "kpis": {
            "total": total,
            "programadas": programadas,
            "ejecutadas": ejecutadas,
            "cerradas": cerradas,
            "alto_critico": alto_critico,
            "vencidas": vencidas,
            "hallazgos": len(hallazgos),
            "hallazgos_abiertos": hallazgos_abiertos,
            "hallazgos_criticos": hallazgos_criticos,
            "evidencias": evidencias,
            "cumplimiento": cumplimiento,
            "riesgo_score": riesgo_score,
            "semaforo": semaforo,
        },
        "charts": {
            "por_tipo": conteo(lambda x: x.tipo_inspeccion),
            "por_estado": conteo(lambda x: x.estado),
            "por_riesgo": conteo(lambda x: x.nivel_riesgo),
            "por_area": conteo(lambda x: x.area.nombre if x.area else "Sin área"),
            "por_cargo": conteo(lambda x: x.cargo.nombre if x.cargo else "Sin cargo"),
            "por_resultado": conteo(lambda x: x.resultado),
        },
        "alertas": {
            "vencidas": vencidas,
            "alto_critico": alto_critico,
            "hallazgos_abiertos": hallazgos_abiertos,
            "hallazgos_criticos": hallazgos_criticos,
            "sin_evidencia": max(total - evidencias, 0) if total else 0,
        },
        "recomendaciones": recomendaciones,
    }




# ============================================================
# FASE 1.1.8.3 — EXPORTACIONES PDF / EXCEL INSPECCIONES SST
# Enterprise: Excel General, PDF General, Acta PDF, PDF Individual,
# Hallazgos Excel/PDF, Seguimientos PDF y Dashboard Ejecutivo PDF.
# ============================================================

def _fmt(value):
    if value is None:
        return ""
    try:
        if hasattr(value, "strftime"):
            return value.strftime("%Y-%m-%d %H:%M") if getattr(value, "hour", None) is not None else value.strftime("%Y-%m-%d")
    except Exception:
        pass
    return str(value)


def _safe_filename(name: str) -> str:
    return "".join(c if c.isalnum() or c in "._-" else "_" for c in name)


def _excel_response(wb, filename: str):
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_safe_filename(filename)}"'},
    )


def _pdf_response(buffer: io.BytesIO, filename: str):
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{_safe_filename(filename)}"'},
    )


def _xlsx_header(ws, headers):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, _ in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(idx)].width = 22
    ws.freeze_panes = "A2"


def _pdf_doc(title: str, subtitle: str = ""):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Spacer

    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="TitleSST", parent=styles["Title"], alignment=TA_CENTER, fontSize=16, leading=20, textColor=colors.HexColor("#0f172a")))
    styles.add(ParagraphStyle(name="SubSST", parent=styles["Normal"], alignment=TA_CENTER, fontSize=9, leading=12, textColor=colors.HexColor("#475569")))
    styles.add(ParagraphStyle(name="H2SST", parent=styles["Heading2"], alignment=TA_LEFT, fontSize=11, leading=14, textColor=colors.HexColor("#0f766e")))
    elements = [Paragraph(title, styles["TitleSST"])]
    if subtitle:
        elements.append(Paragraph(subtitle, styles["SubSST"]))
    elements.append(Spacer(1, 0.25 * cm))
    return buffer, elements, styles


def _pdf_table(data, col_widths=None, repeat=1):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle
    tbl = Table(data, colWidths=col_widths, repeatRows=repeat)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return tbl


def _firma_flowable(firma_base64: str | None, width=120, height=45):
    if not firma_base64:
        return "Sin firma"
    try:
        import base64
        from reportlab.lib.utils import ImageReader
        from reportlab.platypus import Image
        raw = firma_base64.split(",", 1)[1] if "," in firma_base64 else firma_base64
        data = base64.b64decode(raw)
        img = Image(ImageReader(io.BytesIO(data)), width=width, height=height)
        return img
    except Exception:
        return "Firma registrada"


def _inspecciones_filtradas_export(db: Session, empresa_id=None, sede_id=None, area_id=None, cargo_id=None, empleado_id=None, estado=None, riesgo=None, q=None):
    return _query_inspecciones(db, empresa_id, sede_id, area_id, cargo_id, empleado_id, estado, riesgo, q).filter(InspeccionSST.activo.is_(True)).all()


@router.get("/exportaciones/excel-general")
def exportar_excel_general_inspecciones(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    cargo_id: int | None = Query(default=None),
    empleado_id: int | None = Query(default=None),
    estado: str | None = Query(default=None),
    riesgo: str | None = Query(default=None),
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Inspecciones SST"
    headers = ["ID", "Código", "Empresa", "Sede", "Área", "Cargo", "Empleado", "Tipo", "Título", "Lugar", "Responsable", "Fecha programada", "Fecha inspección", "Estado", "Resultado", "Riesgo", "Cumplimiento %", "Hallazgos", "Hallazgos abiertos", "Evidencias", "Cierre digital"]
    _xlsx_header(ws, headers)
    for item in _inspecciones_filtradas_export(db, empresa_id, sede_id, area_id, cargo_id, empleado_id, estado, riesgo, q):
        resp = _inspeccion_to_response(db, item)
        ws.append([item.id, item.codigo, resp.empresa_nombre, resp.sede_nombre, resp.area_nombre, resp.cargo_nombre, resp.empleado_nombre, item.tipo_inspeccion, item.titulo, item.lugar, item.responsable, _fmt(item.fecha_programada), _fmt(item.fecha_inspeccion), item.estado, item.resultado, item.nivel_riesgo, float(item.cumplimiento or 0), resp.total_hallazgos, resp.hallazgos_abiertos, resp.total_evidencias, "SI" if item.cierre_digital else "NO"])
    return _excel_response(wb, "inspecciones_sst_general.xlsx")


@router.get("/exportaciones/pdf-general")
def exportar_pdf_general_inspecciones(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    cargo_id: int | None = Query(default=None),
    empleado_id: int | None = Query(default=None),
    estado: str | None = Query(default=None),
    riesgo: str | None = Query(default=None),
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Spacer
    buffer, elements, styles = _pdf_doc("Reporte General de Inspecciones SST", f"Generado: {_fmt(datetime.utcnow())}")
    rows = [["Código", "Empresa", "Sede/Área", "Fecha", "Estado", "Riesgo", "Cumpl.", "Hallazgos"]]
    items = _inspecciones_filtradas_export(db, empresa_id, sede_id, area_id, cargo_id, empleado_id, estado, riesgo, q)
    for item in items:
        resp = _inspeccion_to_response(db, item)
        rows.append([item.codigo, resp.empresa_nombre or "", f"{resp.sede_nombre or ''} / {resp.area_nombre or ''}", _fmt(item.fecha_inspeccion), item.estado or "", item.nivel_riesgo or "", f"{float(item.cumplimiento or 0):.1f}%", str(resp.total_hallazgos)])
    elements.append(_pdf_table(rows, [2.1*cm, 4.0*cm, 4.6*cm, 2.1*cm, 2.2*cm, 2.0*cm, 1.5*cm, 1.8*cm]))
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm)
    doc.build(elements)
    return _pdf_response(buffer, "inspecciones_sst_general.pdf")


def _hallazgos_query_export(db: Session, empresa_id=None, inspeccion_id=None, estado=None, riesgo=None):
    qh = db.query(InspeccionHallazgoSST).join(InspeccionSST, InspeccionSST.id == InspeccionHallazgoSST.inspeccion_id).filter(InspeccionHallazgoSST.activo.is_(True))
    if empresa_id:
        qh = qh.filter(InspeccionHallazgoSST.empresa_id == empresa_id)
    if inspeccion_id:
        qh = qh.filter(InspeccionHallazgoSST.inspeccion_id == inspeccion_id)
    if estado:
        qh = qh.filter(func.upper(InspeccionHallazgoSST.estado) == estado.upper().strip())
    if riesgo:
        qh = qh.filter(func.upper(InspeccionHallazgoSST.nivel_riesgo) == riesgo.upper().strip())
    return qh.order_by(InspeccionHallazgoSST.id.desc()).all()


@router.get("/exportaciones/hallazgos-excel")
def exportar_hallazgos_excel(
    empresa_id: int | None = Query(default=None),
    inspeccion_id: int | None = Query(default=None),
    estado: str | None = Query(default=None),
    riesgo: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Hallazgos SST"
    _xlsx_header(ws, ["ID", "Inspección", "Código", "Empresa", "Descripción", "Tipo", "Riesgo", "Acción recomendada", "Responsable", "Compromiso", "Cierre", "Estado", "Observaciones"])
    for h in _hallazgos_query_export(db, empresa_id, inspeccion_id, estado, riesgo):
        insp = h.inspeccion
        ws.append([h.id, h.inspeccion_id, insp.codigo if insp else "", h.empresa.nombre if h.empresa else "", h.descripcion, h.tipo_hallazgo, h.nivel_riesgo, h.accion_recomendada, h.responsable, _fmt(h.fecha_compromiso), _fmt(h.fecha_cierre), h.estado, h.observaciones])
    return _excel_response(wb, "hallazgos_inspecciones_sst.xlsx")


@router.get("/exportaciones/hallazgos-pdf")
def exportar_hallazgos_pdf(
    empresa_id: int | None = Query(default=None),
    inspeccion_id: int | None = Query(default=None),
    estado: str | None = Query(default=None),
    riesgo: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate
    buffer, elements, styles = _pdf_doc("Reporte de Hallazgos de Inspecciones SST", f"Generado: {_fmt(datetime.utcnow())}")
    rows = [["Inspección", "Descripción", "Tipo", "Riesgo", "Responsable", "Compromiso", "Estado"]]
    for h in _hallazgos_query_export(db, empresa_id, inspeccion_id, estado, riesgo):
        rows.append([h.inspeccion.codigo if h.inspeccion else str(h.inspeccion_id), h.descripcion or "", h.tipo_hallazgo or "", h.nivel_riesgo or "", h.responsable or "", _fmt(h.fecha_compromiso), h.estado or ""])
    elements.append(_pdf_table(rows, [2.2*cm, 8.0*cm, 3.0*cm, 2.0*cm, 3.5*cm, 2.4*cm, 2.4*cm]))
    SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm).build(elements)
    return _pdf_response(buffer, "hallazgos_inspecciones_sst.pdf")


@router.get("/exportaciones/seguimientos-pdf")
def exportar_seguimientos_pdf(
    empresa_id: int | None = Query(default=None),
    inspeccion_id: int | None = Query(default=None),
    hallazgo_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate
    buffer, elements, styles = _pdf_doc("Reporte de Seguimientos de Hallazgos SST", f"Generado: {_fmt(datetime.utcnow())}")
    qs = db.query(InspeccionHallazgoSeguimientoSST).join(InspeccionHallazgoSST, InspeccionHallazgoSST.id == InspeccionHallazgoSeguimientoSST.hallazgo_id).join(InspeccionSST, InspeccionSST.id == InspeccionHallazgoSST.inspeccion_id).filter(InspeccionHallazgoSeguimientoSST.activo.is_(True))
    if empresa_id:
        qs = qs.filter(InspeccionSST.empresa_id == empresa_id)
    if inspeccion_id:
        qs = qs.filter(InspeccionHallazgoSST.inspeccion_id == inspeccion_id)
    if hallazgo_id:
        qs = qs.filter(InspeccionHallazgoSeguimientoSST.hallazgo_id == hallazgo_id)
    rows = [["Fecha", "Inspección", "Hallazgo", "Avance", "Comentario"]]
    for s in qs.order_by(InspeccionHallazgoSeguimientoSST.fecha_registro.desc()).all():
        h = s.hallazgo
        insp = h.inspeccion if h else None
        rows.append([_fmt(s.fecha_registro), insp.codigo if insp else "", h.descripcion if h else "", f"{s.porcentaje_avance or 0}%", s.comentario or ""])
    elements.append(_pdf_table(rows, [3*cm, 2.5*cm, 7*cm, 2*cm, 10*cm]))
    SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1*cm, leftMargin=1*cm, topMargin=1*cm, bottomMargin=1*cm).build(elements)
    return _pdf_response(buffer, "seguimientos_inspecciones_sst.pdf")


@router.get("/exportaciones/dashboard-ejecutivo-pdf")
def exportar_dashboard_ejecutivo_pdf(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    cargo_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Spacer, Paragraph
    resumen = dashboard_inspecciones(empresa_id, sede_id, area_id, cargo_id, db, usuario)
    k = resumen["kpis"]
    buffer, elements, styles = _pdf_doc("Dashboard Ejecutivo de Inspecciones SST", f"Generado: {_fmt(datetime.utcnow())}")
    rows = [["Indicador", "Valor"], ["Total inspecciones", k.get("total", 0)], ["Programadas", k.get("programadas", 0)], ["Ejecutadas", k.get("ejecutadas", 0)], ["Cerradas", k.get("cerradas", 0)], ["Alto / Crítico", k.get("alto_critico", 0)], ["Hallazgos abiertos", k.get("hallazgos_abiertos", 0)], ["Hallazgos críticos", k.get("hallazgos_criticos", 0)], ["Evidencias", k.get("evidencias", 0)], ["Cumplimiento", f"{k.get('cumplimiento', 0)}%"], ["Semáforo", k.get("semaforo", "VERDE")]]
    elements.append(_pdf_table(rows, [8*cm, 8*cm]))
    elements.append(Spacer(1, 0.3*cm))
    elements.append(Paragraph("Recomendaciones PRO", styles["H2SST"]))
    for r in resumen.get("recomendaciones", []):
        elements.append(Paragraph(f"• {r}", styles["Normal"]))
    SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1.5*cm, leftMargin=1.5*cm, topMargin=1.2*cm, bottomMargin=1.2*cm).build(elements)
    return _pdf_response(buffer, "dashboard_ejecutivo_inspecciones_sst.pdf")


@router.get("/exportaciones/{inspeccion_id}/pdf-individual")
def exportar_pdf_individual_inspeccion(inspeccion_id: int, db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Spacer, Paragraph
    item = _query_inspecciones(db).filter(InspeccionSST.id == inspeccion_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Inspección no encontrada")
    resp = _inspeccion_to_response(db, item)
    buffer, elements, styles = _pdf_doc(f"Inspección SST {item.codigo}", item.titulo or "")
    rows = [["Campo", "Valor"], ["Empresa", resp.empresa_nombre or ""], ["Sede", resp.sede_nombre or ""], ["Área", resp.area_nombre or ""], ["Cargo", resp.cargo_nombre or ""], ["Empleado", resp.empleado_nombre or ""], ["Tipo", item.tipo_inspeccion or ""], ["Lugar", item.lugar or ""], ["Responsable", item.responsable or ""], ["Fecha programada", _fmt(item.fecha_programada)], ["Fecha inspección", _fmt(item.fecha_inspeccion)], ["Estado", item.estado or ""], ["Resultado", item.resultado or ""], ["Nivel riesgo", item.nivel_riesgo or ""], ["Cumplimiento", f"{float(item.cumplimiento or 0):.1f}%"], ["Observaciones", item.observaciones or ""]]
    elements.append(_pdf_table(rows, [5*cm, 11*cm], repeat=1))
    elements.append(Spacer(1, 0.3*cm))
    hall = db.query(InspeccionHallazgoSST).filter(InspeccionHallazgoSST.inspeccion_id == item.id, InspeccionHallazgoSST.activo.is_(True)).all()
    elements.append(Paragraph("Hallazgos", styles["H2SST"]))
    hrows = [["Descripción", "Riesgo", "Responsable", "Estado"]] + [[h.descripcion or "", h.nivel_riesgo or "", h.responsable or "", h.estado or ""] for h in hall]
    elements.append(_pdf_table(hrows if len(hrows) > 1 else hrows + [["Sin hallazgos", "", "", ""]], [8*cm, 2.2*cm, 3.5*cm, 2.3*cm]))
    SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1.5*cm, leftMargin=1.5*cm, topMargin=1.2*cm, bottomMargin=1.2*cm).build(elements)
    return _pdf_response(buffer, f"inspeccion_sst_{item.codigo}.pdf")


@router.get("/exportaciones/{inspeccion_id}/acta-pdf")
def exportar_acta_pdf_inspeccion(inspeccion_id: int, db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Spacer, Paragraph
    item = _query_inspecciones(db).filter(InspeccionSST.id == inspeccion_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Inspección no encontrada")
    resp = _inspeccion_to_response(db, item)
    buffer, elements, styles = _pdf_doc("ACTA OFICIAL DE INSPECCIÓN SST", f"Código: {item.codigo} — {item.titulo}")
    datos = [["Empresa", resp.empresa_nombre or "", "Fecha", _fmt(item.fecha_inspeccion)], ["Sede", resp.sede_nombre or "", "Área", resp.area_nombre or ""], ["Lugar", item.lugar or "", "Responsable", item.responsable or ""], ["Resultado", item.resultado or "", "Riesgo", item.nivel_riesgo or ""], ["Cumplimiento", f"{float(item.cumplimiento or 0):.1f}%", "Estado", item.estado or ""]]
    elements.append(_pdf_table(datos, [3.2*cm, 5.2*cm, 3.0*cm, 5.2*cm], repeat=0))
    elements.append(Spacer(1, 0.25*cm))
    elements.append(Paragraph("Observaciones", styles["H2SST"]))
    elements.append(Paragraph(item.observaciones or "Sin observaciones registradas.", styles["Normal"]))
    elements.append(Spacer(1, 0.35*cm))
    firmas = [["Inspector", "Responsable Área", "Responsable SST"], [_firma_flowable(item.firma_inspector), _firma_flowable(item.firma_responsable_area), _firma_flowable(item.firma_sst)], [item.firma_inspector_nombre or "Pendiente", item.firma_responsable_area_nombre or "Pendiente", item.firma_sst_nombre or "Pendiente"], [_fmt(item.firma_inspector_fecha), _fmt(item.firma_responsable_area_fecha), _fmt(item.firma_sst_fecha)]]
    elements.append(_pdf_table(firmas, [5.5*cm, 5.5*cm, 5.5*cm], repeat=1))
    elements.append(Spacer(1, 0.25*cm))
    elements.append(Paragraph("Trazabilidad", styles["H2SST"]))
    elements.append(Paragraph((item.trazabilidad or "Sin trazabilidad registrada.").replace("\n", "<br/>") , styles["Normal"]))
    SimpleDocTemplate(buffer, pagesize=A4, rightMargin=1.5*cm, leftMargin=1.5*cm, topMargin=1.2*cm, bottomMargin=1.2*cm).build(elements)
    return _pdf_response(buffer, f"acta_inspeccion_sst_{item.codigo}.pdf")

@router.post("/{inspeccion_id}/firmas", response_model=InspeccionResponse)
def registrar_firma_inspeccion(
    inspeccion_id: int,
    data: InspeccionFirmaRequest,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    item = db.query(InspeccionSST).filter(InspeccionSST.id == inspeccion_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Inspección no encontrada")
    ahora = datetime.utcnow()
    rol = data.rol_firma
    if rol == "INSPECTOR":
        item.firma_inspector = data.firma_base64
        item.firma_inspector_nombre = data.nombre_firmante
        item.firma_inspector_fecha = ahora
    elif rol == "RESPONSABLE_AREA":
        item.firma_responsable_area = data.firma_base64
        item.firma_responsable_area_nombre = data.nombre_firmante
        item.firma_responsable_area_fecha = ahora
    elif rol == "SST":
        item.firma_sst = data.firma_base64
        item.firma_sst_nombre = data.nombre_firmante
        item.firma_sst_fecha = ahora
    linea = f"[{ahora.isoformat()}] Firma {rol}: {data.nombre_firmante}. {data.observacion or ''}"
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + linea
    if item.estado == "PROGRAMADA":
        item.estado = "EN_PROCESO"
    db.commit()
    db.refresh(item)
    return obtener_inspeccion(item.id, db, usuario)


@router.post("/{inspeccion_id}/cierre-digital", response_model=InspeccionResponse)
def cerrar_digitalmente_inspeccion(
    inspeccion_id: int,
    data: InspeccionCierreDigitalRequest,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    item = db.query(InspeccionSST).filter(InspeccionSST.id == inspeccion_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Inspección no encontrada")
    faltantes = []
    if not item.firma_inspector:
        faltantes.append("Inspector")
    if not item.firma_responsable_area:
        faltantes.append("Responsable Área")
    if not item.firma_sst:
        faltantes.append("SST")
    if faltantes:
        raise HTTPException(status_code=400, detail=f"No se puede cerrar. Faltan firmas: {', '.join(faltantes)}")
    ahora = datetime.utcnow()
    item.cierre_digital = True
    item.cierre_digital_fecha = ahora
    item.cierre_digital_usuario_id = getattr(usuario, "id", None)
    item.estado = "CERRADA"
    linea = f"[{ahora.isoformat()}] Cierre digital por usuario {getattr(usuario, 'id', '')}. {data.observacion or ''}"
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + linea
    db.commit()
    db.refresh(item)
    return obtener_inspeccion(item.id, db, usuario)

@router.get("/{inspeccion_id}", response_model=InspeccionResponse)
def obtener_inspeccion(inspeccion_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = _query_inspecciones(db).filter(InspeccionSST.id == inspeccion_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Inspección no encontrada")
    return _inspeccion_to_response(db, item)


@router.post("/", response_model=InspeccionResponse)
def crear_inspeccion(data: InspeccionCreate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    _validar_empresa(db, data.empresa_id)
    _validar_opcional(db, Sede, data.sede_id, "Sede")
    _validar_opcional(db, Area, data.area_id, "Área")
    _validar_opcional(db, Cargo, data.cargo_id, "Cargo")
    _validar_opcional(db, Empleado, data.empleado_id, "Empleado")
    existe = db.query(InspeccionSST).filter(InspeccionSST.empresa_id == data.empresa_id, func.upper(InspeccionSST.codigo) == data.codigo.upper()).first()
    if existe:
        raise HTTPException(status_code=400, detail="Ya existe una inspección con ese código para la empresa")
    payload = data.model_dump()
    payload["usuario_id"] = payload.get("usuario_id") or getattr(usuario, "id", None)
    item = InspeccionSST(**payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return obtener_inspeccion(item.id, db, usuario)


@router.put("/{inspeccion_id}", response_model=InspeccionResponse)
def actualizar_inspeccion(inspeccion_id: int, data: InspeccionUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = db.query(InspeccionSST).filter(InspeccionSST.id == inspeccion_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Inspección no encontrada")
    if item.cierre_digital:
        raise HTTPException(status_code=400, detail="La inspección tiene cierre digital y no puede modificarse")
    payload = data.model_dump(exclude_unset=True)
    for key, value in payload.items():
        setattr(item, key, value)
    db.commit()
    db.refresh(item)
    return obtener_inspeccion(item.id, db, usuario)


@router.delete("/{inspeccion_id}")
def eliminar_inspeccion(inspeccion_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    item = db.query(InspeccionSST).filter(InspeccionSST.id == inspeccion_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Inspección no encontrada")
    if item.cierre_digital:
        raise HTTPException(status_code=400, detail="La inspección tiene cierre digital y no puede anularse")
    item.activo = False
    item.estado = "ANULADA"
    db.commit()
    return {"ok": True, "message": "Inspección anulada"}


@router.get("/{inspeccion_id}/hallazgos", response_model=list[HallazgoResponse])
def listar_hallazgos(inspeccion_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    items = db.query(InspeccionHallazgoSST).filter(InspeccionHallazgoSST.inspeccion_id == inspeccion_id, InspeccionHallazgoSST.activo.is_(True)).order_by(InspeccionHallazgoSST.id.desc()).all()
    return items


@router.post("/{inspeccion_id}/hallazgos", response_model=HallazgoResponse)
def crear_hallazgo(inspeccion_id: int, data: HallazgoCreate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    inspeccion = db.query(InspeccionSST).filter(InspeccionSST.id == inspeccion_id).first()
    if not inspeccion:
        raise HTTPException(status_code=404, detail="Inspección no encontrada")
    if inspeccion.cierre_digital:
        raise HTTPException(status_code=400, detail="La inspección tiene cierre digital y no permite nuevos hallazgos")
    payload = data.model_dump()

# Compatibilidad Schema ↔ Modelo
    payload.pop("firma_inspector", None)
    payload.pop("firma_inspector_nombre", None)

    payload.pop("firma_responsable_area", None)
    payload.pop("firma_responsable_area_nombre", None)

    payload.pop("firma_sst", None)
    payload.pop("firma_sst_nombre", None)

    payload.pop("cierre_digital", None)
    payload.pop("trazabilidad", None)

    payload["inspeccion_id"] = inspeccion_id
    payload["empresa_id"] = inspeccion.empresa_id

    item = InspeccionHallazgoSST(**payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


@router.put("/hallazgos/{hallazgo_id}", response_model=HallazgoResponse)
def actualizar_hallazgo(hallazgo_id: int, data: HallazgoUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = db.query(InspeccionHallazgoSST).filter(InspeccionHallazgoSST.id == hallazgo_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Hallazgo no encontrado")
    if item.inspeccion and item.inspeccion.cierre_digital:
        raise HTTPException(status_code=400, detail="La inspección tiene cierre digital y no permite modificar hallazgos")
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(item, key, value)
    if item.estado == "CERRADO" and not item.fecha_cierre:
        item.fecha_cierre = date.today()
    db.commit()
    db.refresh(item)
    return item


@router.delete("/hallazgos/{hallazgo_id}")
def eliminar_hallazgo(hallazgo_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    item = db.query(InspeccionHallazgoSST).filter(InspeccionHallazgoSST.id == hallazgo_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Hallazgo no encontrado")
    if item.inspeccion and item.inspeccion.cierre_digital:
        raise HTTPException(status_code=400, detail="La inspección tiene cierre digital y no permite anular hallazgos")
    item.activo = False
    item.estado = "ANULADO"
    db.commit()
    return {"ok": True, "message": "Hallazgo anulado"}


@router.get("/{inspeccion_id}/evidencias")
def listar_evidencias(inspeccion_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    archivos = db.query(ArchivoSST).filter(ArchivoSST.modulo == "INSPECCIONES", ArchivoSST.referencia_id == inspeccion_id, ArchivoSST.activo.is_(True)).order_by(ArchivoSST.fecha_creacion.desc()).all()
    return [_archivo_to_dict(a) for a in archivos]


@router.post("/{inspeccion_id}/evidencias")
def subir_evidencia(
    inspeccion_id: int,
    tipo_evidencia: str = Form(default="EVIDENCIA"),
    descripcion: str = Form(default=""),
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    inspeccion = db.query(InspeccionSST).filter(InspeccionSST.id == inspeccion_id).first()
    if not inspeccion:
        raise HTTPException(status_code=404, detail="Inspección no encontrada")
    if inspeccion.cierre_digital:
        raise HTTPException(status_code=400, detail="La inspección tiene cierre digital y no permite subir evidencias")
    path, original, filename, mime_type, size = _guardar_upload(archivo)
    registro = ArchivoSST(
        empresa_id=inspeccion.empresa_id,
        usuario_id=getattr(usuario, "id", None),
        tipo=(tipo_evidencia or "EVIDENCIA").upper().strip(),
        nombre_original=original,
        nombre_archivo=filename,
        ruta=str(path),
        url=_public_upload_url(path),
        extension=filename.rsplit(".", 1)[-1].lower(),
        mime_type=mime_type,
        tamano_bytes=size,
        modulo="INSPECCIONES",
        referencia_id=inspeccion.id,
        descripcion=descripcion or "Evidencia de inspección SST",
        activo=True,
    )
    db.add(registro)
    db.commit()
    db.refresh(registro)
    return _archivo_to_dict(registro)


@router.delete("/{inspeccion_id}/evidencias/{archivo_id}")
def eliminar_evidencia(inspeccion_id: int, archivo_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    archivo = db.query(ArchivoSST).filter(ArchivoSST.id == archivo_id, ArchivoSST.modulo == "INSPECCIONES", ArchivoSST.referencia_id == inspeccion_id).first()
    if not archivo:
        raise HTTPException(status_code=404, detail="Evidencia no encontrada")
    inspeccion = db.query(InspeccionSST).filter(InspeccionSST.id == inspeccion_id).first()
    if inspeccion and inspeccion.cierre_digital:
        raise HTTPException(status_code=400, detail="La inspección tiene cierre digital y no permite eliminar evidencias")
    archivo.activo = False
    db.commit()
    return {"ok": True, "message": "Evidencia desactivada"}
