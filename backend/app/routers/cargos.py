# ============================================================
# ROUTER CARGOS - ERP SST PRO
# FASE 1.1.4.3 — EXPORTACIÓN PDF / EXCEL
# Archivo: backend/app/routers/cargos.py
# Compatible con Cargos Analytics PRO sin romper CRUD existente.
# ============================================================

from io import BytesIO
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func, or_

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.database import get_db
from app.models.cargo import Cargo
from app.models.empresa import Empresa
from app.models.sede import Sede
from app.models.area import Area
from app.models.empleado import Empleado
from app.models.epp import CargoEPPCatalogo, EPPCatalogo
from app.schemas.cargo_schema import (
    CargoCreate,
    CargoDashboardResponse,
    CargoEPPAsignacionResponse,
    CargoEPPAsignacionUpdate,
    CargoResponse,
    CargoUpdate,
)
from app.schemas.epp_schema import EPPCatalogoResponse
from app.auth.dependencies import require_roles, require_permission
from app.core.default_permissions import PERM_REGISTROS_ELIMINAR, PERM_REPORTES_EXPORTAR

router = APIRouter(prefix="/cargos", tags=["Cargos SST Enterprise 360"])
ROLES_SST = ["SUPER_ADMIN", "ADMIN_EMPRESA", "RESPONSABLE_SST"]
EXPORTAR_REPORTES = require_permission(PERM_REPORTES_EXPORTAR)
ELIMINAR_REGISTROS = require_permission(PERM_REGISTROS_ELIMINAR)


# ============================================================
# Helpers de compatibilidad
# ============================================================
def _texto(valor, defecto=""):
    if valor is None:
        return defecto
    return str(valor).strip() or defecto


def _si_no(valor):
    return "Sí" if bool(valor) else "No"


def _cargo_to_response(cargo: Cargo) -> CargoResponse:
    data = CargoResponse.model_validate(cargo)
    data.empresa_nombre = cargo.empresa.nombre if cargo.empresa else None
    data.sede_nombre = cargo.sede.nombre if cargo.sede else None
    data.area_nombre = cargo.area.nombre if cargo.area else None

    # Alias compatibles con el frontend anterior y con Analytics PRO.
    data.codigo = cargo.codigo_cargo
    data.tipo = cargo.tipo_cargo
    data.proceso = cargo.proceso_asociado
    data.empleados_asociados = cargo.numero_empleados or 0
    data.requiere_examen_medico = bool(cargo.examenes_medicos)
    data.requiere_capacitacion = bool(cargo.capacitaciones_requeridas)
    data.funciones = cargo.perfil_sst
    data.observaciones = cargo.exposicion
    return data


def _payload_compatible(data):
    payload = data.model_dump(exclude_unset=True)

    # Compatibilidad con nombres enviados por versiones previas del frontend.
    alias = {
        "codigo": "codigo_cargo",
        "tipo": "tipo_cargo",
        "proceso": "proceso_asociado",
        "empleados_asociados": "numero_empleados",
        "funciones": "perfil_sst",
    }
    for origen, destino in alias.items():
        if origen in payload and destino not in payload:
            payload[destino] = payload.pop(origen)
        elif origen in payload:
            payload.pop(origen, None)

    # Checkboxes antiguos: si llegan activos y no hay texto, deja una marca clara.
    if payload.pop("requiere_examen_medico", False) and not payload.get("examenes_medicos"):
        payload["examenes_medicos"] = "Requiere examen médico ocupacional según exposición del cargo."
    if payload.pop("requiere_capacitacion", False) and not payload.get("capacitaciones_requeridas"):
        payload["capacitaciones_requeridas"] = "Requiere capacitación SST de acuerdo con funciones y nivel de riesgo."

    # Campo observaciones se mantiene como alias legacy si se necesita.
    payload.pop("observaciones", None)

    if "empresa_id" in payload and payload["empresa_id"] in ["", None]:
        payload.pop("empresa_id", None)
    for key in ["sede_id", "area_id"]:
        if payload.get(key) in ["", 0, "0"]:
            payload[key] = None
    if "numero_empleados" in payload:
        payload["numero_empleados"] = int(payload.get("numero_empleados") or 0)

    return payload


def _validar_relaciones(db: Session, data):
    empresa_id = getattr(data, "empresa_id", None)
    sede_id = getattr(data, "sede_id", None)
    area_id = getattr(data, "area_id", None)
    if empresa_id:
        empresa = db.query(Empresa).filter(Empresa.id == empresa_id).first()
        if not empresa:
            raise HTTPException(status_code=404, detail="Empresa no encontrada")
    if sede_id:
        sede = db.query(Sede).filter(Sede.id == sede_id).first()
        if not sede:
            raise HTTPException(status_code=404, detail="Sede no encontrada")
    if area_id:
        area = db.query(Area).filter(Area.id == area_id).first()
        if not area:
            raise HTTPException(status_code=404, detail="Área no encontrada")


def _validar_empresa_usuario(usuario, empresa_id: int) -> None:
    if str(getattr(usuario, "rol", "")).strip().upper() == "SUPER_ADMIN":
        return
    empresa_usuario_id = getattr(usuario, "empresa_id", None)
    if not empresa_usuario_id:
        raise HTTPException(status_code=403, detail="Usuario sin empresa asignada")
    if int(empresa_usuario_id) != int(empresa_id):
        raise HTTPException(status_code=403, detail="No puede acceder a cargos de otra empresa")


def _obtener_cargo_autorizado(db: Session, cargo_id: int, usuario) -> Cargo:
    cargo = db.query(Cargo).filter(Cargo.id == cargo_id).first()
    if not cargo:
        raise HTTPException(status_code=404, detail="Cargo no encontrado")
    _validar_empresa_usuario(usuario, cargo.empresa_id)
    return cargo


def _respuesta_epp_cargo(db: Session, cargo: Cargo) -> CargoEPPAsignacionResponse:
    asociaciones = (
        db.query(CargoEPPCatalogo)
        .options(joinedload(CargoEPPCatalogo.epp).joinedload(EPPCatalogo.empresa))
        .filter(CargoEPPCatalogo.cargo_id == cargo.id)
        .order_by(CargoEPPCatalogo.id.asc())
        .all()
    )
    epps = []
    for asociacion in asociaciones:
        item = EPPCatalogoResponse.model_validate(asociacion.epp)
        item.empresa_nombre = asociacion.epp.empresa.nombre if asociacion.epp.empresa else None
        epps.append(item)
    return CargoEPPAsignacionResponse(
        cargo_id=cargo.id,
        empresa_id=cargo.empresa_id,
        epp_ids=[asociacion.epp_id for asociacion in asociaciones],
        epps=epps,
    )


def _query_cargos_filtrada(
    db: Session,
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    estado: str | None = None,
    riesgo: str | None = None,
    tipo: str | None = None,
    q: str | None = None,
):
    query = db.query(Cargo)
    if empresa_id:
        query = query.filter(Cargo.empresa_id == empresa_id)
    if sede_id:
        query = query.filter(Cargo.sede_id == sede_id)
    if area_id:
        query = query.filter(Cargo.area_id == area_id)
    if estado in ["ACTIVO", "INACTIVO"]:
        query = query.filter(Cargo.activo == (estado == "ACTIVO"))
    if riesgo:
        query = query.filter(func.upper(Cargo.nivel_riesgo) == riesgo.upper())
    if tipo:
        query = query.filter(func.upper(Cargo.tipo_cargo) == tipo.upper())
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(
                Cargo.nombre.ilike(like),
                Cargo.codigo_cargo.ilike(like),
                Cargo.descripcion.ilike(like),
                Cargo.proceso_asociado.ilike(like),
            )
        )
    return query.order_by(Cargo.id.desc())


def _nombre_archivo(prefijo: str, extension: str):
    fecha = datetime.now().strftime("%Y%m%d_%H%M")
    return f"{prefijo}_{fecha}.{extension}"


# ============================================================
# CRUD + Analytics
# ============================================================
@router.post("/", response_model=CargoResponse)
def crear_cargo(data: CargoCreate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    _validar_relaciones(db, data)
    cargo = Cargo(**_payload_compatible(data))
    db.add(cargo)
    db.commit()
    db.refresh(cargo)
    return _cargo_to_response(cargo)


@router.get("/", response_model=list[CargoResponse])
def listar_cargos(
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    estado: str | None = None,
    riesgo: str | None = None,
    tipo: str | None = None,
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    cargos = _query_cargos_filtrada(db, empresa_id, sede_id, area_id, estado, riesgo, tipo, q).all()
    return [_cargo_to_response(c) for c in cargos]


@router.get("/dashboard", response_model=CargoDashboardResponse)
def dashboard_cargos(
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    cargos = _query_cargos_filtrada(db, empresa_id, sede_id, area_id).all()
    total = len(cargos)
    activos = sum(1 for c in cargos if c.activo)
    inactivos = total - activos
    alto_critico = sum(1 for c in cargos if (c.nivel_riesgo or "").upper() in ["ALTO", "CRITICO", "CRÍTICO"])
    requieren_epp = sum(1 for c in cargos if c.requiere_epp)
    sin_area = sum(1 for c in cargos if not c.area_id)
    sin_codigo = sum(1 for c in cargos if not c.codigo_cargo)
    sin_examenes = sum(1 for c in cargos if not c.examenes_medicos)
    sin_capacitaciones = sum(1 for c in cargos if not c.capacitaciones_requeridas)
    empleados_bd = db.query(func.count(Empleado.id)).filter(Empleado.cargo_id.in_([c.id for c in cargos])).scalar() if cargos else 0
    empleados_manual = sum(int(c.numero_empleados or 0) for c in cargos)
    empleados = empleados_bd or empleados_manual

    def agrupar(attr, default="Sin dato"):
        tmp = {}
        for c in cargos:
            key = getattr(c, attr, None) or default
            tmp[key] = tmp.get(key, 0) + 1
        return [{"nombre": k, "total": v} for k, v in sorted(tmp.items(), key=lambda x: x[1], reverse=True)]

    riesgos = agrupar("nivel_riesgo", "Sin riesgo")
    tipos = agrupar("tipo_cargo", "Sin tipo")
    areas = {}
    for c in cargos:
        key = c.area.nombre if c.area else "Sin área"
        areas[key] = areas.get(key, 0) + 1
    cargos_por_area = [{"nombre": k, "total": v} for k, v in sorted(areas.items(), key=lambda x: x[1], reverse=True)]

    penalizaciones = alto_critico + sin_area + sin_codigo + sin_examenes + sin_capacitaciones
    indice = 100 if total == 0 else max(0, min(100, round(100 - (penalizaciones / max(total * 5, 1)) * 100)))

    def peso(c):
        return (
            (3 if (c.nivel_riesgo or "").upper() in ["CRITICO", "CRÍTICO"] else 2 if (c.nivel_riesgo or "").upper() == "ALTO" else 0)
            + (1 if not c.area_id else 0)
            + (1 if not c.codigo_cargo else 0)
            + (1 if c.requiere_epp and not c.epp_requerido else 0)
        )

    cargos_prioritarios = []
    for c in sorted(cargos, key=peso, reverse=True)[:5]:
        cargos_prioritarios.append({
            "id": c.id,
            "nombre": c.nombre,
            "riesgo": c.nivel_riesgo or "MEDIO",
            "empleados": c.numero_empleados or 0,
            "area": c.area.nombre if c.area else "Sin área",
        })

    recomendaciones = []
    if alto_critico:
        recomendaciones.append("Priorizar controles para cargos con riesgo alto o crítico.")
    if sin_area:
        recomendaciones.append("Asignar área a los cargos pendientes para mejorar trazabilidad del SG-SST.")
    if sin_codigo:
        recomendaciones.append("Normalizar codificación de cargos para trazabilidad documental.")
    if requieren_epp:
        recomendaciones.append("Verificar matriz EPP para cargos con elementos de protección obligatorios.")
    if sin_examenes:
        recomendaciones.append("Definir exámenes médicos ocupacionales según exposición del cargo.")
    if sin_capacitaciones:
        recomendaciones.append("Asignar capacitaciones obligatorias por cargo y nivel de riesgo.")
    if not recomendaciones:
        recomendaciones.append("La estructura de cargos se encuentra controlada. Mantener seguimiento preventivo.")

    return CargoDashboardResponse(
        total_cargos=total,
        activos=activos,
        inactivos=inactivos,
        alto_critico=alto_critico,
        riesgo_alto_critico=alto_critico,
        requieren_epp=requieren_epp,
        requieren_examen_medico=total - sin_examenes,
        requieren_capacitacion=total - sin_capacitaciones,
        empleados_asociados=empleados or 0,
        sin_area=sin_area,
        sin_codigo=sin_codigo,
        sin_examenes=sin_examenes,
        sin_capacitaciones=sin_capacitaciones,
        indice_gestion=indice,
        cargos_por_riesgo=riesgos,
        distribucion_riesgo=riesgos,
        cargos_por_tipo=tipos,
        distribucion_tipo=tipos,
        cargos_por_area=cargos_por_area,
        cargo_prioritario=cargos_prioritarios[0] if cargos_prioritarios else None,
        cargos_prioritarios=cargos_prioritarios,
        recomendaciones=recomendaciones,
    )


# Alias para frontend que consulta /dashboard/resumen.
@router.get("/dashboard/resumen", response_model=CargoDashboardResponse)
def dashboard_cargos_resumen(
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    return dashboard_cargos(empresa_id=empresa_id, sede_id=sede_id, area_id=area_id, db=db, usuario=usuario)


# ============================================================
# FASE 1.1.4.3 — Exportación Excel / PDF
# IMPORTANTE: estas rutas van antes de /{cargo_id}
# ============================================================
@router.get("/export/excel")
def exportar_cargos_excel(
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    estado: str | None = None,
    riesgo: str | None = None,
    tipo: str | None = None,
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    cargos = _query_cargos_filtrada(db, empresa_id, sede_id, area_id, estado, riesgo, tipo, q).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Cargos SST"

    ws.merge_cells("A1:Q1")
    ws["A1"] = "ERP SST PRO - MATRIZ DE CARGOS SST"
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=15)
    ws["A1"].fill = PatternFill("solid", fgColor="123A7A")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:Q2")
    ws["A2"] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}  |  Total registros: {len(cargos)}"
    ws["A2"].font = Font(italic=True, color="475569")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "ID", "Empresa", "Sede", "Área", "Código", "Cargo", "Tipo", "Riesgo", "Exposición",
        "Proceso", "Empleados", "Requiere EPP", "EPP requerido", "Exámenes médicos",
        "Capacitaciones", "Estado", "Fecha creación",
    ]
    ws.append(headers)
    header_row = 3
    header_fill = PatternFill("solid", fgColor="EAF2FF")
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="0F172A")
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for c in cargos:
        ws.append([
            c.id,
            c.empresa.nombre if c.empresa else "Sin empresa",
            c.sede.nombre if c.sede else "Sin sede",
            c.area.nombre if c.area else "Sin área",
            c.codigo_cargo or "SIN-CÓDIGO",
            c.nombre,
            c.tipo_cargo or "OPERATIVO",
            c.nivel_riesgo or "MEDIO",
            c.exposicion or "Sin exposición",
            c.proceso_asociado or "Sin proceso",
            c.numero_empleados or 0,
            _si_no(c.requiere_epp),
            c.epp_requerido or "",
            c.examenes_medicos or "",
            c.capacitaciones_requeridas or "",
            "ACTIVO" if c.activo else "INACTIVO",
            c.fecha_creacion.strftime("%d/%m/%Y") if c.fecha_creacion else "",
        ])

    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [8, 24, 20, 20, 16, 28, 16, 14, 16, 24, 12, 14, 30, 35, 35, 14, 16]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:Q{max(3, ws.max_row)}"

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = _nombre_archivo("cargos_sst", "xlsx")
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/pdf")
def exportar_cargos_pdf(
    empresa_id: int | None = None,
    sede_id: int | None = None,
    area_id: int | None = None,
    estado: str | None = None,
    riesgo: str | None = None,
    tipo: str | None = None,
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(EXPORTAR_REPORTES),
):
    cargos = _query_cargos_filtrada(db, empresa_id, sede_id, area_id, estado, riesgo, tipo, q).all()
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=1 * cm, leftMargin=1 * cm, topMargin=1 * cm, bottomMargin=1 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("CargoTitle", parent=styles["Title"], textColor=colors.HexColor("#123A7A"), alignment=TA_CENTER, fontSize=16)
    small_style = ParagraphStyle("CargoSmall", parent=styles["Normal"], fontSize=7, leading=9, alignment=TA_LEFT)

    elements = [
        Paragraph("ERP SST PRO - Reporte Ejecutivo de Cargos SST", title_style),
        Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Total registros: {len(cargos)}", styles["Normal"]),
        Spacer(1, 0.25 * cm),
    ]

    data = [["Código", "Cargo", "Empresa", "Sede", "Área", "Tipo", "Riesgo", "Empl.", "EPP", "Estado"]]
    for c in cargos:
        data.append([
            Paragraph(_texto(c.codigo_cargo, "SIN-CÓDIGO"), small_style),
            Paragraph(_texto(c.nombre), small_style),
            Paragraph(_texto(c.empresa.nombre if c.empresa else None, "Sin empresa"), small_style),
            Paragraph(_texto(c.sede.nombre if c.sede else None, "Sin sede"), small_style),
            Paragraph(_texto(c.area.nombre if c.area else None, "Sin área"), small_style),
            Paragraph(_texto(c.tipo_cargo, "OPERATIVO"), small_style),
            Paragraph(_texto(c.nivel_riesgo, "MEDIO"), small_style),
            str(c.numero_empleados or 0),
            _si_no(c.requiere_epp),
            "ACTIVO" if c.activo else "INACTIVO",
        ])

    if len(data) == 1:
        data.append(["Sin registros", "", "", "", "", "", "", "", "", ""])

    table = Table(data, repeatRows=1, colWidths=[2.4 * cm, 4.4 * cm, 4 * cm, 3.4 * cm, 3.4 * cm, 2.5 * cm, 2.2 * cm, 1.4 * cm, 1.4 * cm, 2.2 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123A7A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("FONTSIZE", (0, 1), (-1, -1), 7),
    ]))
    elements.append(table)
    doc.build(elements)
    stream.seek(0)

    filename = _nombre_archivo("reporte_cargos_sst", "pdf")
    return StreamingResponse(stream, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})


# ============================================================
# Obtener, ficha PDF, actualizar y desactivar
# ============================================================
@router.get("/{cargo_id}/epp", response_model=CargoEPPAsignacionResponse)
def obtener_epp_cargo(cargo_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    cargo = _obtener_cargo_autorizado(db, cargo_id, usuario)
    return _respuesta_epp_cargo(db, cargo)


@router.put("/{cargo_id}/epp", response_model=CargoEPPAsignacionResponse)
def actualizar_epp_cargo(
    cargo_id: int,
    data: CargoEPPAsignacionUpdate,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    cargo = _obtener_cargo_autorizado(db, cargo_id, usuario)
    epp_ids = data.epp_ids
    epps = []
    if epp_ids:
        epps = (
            db.query(EPPCatalogo)
            .filter(
                EPPCatalogo.id.in_(epp_ids),
                EPPCatalogo.empresa_id == cargo.empresa_id,
                EPPCatalogo.activo == True,
            )
            .all()
        )
        if len(epps) != len(epp_ids):
            raise HTTPException(status_code=404, detail="Uno o más EPP no existen, están inactivos o pertenecen a otra empresa")

    db.query(CargoEPPCatalogo).filter(CargoEPPCatalogo.cargo_id == cargo.id).delete(synchronize_session=False)
    db.add_all(
        [
            CargoEPPCatalogo(empresa_id=cargo.empresa_id, cargo_id=cargo.id, epp_id=epp_id)
            for epp_id in epp_ids
        ]
    )
    cargo.requiere_epp = bool(epp_ids)
    db.commit()
    return _respuesta_epp_cargo(db, cargo)


@router.get("/{cargo_id}", response_model=CargoResponse)
def obtener_cargo(cargo_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    cargo = db.query(Cargo).filter(Cargo.id == cargo_id).first()
    if not cargo:
        raise HTTPException(status_code=404, detail="Cargo no encontrado")
    return _cargo_to_response(cargo)


@router.get("/{cargo_id}/export/pdf")
def exportar_ficha_cargo_pdf(cargo_id: int, db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    cargo = db.query(Cargo).filter(Cargo.id == cargo_id).first()
    if not cargo:
        raise HTTPException(status_code=404, detail="Cargo no encontrado")

    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, rightMargin=1.4 * cm, leftMargin=1.4 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("FichaTitle", parent=styles["Title"], textColor=colors.HexColor("#123A7A"), alignment=TA_CENTER, fontSize=17)
    label_style = ParagraphStyle("Label", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#475569"), leading=10)
    value_style = ParagraphStyle("Value", parent=styles["Normal"], fontSize=9, leading=11)

    def p(text, style=value_style):
        return Paragraph(_texto(text, "—"), style)

    elements = [
        Paragraph("Ficha Técnica del Cargo SST", title_style),
        Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]),
        Spacer(1, 0.3 * cm),
    ]

    general = [
        [p("Empresa", label_style), p(cargo.empresa.nombre if cargo.empresa else "Sin empresa"), p("Sede", label_style), p(cargo.sede.nombre if cargo.sede else "Sin sede")],
        [p("Área", label_style), p(cargo.area.nombre if cargo.area else "Sin área"), p("Código", label_style), p(cargo.codigo_cargo or "SIN-CÓDIGO")],
        [p("Cargo", label_style), p(cargo.nombre), p("Tipo", label_style), p(cargo.tipo_cargo or "OPERATIVO")],
        [p("Riesgo", label_style), p(cargo.nivel_riesgo or "MEDIO"), p("Exposición", label_style), p(cargo.exposicion or "Sin exposición")],
        [p("Empleados", label_style), p(str(cargo.numero_empleados or 0)), p("Estado", label_style), p("ACTIVO" if cargo.activo else "INACTIVO")],
    ]
    table = Table(general, colWidths=[3 * cm, 6 * cm, 3 * cm, 6 * cm])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F1F5F9")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.35 * cm))

    bloques = [
        ("Descripción", cargo.descripcion),
        ("Proceso asociado", cargo.proceso_asociado),
        ("EPP requerido", cargo.epp_requerido if cargo.requiere_epp else "No registra EPP obligatorio."),
        ("Exámenes médicos", cargo.examenes_medicos),
        ("Capacitaciones requeridas", cargo.capacitaciones_requeridas),
        ("Perfil SST", cargo.perfil_sst),
        ("Competencias", cargo.competencias),
    ]
    for titulo, contenido in bloques:
        elements.append(Paragraph(titulo, ParagraphStyle(f"h-{titulo}", parent=styles["Heading3"], textColor=colors.HexColor("#123A7A"), fontSize=10)))
        elements.append(Paragraph(_texto(contenido, "Sin información registrada."), value_style))
        elements.append(Spacer(1, 0.18 * cm))

    doc.build(elements)
    stream.seek(0)
    codigo = (cargo.codigo_cargo or f"cargo_{cargo.id}").replace(" ", "_")
    filename = f"ficha_cargo_sst_{codigo}.pdf"
    return StreamingResponse(stream, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.put("/{cargo_id}", response_model=CargoResponse)
def actualizar_cargo(cargo_id: int, data: CargoUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    cargo = db.query(Cargo).filter(Cargo.id == cargo_id).first()
    if not cargo:
        raise HTTPException(status_code=404, detail="Cargo no encontrado")
    _validar_relaciones(db, data)
    payload = _payload_compatible(data)
    for key, value in payload.items():
        setattr(cargo, key, value)
    db.commit()
    db.refresh(cargo)
    return _cargo_to_response(cargo)


@router.patch("/{cargo_id}/estado", response_model=CargoResponse)
def cambiar_estado_cargo(cargo_id: int, activo: bool, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    cargo = db.query(Cargo).filter(Cargo.id == cargo_id).first()
    if not cargo:
        raise HTTPException(status_code=404, detail="Cargo no encontrado")
    cargo.activo = activo
    db.commit()
    db.refresh(cargo)
    return _cargo_to_response(cargo)


@router.delete("/{cargo_id}")
def eliminar_cargo(cargo_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    cargo = db.query(Cargo).filter(Cargo.id == cargo_id).first()
    if not cargo:
        raise HTTPException(status_code=404, detail="Cargo no encontrado")
    cargo.activo = False
    db.commit()
    return {"mensaje": "Cargo desactivado correctamente"}
