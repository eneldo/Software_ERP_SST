# ============================================================
# ROUTER CAPA SST ENTERPRISE - ERP SST PRO
# FASE 1.1.8.7.1 — Optimización CAPA Enterprise
# Archivo: backend/app/routers/capa.py
# ============================================================

from datetime import date, datetime
from pathlib import Path
import io
import os
import uuid

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, joinedload

from app.auth.dependencies import require_roles, require_permission
from app.core.default_permissions import PERM_REGISTROS_ELIMINAR, PERM_REPORTES_EXPORTAR
from app.database import get_db
from app.models.archivo_sst import ArchivoSST
from app.models.area import Area
from app.models.capa import CapaSST, CapaSeguimientoSST
from app.models.cargo import Cargo
from app.models.empleado import Empleado
from app.models.empresa import Empresa
from app.models.inspeccion import InspeccionHallazgoSST, InspeccionSST
from app.models.sede import Sede
from app.schemas.capa_schema import (
    CapaCierreRequest,
    CapaCreate,
    CapaDashboardResponse,
    CapaResponse,
    CapaSeguimientoCreate,
    CapaSeguimientoResponse,
    CapaSeguimientoUpdate,
    CapaUpdate,
)

router = APIRouter(prefix="/capa", tags=["CAPA SST Enterprise"])
ROLES_SST = ["SUPER_ADMIN", "ADMIN_EMPRESA", "RESPONSABLE_SST"]
EXPORTAR_REPORTES = require_permission(PERM_REPORTES_EXPORTAR)
ELIMINAR_REGISTROS = require_permission(PERM_REGISTROS_ELIMINAR)

UPLOAD_ROOT = Path(os.getenv("UPLOAD_DIR", "app/uploads")).resolve()
CAPA_UPLOAD_DIR = UPLOAD_ROOT / "capa"
CAPA_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXT = {"pdf", "jpg", "jpeg", "png", "webp", "xlsx", "xls", "csv", "docx", "doc"}
ALLOWED_MIME_PREFIX = {"image/", "application/pdf", "application/vnd", "application/msword"}
MAX_UPLOAD_MB = 25


def _upper(value, default=None):
    if value is None:
        return default
    value = str(value).strip().upper()
    return value if value else default


def _es_super_admin(usuario) -> bool:
    return str(getattr(usuario, "rol", "") or "").upper() == "SUPER_ADMIN"


def _empresa_usuario_id(usuario) -> int | None:
    return getattr(usuario, "empresa_id", None)


def _validar_empresa_usuario(usuario, empresa_id: int | None) -> None:
    """Aislamiento multiempresa: SUPER_ADMIN ve todo; otros roles solo su empresa."""
    if not empresa_id or _es_super_admin(usuario):
        return
    usuario_empresa_id = _empresa_usuario_id(usuario)
    if usuario_empresa_id is not None and int(usuario_empresa_id) != int(empresa_id):
        raise HTTPException(status_code=403, detail="No tiene permisos sobre esta empresa")


def _filtrar_empresa_usuario(query, usuario, model):
    if usuario is None or _es_super_admin(usuario):
        return query
    usuario_empresa_id = _empresa_usuario_id(usuario)
    if usuario_empresa_id is None:
        raise HTTPException(status_code=403, detail="Usuario sin empresa asignada para operación multiempresa")
    return query.filter(model.empresa_id == usuario_empresa_id)


def _public_upload_url(file_path: Path) -> str:
    try:
        rel = file_path.resolve().relative_to(UPLOAD_ROOT)
        return "/uploads/" + rel.as_posix()
    except Exception:
        return "/uploads/capa/" + file_path.name


def _image_to_webp_bytes(image, max_size: tuple[int, int], quality: int) -> bytes:
    from PIL import Image

    img = image.copy()
    if img.mode not in ("RGB", "RGBA"):
        img = img.convert("RGB")
    img.thumbnail(max_size)
    out = io.BytesIO()
    img.save(out, format="WEBP", quality=quality, method=6, optimize=True)
    return out.getvalue()


def _guardar_upload(upload: UploadFile) -> tuple[Path, str, str, str, int]:
    """
    Guarda evidencias CAPA con optimización Enterprise.

    Imágenes:
    - Archivo principal: WEBP máximo 1920x1080.
    - Preview: WEBP máximo 1280x900 con sufijo _preview.
    - Miniatura: WEBP 220x220 con sufijo _thumb.

    Documentos:
    - Se guardan sin transformación.
    - El frontend muestra icono o visor PDF según mime_type.
    """
    original = upload.filename or "evidencia_capa"
    original = Path(original).name.replace("\x00", "")
    if ".." in original or "/" in original or "\\" in original:
        raise HTTPException(status_code=400, detail="Nombre de archivo no permitido")
    extension = original.rsplit(".", 1)[-1].lower() if "." in original else "bin"
    if extension not in ALLOWED_EXT:
        raise HTTPException(status_code=400, detail="Archivo no permitido para CAPA")

    content = upload.file.read()
    if len(content) > MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(status_code=400, detail=f"El archivo supera {MAX_UPLOAD_MB} MB")

    mime_type = upload.content_type or "application/octet-stream"
    office_ext = {"xlsx", "xls", "csv", "docx", "doc"}
    if not any(mime_type.startswith(prefix) for prefix in ALLOWED_MIME_PREFIX) and extension not in office_ext:
        raise HTTPException(status_code=400, detail="Tipo de archivo no permitido")

    uid = uuid.uuid4().hex

    if extension in {"jpg", "jpeg", "png", "webp"}:
        try:
            from PIL import Image
            image = Image.open(io.BytesIO(content))
            main_bytes = _image_to_webp_bytes(image, (1920, 1080), 82)
            preview_bytes = _image_to_webp_bytes(image, (1280, 900), 78)
            thumb_bytes = _image_to_webp_bytes(image, (220, 220), 72)

            filename = f"{uid}.webp"
            path = CAPA_UPLOAD_DIR / filename
            path.write_bytes(main_bytes)
            (CAPA_UPLOAD_DIR / f"{uid}_preview.webp").write_bytes(preview_bytes)
            (CAPA_UPLOAD_DIR / f"{uid}_thumb.webp").write_bytes(thumb_bytes)
            return path, original, filename, "image/webp", len(main_bytes)
        except Exception:
            pass

    filename = f"{uid}.{extension}"
    path = CAPA_UPLOAD_DIR / filename
    path.write_bytes(content)
    return path, original, filename, mime_type, len(content)


def _archivo_variant_url(archivo: ArchivoSST, suffix: str) -> str | None:
    if not archivo.nombre_archivo or not str(archivo.mime_type or "").startswith("image/"):
        return None
    path = Path(archivo.ruta or "")
    base_dir = path.parent if path.parent else CAPA_UPLOAD_DIR
    stem = Path(archivo.nombre_archivo).stem
    variant = base_dir / f"{stem}_{suffix}.webp"
    if variant.exists():
        return _public_upload_url(variant)
    return None


def _archivo_to_dict(archivo: ArchivoSST):
    preview_url = _archivo_variant_url(archivo, "preview") or archivo.url
    thumbnail_url = _archivo_variant_url(archivo, "thumb")
    return {
        "id": archivo.id,
        "empresa_id": archivo.empresa_id,
        "usuario_id": archivo.usuario_id,
        "tipo": archivo.tipo,
        "nombre_original": archivo.nombre_original,
        "nombre_archivo": archivo.nombre_archivo,
        "ruta": archivo.ruta,
        "url": archivo.url,
        "preview_url": preview_url,
        "thumbnail_url": thumbnail_url,
        "extension": archivo.extension,
        "mime_type": archivo.mime_type,
        "tamano_bytes": archivo.tamano_bytes,
        "modulo": archivo.modulo,
        "referencia_id": archivo.referencia_id,
        "descripcion": archivo.descripcion,
        "activo": archivo.activo,
        "fecha_creacion": archivo.fecha_creacion,
    }


def _validar_empresa(db: Session, empresa_id: int):
    empresa = db.query(Empresa).filter(Empresa.id == empresa_id).first()
    if not empresa:
        raise HTTPException(status_code=404, detail="Empresa no encontrada")
    return empresa


def _validar_opcional(db: Session, model, item_id: int | None, label: str):
    if not item_id:
        return None
    item = db.query(model).filter(model.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail=f"{label} no encontrado")
    return item


def _query_capas(db: Session, empresa_id=None, sede_id=None, area_id=None, estado=None, prioridad=None, tipo_accion=None, origen=None, q=None, usuario=None):
    query = db.query(CapaSST).options(
        joinedload(CapaSST.empresa),
        joinedload(CapaSST.sede),
        joinedload(CapaSST.area),
        joinedload(CapaSST.cargo),
        joinedload(CapaSST.empleado),
        joinedload(CapaSST.inspeccion),
        joinedload(CapaSST.hallazgo),
    ).filter(CapaSST.activo.is_(True))
    query = _filtrar_empresa_usuario(query, usuario, CapaSST)
    if empresa_id:
        _validar_empresa_usuario(usuario, empresa_id)
        query = query.filter(CapaSST.empresa_id == empresa_id)
    if sede_id:
        query = query.filter(CapaSST.sede_id == sede_id)
    if area_id:
        query = query.filter(CapaSST.area_id == area_id)
    if estado:
        query = query.filter(func.upper(CapaSST.estado) == estado.upper().strip())
    if prioridad:
        query = query.filter(func.upper(CapaSST.prioridad) == prioridad.upper().strip())
    if tipo_accion:
        query = query.filter(func.upper(CapaSST.tipo_accion) == tipo_accion.upper().strip())
    if origen:
        query = query.filter(func.upper(CapaSST.origen) == origen.upper().strip())
    if q:
        like = f"%{q.strip()}%"
        query = query.filter(or_(CapaSST.codigo.ilike(like), CapaSST.titulo.ilike(like), CapaSST.descripcion.ilike(like), CapaSST.responsable.ilike(like)))
    return query.order_by(CapaSST.id.desc())


def _capa_to_response(db: Session, item: CapaSST):
    data = CapaResponse.model_validate(item)
    data.empresa_nombre = item.empresa.nombre if item.empresa else None
    data.sede_nombre = item.sede.nombre if item.sede else None
    data.area_nombre = item.area.nombre if item.area else None
    data.cargo_nombre = item.cargo.nombre if item.cargo else None
    if item.empleado:
        data.empleado_nombre = f"{item.empleado.nombres} {item.empleado.apellidos}".strip()
    data.inspeccion_codigo = item.inspeccion.codigo if item.inspeccion else None
    data.hallazgo_descripcion = item.hallazgo.descripcion if item.hallazgo else None
    data.total_seguimientos = db.query(func.count(CapaSeguimientoSST.id)).filter(CapaSeguimientoSST.capa_id == item.id, CapaSeguimientoSST.activo.is_(True)).scalar() or 0
    data.total_evidencias = db.query(func.count(ArchivoSST.id)).filter(ArchivoSST.modulo == "CAPA", ArchivoSST.referencia_id == item.id, ArchivoSST.activo.is_(True)).scalar() or 0
    if item.fecha_compromiso and item.estado not in ["CERRADA", "ANULADA"]:
        diff = (item.fecha_compromiso - date.today()).days
        data.dias_vencimiento = diff
        data.vencida = diff < 0
    return data


def _agregar_traza(item: CapaSST, texto: str):
    ahora = datetime.utcnow().isoformat()
    linea = f"[{ahora}] {texto}"
    item.trazabilidad = (item.trazabilidad + "\n" if item.trazabilidad else "") + linea


def _sincronizar_cierre_automatico(db: Session, item: CapaSST, usuario_id=None):
    """Cierre automático controlado: solo cierra si avance=100, tiene evidencia y está verificada/cerrada."""
    if not item or item.estado in ["CERRADA", "ANULADA"]:
        return
    total_evidencias = db.query(func.count(ArchivoSST.id)).filter(
        ArchivoSST.modulo == "CAPA",
        ArchivoSST.referencia_id == item.id,
        ArchivoSST.activo.is_(True),
    ).scalar() or 0
    if float(item.avance or 0) >= 100:
        if total_evidencias > 0 and (item.efectiva is True or item.verificacion_eficacia):
            item.estado = "CERRADA"
            item.fecha_cierre = item.fecha_cierre or date.today()
            item.efectiva = True if item.efectiva is None else item.efectiva
            _agregar_traza(item, f"Cierre automático CAPA Enterprise por avance 100%, evidencia y verificación. Usuario {usuario_id or ''}.")
        elif item.estado not in ["VERIFICACION", "CERRADA"]:
            item.estado = "VERIFICACION"
            _agregar_traza(item, "CAPA enviada automáticamente a VERIFICACION por avance 100%.")



@router.post("/desde-hallazgo/{hallazgo_id}", response_model=CapaResponse)
def crear_capa_desde_hallazgo(
    hallazgo_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    hallazgo = db.query(InspeccionHallazgoSST).filter(InspeccionHallazgoSST.id == hallazgo_id).first()
    if hallazgo:
        _validar_empresa_usuario(usuario, hallazgo.empresa_id)
    if not hallazgo:
        raise HTTPException(status_code=404, detail="Hallazgo no encontrado")
    existe = db.query(CapaSST).filter(
        CapaSST.hallazgo_id == hallazgo.id,
        CapaSST.activo.is_(True),
    ).first()
    if existe:
        return obtener_capa(existe.id, db, usuario)

    consecutivo = (db.query(func.count(CapaSST.id)).filter(CapaSST.empresa_id == hallazgo.empresa_id).scalar() or 0) + 1
    codigo = f"CAPA-HALL-{hallazgo.id:04d}-{consecutivo:03d}"
    item = CapaSST(
        empresa_id=hallazgo.empresa_id,
        inspeccion_id=hallazgo.inspeccion_id,
        hallazgo_id=hallazgo.id,
        usuario_id=getattr(usuario, "id", None),
        codigo=codigo,
        titulo=f"CAPA por hallazgo #{hallazgo.id}",
        descripcion=hallazgo.descripcion,
        tipo_accion="CORRECTIVA",
        origen="HALLAZGO_SST",
        prioridad="CRITICA" if hallazgo.nivel_riesgo == "CRITICO" else "ALTA" if hallazgo.nivel_riesgo == "ALTO" else "MEDIA",
        estado="ABIERTA",
        responsable=hallazgo.responsable,
        fecha_apertura=date.today(),
        fecha_compromiso=hallazgo.fecha_compromiso,
        avance=0,
        causa_raiz="Pendiente análisis causa raíz.",
        accion_correctiva=hallazgo.accion_recomendada,
        observaciones="CAPA generada automáticamente desde hallazgo SST.",
        activo=True,
    )
    _agregar_traza(item, f"CAPA generada desde hallazgo {hallazgo.id} por usuario {getattr(usuario, 'id', '')}.")
    db.add(item)
    db.commit()
    db.refresh(item)
    return obtener_capa(item.id, db, usuario)


@router.get("/", response_model=list[CapaResponse])
def listar_capas(
    empresa_id: int | None = Query(default=None),
    sede_id: int | None = Query(default=None),
    area_id: int | None = Query(default=None),
    estado: str | None = Query(default=None),
    prioridad: str | None = Query(default=None),
    tipo_accion: str | None = Query(default=None),
    origen: str | None = Query(default=None),
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    items = _query_capas(db, empresa_id, sede_id, area_id, estado, prioridad, tipo_accion, origen, q, usuario=usuario).all()
    return [_capa_to_response(db, item) for item in items]


@router.get("/dashboard/resumen", response_model=CapaDashboardResponse)
def dashboard_capas(empresa_id: int | None = Query(default=None), db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    capas = _query_capas(db, empresa_id=empresa_id, usuario=usuario).all()
    total = len(capas)
    abiertas = sum(1 for c in capas if c.estado not in ["CERRADA", "ANULADA"])
    cerradas = sum(1 for c in capas if c.estado == "CERRADA")
    en_ejecucion = sum(1 for c in capas if c.estado in ["PLANIFICADA", "EN_EJECUCION", "VERIFICACION"])
    hoy = date.today()
    vencidas = sum(1 for c in capas if c.fecha_compromiso and c.fecha_compromiso < hoy and c.estado not in ["CERRADA", "ANULADA"])
    criticas = sum(1 for c in capas if c.prioridad == "CRITICA")
    cumplimiento = round((cerradas / total) * 100, 1) if total else 0
    avance_promedio = round(sum(float(c.avance or 0) for c in capas) / total, 1) if total else 0

    def conteo(attr):
        data = {}
        for item in capas:
            key = attr(item) or "Sin dato"
            data[key] = data.get(key, 0) + 1
        return [{"name": k, "value": v} for k, v in sorted(data.items(), key=lambda x: x[1], reverse=True)[:8]]

    riesgo_score = 0
    riesgo_score += 45 if vencidas else 0
    riesgo_score += 25 if criticas else 0
    riesgo_score += 20 if abiertas else 0
    riesgo_score += 10 if total and cumplimiento < 60 else 0
    semaforo = "ROJO" if riesgo_score >= 60 else "AMARILLO" if riesgo_score >= 25 else "VERDE"

    recomendaciones = []
    if vencidas:
        recomendaciones.append("Priorizar CAPA vencidas y reasignar responsables con fecha compromiso vigente.")
    if criticas:
        recomendaciones.append("Realizar revisión gerencial de acciones con prioridad crítica.")
    if abiertas:
        recomendaciones.append("Actualizar avance y seguimientos periódicos de CAPA abiertas.")
    if not recomendaciones:
        recomendaciones.append("Gestión CAPA estable. Mantener seguimiento preventivo mensual.")

    return {
        "kpis": {
            "total": total,
            "abiertas": abiertas,
            "cerradas": cerradas,
            "en_ejecucion": en_ejecucion,
            "vencidas": vencidas,
            "criticas": criticas,
            "cumplimiento": cumplimiento,
            "avance_promedio": avance_promedio,
            "riesgo_score": riesgo_score,
            "semaforo": semaforo,
        },
        "charts": {
            "por_tipo": conteo(lambda x: x.tipo_accion),
            "por_estado": conteo(lambda x: x.estado),
            "por_prioridad": conteo(lambda x: x.prioridad),
            "por_origen": conteo(lambda x: x.origen),
            "por_area": conteo(lambda x: x.area.nombre if x.area else "Sin área"),
        },
        "alertas": {"vencidas": vencidas, "criticas": criticas, "abiertas": abiertas},
        "recomendaciones": recomendaciones,
    }


@router.post("/", response_model=CapaResponse)
def crear_capa(data: CapaCreate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    _validar_empresa_usuario(usuario, data.empresa_id)
    _validar_empresa(db, data.empresa_id)
    _validar_opcional(db, Sede, data.sede_id, "Sede")
    _validar_opcional(db, Area, data.area_id, "Área")
    _validar_opcional(db, Cargo, data.cargo_id, "Cargo")
    _validar_opcional(db, Empleado, data.empleado_id, "Empleado")
    _validar_opcional(db, InspeccionSST, data.inspeccion_id, "Inspección")
    _validar_opcional(db, InspeccionHallazgoSST, data.hallazgo_id, "Hallazgo")
    existe = db.query(CapaSST).filter(CapaSST.empresa_id == data.empresa_id, func.upper(CapaSST.codigo) == data.codigo.upper()).first()
    if existe:
        raise HTTPException(status_code=400, detail="Ya existe una CAPA con ese código para la empresa")
    payload = data.model_dump()
    payload["usuario_id"] = payload.get("usuario_id") or getattr(usuario, "id", None)
    payload["fecha_apertura"] = payload.get("fecha_apertura") or date.today()
    item = CapaSST(**payload)
    _agregar_traza(item, f"CAPA creada por usuario {getattr(usuario, 'id', '')}. Estado {item.estado}.")
    db.add(item)
    db.commit()
    db.refresh(item)
    return obtener_capa(item.id, db, usuario)


@router.get("/{capa_id}", response_model=CapaResponse)
def obtener_capa(capa_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = _query_capas(db, usuario=usuario).filter(CapaSST.id == capa_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="CAPA no encontrada")
    return _capa_to_response(db, item)


@router.put("/{capa_id}", response_model=CapaResponse)
def actualizar_capa(capa_id: int, data: CapaUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = db.query(CapaSST).filter(CapaSST.id == capa_id, CapaSST.activo.is_(True)).first()
    if not item:
        raise HTTPException(status_code=404, detail="CAPA no encontrada")
    _validar_empresa_usuario(usuario, item.empresa_id)
    if item.estado == "CERRADA":
        raise HTTPException(status_code=400, detail="La CAPA está cerrada y no permite modificaciones")
    payload = data.model_dump(exclude_unset=True)
    for key, value in payload.items():
        setattr(item, key, value)
    _agregar_traza(item, f"CAPA actualizada por usuario {getattr(usuario, 'id', '')}.")
    _sincronizar_cierre_automatico(db, item, getattr(usuario, "id", None))
    db.commit()
    db.refresh(item)
    return obtener_capa(item.id, db, usuario)


@router.delete("/{capa_id}")
def eliminar_capa(capa_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    item = db.query(CapaSST).filter(CapaSST.id == capa_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="CAPA no encontrada")
    _validar_empresa_usuario(usuario, item.empresa_id)
    item.activo = False
    item.estado = "ANULADA"
    _agregar_traza(item, f"CAPA anulada por usuario {getattr(usuario, 'id', '')}.")
    db.commit()
    return {"ok": True, "message": "CAPA anulada"}


@router.post("/{capa_id}/cerrar", response_model=CapaResponse)
def cerrar_capa(capa_id: int, data: CapaCierreRequest, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = db.query(CapaSST).filter(CapaSST.id == capa_id, CapaSST.activo.is_(True)).first()
    if not item:
        raise HTTPException(status_code=404, detail="CAPA no encontrada")
    _validar_empresa_usuario(usuario, item.empresa_id)
    total_seguimientos = db.query(func.count(CapaSeguimientoSST.id)).filter(CapaSeguimientoSST.capa_id == item.id, CapaSeguimientoSST.activo.is_(True)).scalar() or 0
    total_evidencias = db.query(func.count(ArchivoSST.id)).filter(ArchivoSST.modulo == "CAPA", ArchivoSST.referencia_id == item.id, ArchivoSST.activo.is_(True)).scalar() or 0
    if float(item.avance or 0) < 100:
        raise HTTPException(status_code=400, detail="No se puede cerrar. La CAPA debe tener avance del 100%")
    if total_seguimientos == 0:
        raise HTTPException(status_code=400, detail="No se puede cerrar. Debe registrar al menos un seguimiento")
    if total_evidencias == 0:
        raise HTTPException(status_code=400, detail="No se puede cerrar. Debe adjuntar al menos una evidencia")
    item.estado = "CERRADA"
    item.fecha_cierre = date.today()
    item.efectiva = data.efectiva
    item.verificacion_eficacia = data.verificacion_eficacia
    _agregar_traza(item, f"CAPA cerrada digitalmente por usuario {getattr(usuario, 'id', '')}. {data.observacion or ''}")
    db.commit()
    db.refresh(item)
    return obtener_capa(item.id, db, usuario)


@router.get("/{capa_id}/seguimientos", response_model=list[CapaSeguimientoResponse])
def listar_seguimientos(capa_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    items = db.query(CapaSeguimientoSST).filter(CapaSeguimientoSST.capa_id == capa_id, CapaSeguimientoSST.activo.is_(True)).order_by(CapaSeguimientoSST.fecha_seguimiento.desc(), CapaSeguimientoSST.id.desc()).all()
    return items


@router.post("/{capa_id}/seguimientos", response_model=CapaSeguimientoResponse)
def crear_seguimiento(capa_id: int, data: CapaSeguimientoCreate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    capa = db.query(CapaSST).filter(CapaSST.id == capa_id, CapaSST.activo.is_(True)).first()
    if not capa:
        raise HTTPException(status_code=404, detail="CAPA no encontrada")
    _validar_empresa_usuario(usuario, capa.empresa_id)
    if capa.estado == "CERRADA":
        raise HTTPException(status_code=400, detail="La CAPA está cerrada y no permite seguimientos")
    payload = data.model_dump()
    payload["capa_id"] = capa_id
    payload["empresa_id"] = capa.empresa_id
    payload["usuario_id"] = getattr(usuario, "id", None)
    item = CapaSeguimientoSST(**payload)
    capa.avance = max(float(capa.avance or 0), float(item.avance or 0))
    if float(capa.avance or 0) >= 100:
        capa.estado = "VERIFICACION"
        if str(item.resultado or "").upper() in {"VERIFICADO", "CERRADO", "COMPLETADO", "FINALIZADO", "EFECTIVO"}:
            capa.verificacion_eficacia = capa.verificacion_eficacia or item.comentario
            capa.efectiva = True
    elif capa.estado == "ABIERTA":
        capa.estado = "EN_EJECUCION"
    _agregar_traza(capa, f"Seguimiento registrado por {item.responsable or 'usuario ' + str(getattr(usuario, 'id', ''))}. Avance {item.avance}%.")
    _sincronizar_cierre_automatico(db, capa, getattr(usuario, "id", None))
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


@router.put("/seguimientos/{seguimiento_id}", response_model=CapaSeguimientoResponse)
def actualizar_seguimiento(seguimiento_id: int, data: CapaSeguimientoUpdate, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    item = db.query(CapaSeguimientoSST).filter(CapaSeguimientoSST.id == seguimiento_id, CapaSeguimientoSST.activo.is_(True)).first()
    if not item:
        raise HTTPException(status_code=404, detail="Seguimiento CAPA no encontrado")
    capa = db.query(CapaSST).filter(CapaSST.id == item.capa_id).first()
    if capa:
        _validar_empresa_usuario(usuario, capa.empresa_id)
    if capa and capa.estado == "CERRADA":
        raise HTTPException(status_code=400, detail="La CAPA está cerrada y no permite modificar seguimientos")
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(item, key, value)
    if capa and data.avance is not None:
        capa.avance = max(float(capa.avance or 0), float(data.avance or 0))
        _agregar_traza(capa, f"Seguimiento actualizado por usuario {getattr(usuario, 'id', '')}.")
    db.commit()
    db.refresh(item)
    return item


@router.delete("/seguimientos/{seguimiento_id}")
def eliminar_seguimiento(seguimiento_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    item = db.query(CapaSeguimientoSST).filter(CapaSeguimientoSST.id == seguimiento_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Seguimiento CAPA no encontrado")
    item.activo = False
    db.commit()
    return {"ok": True, "message": "Seguimiento desactivado"}


@router.get("/{capa_id}/evidencias")
def listar_evidencias(capa_id: int, db: Session = Depends(get_db), usuario=Depends(require_roles(ROLES_SST))):
    archivos = db.query(ArchivoSST).filter(ArchivoSST.modulo == "CAPA", ArchivoSST.referencia_id == capa_id, ArchivoSST.activo.is_(True)).order_by(ArchivoSST.fecha_creacion.desc()).all()
    return [_archivo_to_dict(a) for a in archivos]


@router.post("/{capa_id}/evidencias")
def subir_evidencia(
    capa_id: int,
    tipo_evidencia: str = Form(default="EVIDENCIA_CAPA"),
    descripcion: str = Form(default=""),
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db),
    usuario=Depends(require_roles(ROLES_SST)),
):
    capa = db.query(CapaSST).filter(CapaSST.id == capa_id, CapaSST.activo.is_(True)).first()
    if not capa:
        raise HTTPException(status_code=404, detail="CAPA no encontrada")
    _validar_empresa_usuario(usuario, capa.empresa_id)
    if capa.estado == "CERRADA":
        raise HTTPException(status_code=400, detail="La CAPA está cerrada y no permite subir evidencias")
    path, original, filename, mime_type, size = _guardar_upload(archivo)
    registro = ArchivoSST(
        empresa_id=capa.empresa_id,
        usuario_id=getattr(usuario, "id", None),
        tipo=(tipo_evidencia or "EVIDENCIA_CAPA").upper().strip(),
        nombre_original=original,
        nombre_archivo=filename,
        ruta=str(path),
        url=_public_upload_url(path),
        extension=filename.rsplit(".", 1)[-1].lower(),
        mime_type=mime_type,
        tamano_bytes=size,
        modulo="CAPA",
        referencia_id=capa.id,
        descripcion=descripcion or "Evidencia CAPA",
        activo=True,
    )
    _agregar_traza(capa, f"Evidencia CAPA adjuntada: {original}.")
    db.add(registro)
    db.flush()
    _sincronizar_cierre_automatico(db, capa, getattr(usuario, "id", None))
    db.commit()
    db.refresh(registro)
    return _archivo_to_dict(registro)


@router.delete("/{capa_id}/evidencias/{archivo_id}")
def eliminar_evidencia(capa_id: int, archivo_id: int, db: Session = Depends(get_db), usuario=Depends(ELIMINAR_REGISTROS)):
    archivo = db.query(ArchivoSST).filter(ArchivoSST.id == archivo_id, ArchivoSST.modulo == "CAPA", ArchivoSST.referencia_id == capa_id).first()
    if not archivo:
        raise HTTPException(status_code=404, detail="Evidencia CAPA no encontrada")
    capa = db.query(CapaSST).filter(CapaSST.id == capa_id).first()
    if capa:
        _validar_empresa_usuario(usuario, capa.empresa_id)
        if capa.estado == "CERRADA":
            raise HTTPException(status_code=400, detail="La CAPA está cerrada y no permite eliminar evidencias")
    archivo.activo = False
    db.commit()
    return {"ok": True, "message": "Evidencia CAPA desactivada"}


def _excel_response(workbook, filename: str):
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _pdf_response(buffer: io.BytesIO, filename: str):
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.get("/exportaciones/excel-general")
def exportar_excel_general(empresa_id: int | None = Query(default=None), db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "CAPA SST"
    headers = ["Código", "Título", "Empresa", "Tipo", "Origen", "Prioridad", "Estado", "Responsable", "Compromiso", "Avance", "Efectiva"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for c in _query_capas(db, empresa_id=empresa_id, usuario=usuario).all():
        ws.append([c.codigo, c.titulo, c.empresa.nombre if c.empresa else "", c.tipo_accion, c.origen, c.prioridad, c.estado, c.responsable or "", str(c.fecha_compromiso or ""), float(c.avance or 0), "SI" if c.efectiva else "NO" if c.efectiva is False else ""])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in col) + 2, 45)
    return _excel_response(wb, "capas_sst_general.xlsx")


@router.get("/exportaciones/pdf-general")
def exportar_pdf_general(empresa_id: int | None = Query(default=None), db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=24, leftMargin=24, topMargin=28, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = [Paragraph("Reporte General CAPA SST", styles["Title"]), Spacer(1, 12)]
    rows = [["Código", "Título", "Tipo", "Prioridad", "Estado", "Responsable", "Avance"]]
    for c in _query_capas(db, empresa_id=empresa_id, usuario=usuario).all():
        rows.append([c.codigo, c.titulo[:45], c.tipo_accion, c.prioridad, c.estado, c.responsable or "", f"{float(c.avance or 0)}%"])
    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0f766e")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .3, colors.HexColor("#cbd5e1")), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,-1), 8)]))
    story.append(table)
    doc.build(story)
    return _pdf_response(buffer, "capas_sst_general.pdf")


@router.get("/exportaciones/dashboard-pdf")
def exportar_dashboard_pdf(empresa_id: int | None = Query(default=None), db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    resumen = dashboard_capas(empresa_id, db, usuario)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48)
    styles = getSampleStyleSheet()
    story = [Paragraph("Dashboard Ejecutivo CAPA SST", styles["Title"]), Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]), Spacer(1, 14)]
    rows = [["Indicador", "Valor"]] + [[k.replace("_", " ").title(), str(v)] for k, v in resumen["kpis"].items()]
    table = Table(rows, colWidths=[280, 180])
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0f766e")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#cbd5e1")), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold")]))
    story.extend([table, Spacer(1, 18), Paragraph("Recomendaciones PRO", styles["Heading2"])])
    for rec in resumen["recomendaciones"]:
        story.append(Paragraph(f"• {rec}", styles["Normal"]))
    doc.build(story)
    return _pdf_response(buffer, "dashboard_capa_sst.pdf")


@router.get("/exportaciones/{capa_id}/pdf-individual")
def exportar_pdf_individual_capa(capa_id: int, db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    item = db.query(CapaSST).options(joinedload(CapaSST.empresa), joinedload(CapaSST.hallazgo), joinedload(CapaSST.inspeccion)).filter(CapaSST.id == capa_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="CAPA no encontrada")
    evidencias = db.query(ArchivoSST).filter(ArchivoSST.modulo == "CAPA", ArchivoSST.referencia_id == item.id, ArchivoSST.activo.is_(True)).all()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    story = [Paragraph(f"CAPA SST {item.codigo}", styles["Title"]), Paragraph(item.titulo or "", styles["Normal"]), Spacer(1, 12)]
    rows = [
        ["Empresa", item.empresa.nombre if item.empresa else ""], ["Origen", item.origen], ["Tipo", item.tipo_accion],
        ["Prioridad", item.prioridad], ["Estado", item.estado], ["Responsable", item.responsable or ""],
        ["Compromiso", str(item.fecha_compromiso or "")], ["Cierre", str(item.fecha_cierre or "")],
        ["Avance", f"{float(item.avance or 0)}%"], ["Efectiva", "SI" if item.efectiva else "NO" if item.efectiva is False else "PENDIENTE"],
    ]
    table = Table([["Campo", "Valor"]] + rows, colWidths=[150, 360])
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0f766e")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#cbd5e1")), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold")]))
    story.extend([table, Spacer(1, 14), Paragraph("Descripción", styles["Heading2"]), Paragraph(item.descripcion or "", styles["Normal"]), Spacer(1, 10), Paragraph("Causa raíz", styles["Heading2"]), Paragraph(item.causa_raiz or "Sin registrar", styles["Normal"]), Spacer(1, 10), Paragraph("Acciones", styles["Heading2"]), Paragraph(f"<b>Inmediata:</b> {item.accion_inmediata or 'Sin registrar'}", styles["Normal"]), Paragraph(f"<b>Correctiva:</b> {item.accion_correctiva or 'Sin registrar'}", styles["Normal"]), Paragraph(f"<b>Preventiva:</b> {item.accion_preventiva or 'Sin registrar'}", styles["Normal"]), Spacer(1, 10), Paragraph("Evidencias", styles["Heading2"])])
    if evidencias:
        for ev in evidencias:
            story.append(Paragraph(f"• {ev.nombre_original} — {ev.descripcion or ''}", styles["Normal"]))
    else:
        story.append(Paragraph("Sin evidencias registradas.", styles["Normal"]))
    doc.build(story)
    return _pdf_response(buffer, f"capa_sst_{item.codigo}.pdf")


@router.get("/exportaciones/{capa_id}/acta-pdf")
def exportar_acta_capa(capa_id: int, db: Session = Depends(get_db), usuario=Depends(EXPORTAR_REPORTES)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    item = db.query(CapaSST).options(joinedload(CapaSST.empresa)).filter(CapaSST.id == capa_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="CAPA no encontrada")
    seguimientos = db.query(CapaSeguimientoSST).filter(CapaSeguimientoSST.capa_id == capa_id, CapaSeguimientoSST.activo.is_(True)).all()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    story = [Paragraph("ACTA OFICIAL CAPA SST", styles["Title"]), Paragraph(f"Código: {item.codigo} — {item.titulo}", styles["Normal"]), Spacer(1, 12)]
    rows = [["Empresa", item.empresa.nombre if item.empresa else "", "Estado", item.estado], ["Tipo", item.tipo_accion, "Prioridad", item.prioridad], ["Responsable", item.responsable or "", "Avance", f"{float(item.avance or 0)}%"], ["Compromiso", str(item.fecha_compromiso or ""), "Cierre", str(item.fecha_cierre or "")]]
    table = Table(rows, colWidths=[90, 190, 90, 130])
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0f766e")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#cbd5e1"))]))
    story.append(table)
    story.extend([Spacer(1, 14), Paragraph("Descripción", styles["Heading2"]), Paragraph(item.descripcion or "", styles["Normal"]), Spacer(1, 12), Paragraph("Causa raíz", styles["Heading2"]), Paragraph(item.causa_raiz or "Sin registrar", styles["Normal"]), Spacer(1, 12), Paragraph("Seguimientos", styles["Heading2"])])
    rows = [["Fecha", "Responsable", "Avance", "Comentario"]]
    for s in seguimientos:
        rows.append([str(s.fecha_seguimiento), s.responsable or "", f"{float(s.avance or 0)}%", s.comentario[:80]])
    st = Table(rows, repeatRows=1)
    st.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0f766e")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .3, colors.HexColor("#cbd5e1")), ("FONTSIZE", (0,0), (-1,-1), 8)]))
    story.append(st)
    story.extend([Spacer(1, 12), Paragraph("Trazabilidad", styles["Heading2"]), Paragraph((item.trazabilidad or "Sin trazabilidad").replace("\n", "<br/>"), styles["Normal"])])
    doc.build(story)
    return _pdf_response(buffer, f"acta_capa_sst_{item.codigo}.pdf")
