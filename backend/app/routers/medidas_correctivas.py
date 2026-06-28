from __future__ import annotations
from datetime import date, datetime
from pathlib import Path
import io, os, uuid
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, joinedload
from app.auth.dependencies import require_roles
from app.database import get_db
from app.models.archivo_sst import ArchivoSST
from app.models.area import Area
from app.models.capa import CapaSST, CapaSeguimientoSST
from app.models.cargo import Cargo
from app.models.empleado import Empleado
from app.models.empresa import Empresa
from app.models.inspeccion import InspeccionHallazgoSST, InspeccionSST
from app.models.sede import Sede
from app.schemas.medidas_correctivas_schema import *

router=APIRouter(prefix='/medidas-correctivas', tags=['Medidas Correctivas SST Enterprise'])
ROLES_SST=['SUPER_ADMIN','ADMIN_EMPRESA','RESPONSABLE_SST']
UPLOAD_ROOT=Path(os.getenv('UPLOAD_DIR','app/uploads')).resolve(); MEDIDAS_UPLOAD_DIR=UPLOAD_ROOT/'medidas-correctivas'; MEDIDAS_UPLOAD_DIR.mkdir(parents=True,exist_ok=True)
ALLOWED_EXT={'pdf','jpg','jpeg','png','webp','xlsx','xls','csv','docx','doc'}; MAX_UPLOAD_MB=25

def _validar_empresa_usuario(usuario, empresa_id):
    if not empresa_id or getattr(usuario,'rol',None)=='SUPER_ADMIN': return
    if getattr(usuario,'empresa_id',None) and int(getattr(usuario,'empresa_id'))==int(empresa_id): return
    raise HTTPException(status_code=403, detail='No tiene permisos sobre esta empresa')
def _upper(v,default=None):
    v=str(v or '').strip().upper(); return v or default
def _public_upload_url(p:Path):
    try: return '/uploads/'+p.resolve().relative_to(UPLOAD_ROOT).as_posix()
    except Exception: return '/uploads/medidas-correctivas/'+p.name
def _validar_empresa(db,empresa_id):
    e=db.query(Empresa).filter(Empresa.id==empresa_id).first()
    if not e: raise HTTPException(status_code=404, detail='Empresa no encontrada')
    return e
def _validar_opcional(db, model, item_id, label):
    if not item_id: return None
    item=db.query(model).filter(model.id==item_id).first()
    if not item: raise HTTPException(status_code=404, detail=f'{label} no encontrado')
    return item

def _image_to_webp_bytes(image,max_size,quality):
    img=image.copy();
    if img.mode not in ('RGB','RGBA'): img=img.convert('RGB')
    img.thumbnail(max_size); out=io.BytesIO(); img.save(out,format='WEBP',quality=quality,method=6,optimize=True); return out.getvalue()
def _guardar_upload(upload:UploadFile):
    original=upload.filename or 'evidencia_medida_correctiva'
    if '..' in original or '/' in original or '\\' in original: raise HTTPException(status_code=400, detail='Nombre de archivo no permitido')
    ext=original.rsplit('.',1)[-1].lower() if '.' in original else 'bin'
    if ext not in ALLOWED_EXT: raise HTTPException(status_code=400, detail='Archivo no permitido')
    content=upload.file.read()
    if len(content)>MAX_UPLOAD_MB*1024*1024: raise HTTPException(status_code=400, detail=f'El archivo supera {MAX_UPLOAD_MB} MB')
    uid=uuid.uuid4().hex
    if ext in {'jpg','jpeg','png','webp'}:
        from PIL import Image, ImageOps
        image=ImageOps.exif_transpose(Image.open(io.BytesIO(content)))
        main=_image_to_webp_bytes(image,(1920,1080),82); prev=_image_to_webp_bytes(image,(1280,900),78); thumb=_image_to_webp_bytes(image,(260,260),72)
        filename=f'{uid}.webp'; path=MEDIDAS_UPLOAD_DIR/filename; path.write_bytes(main); (MEDIDAS_UPLOAD_DIR/f'{uid}_preview.webp').write_bytes(prev); (MEDIDAS_UPLOAD_DIR/f'{uid}_thumb.webp').write_bytes(thumb)
        return path, original, filename, 'image/webp', len(main)
    filename=f'{uid}.{ext}'; path=MEDIDAS_UPLOAD_DIR/filename; path.write_bytes(content); return path, original, filename, upload.content_type or 'application/octet-stream', len(content)

def _archivo_to_dict(a):
    stem=Path(a.nombre_archivo or '').stem; base=Path(a.ruta or '').parent if a.ruta else MEDIDAS_UPLOAD_DIR
    prev=base/f'{stem}_preview.webp'; thumb=base/f'{stem}_thumb.webp'
    return {'id':a.id,'empresa_id':a.empresa_id,'usuario_id':a.usuario_id,'tipo':a.tipo,'nombre_original':a.nombre_original,'nombre_archivo':a.nombre_archivo,'ruta':a.ruta,'url':a.url,'preview_url':_public_upload_url(prev) if prev.exists() else a.url,'thumbnail_url':_public_upload_url(thumb) if thumb.exists() else None,'extension':a.extension,'mime_type':a.mime_type,'tamano_bytes':a.tamano_bytes,'modulo':a.modulo,'referencia_id':a.referencia_id,'descripcion':a.descripcion,'activo':a.activo,'fecha_creacion':a.fecha_creacion}

def _query_medidas(db, empresa_id=None, sede_id=None, area_id=None, estado=None, prioridad=None, tipo_accion=None, origen=None, q=None, usuario=None):
    query=db.query(CapaSST).options(joinedload(CapaSST.empresa),joinedload(CapaSST.sede),joinedload(CapaSST.area),joinedload(CapaSST.cargo),joinedload(CapaSST.empleado),joinedload(CapaSST.inspeccion),joinedload(CapaSST.hallazgo)).filter(CapaSST.activo.is_(True))
    if usuario and getattr(usuario,'rol',None)!='SUPER_ADMIN' and getattr(usuario,'empresa_id',None): query=query.filter(CapaSST.empresa_id==getattr(usuario,'empresa_id'))
    if empresa_id: query=query.filter(CapaSST.empresa_id==empresa_id)
    if sede_id: query=query.filter(CapaSST.sede_id==sede_id)
    if area_id: query=query.filter(CapaSST.area_id==area_id)
    if estado: query=query.filter(func.upper(CapaSST.estado)==estado.upper().strip())
    if prioridad: query=query.filter(func.upper(CapaSST.prioridad)==prioridad.upper().strip())
    if tipo_accion: query=query.filter(func.upper(CapaSST.tipo_accion)==tipo_accion.upper().strip())
    if origen: query=query.filter(func.upper(CapaSST.origen)==origen.upper().strip())
    if q:
        like=f'%{q.strip()}%'; query=query.filter(or_(CapaSST.codigo.ilike(like),CapaSST.titulo.ilike(like),CapaSST.descripcion.ilike(like),CapaSST.responsable.ilike(like)))
    return query.order_by(CapaSST.id.desc())

def _medida_to_response(db,item):
    data=MedidaCorrectivaResponse.model_validate(item)
    data.empresa_nombre=item.empresa.nombre if item.empresa else None; data.sede_nombre=item.sede.nombre if item.sede else None; data.area_nombre=item.area.nombre if item.area else None; data.cargo_nombre=item.cargo.nombre if item.cargo else None
    data.empleado_nombre=f'{item.empleado.nombres} {item.empleado.apellidos}'.strip() if item.empleado else None
    data.inspeccion_codigo=item.inspeccion.codigo if item.inspeccion else None; data.hallazgo_descripcion=item.hallazgo.descripcion if item.hallazgo else None
    data.total_seguimientos=db.query(func.count(CapaSeguimientoSST.id)).filter(CapaSeguimientoSST.capa_id==item.id,CapaSeguimientoSST.activo.is_(True)).scalar() or 0
    data.total_evidencias=db.query(func.count(ArchivoSST.id)).filter(ArchivoSST.modulo.in_(['CAPA','MEDIDAS_CORRECTIVAS']),ArchivoSST.referencia_id==item.id,ArchivoSST.activo.is_(True)).scalar() or 0
    if item.fecha_compromiso and item.estado not in ['CERRADA','ANULADA']:
        diff=(item.fecha_compromiso-date.today()).days; data.dias_vencimiento=diff; data.vencida=diff<0
    return data

def _traza(item,texto): item.trazabilidad=(item.trazabilidad+'\n' if item.trazabilidad else '')+f'[{datetime.utcnow().isoformat()}] {texto}'

def _validar_cierre(db,item):
    seg=db.query(func.count(CapaSeguimientoSST.id)).filter(CapaSeguimientoSST.capa_id==item.id,CapaSeguimientoSST.activo.is_(True)).scalar() or 0
    evi=db.query(func.count(ArchivoSST.id)).filter(ArchivoSST.modulo.in_(['CAPA','MEDIDAS_CORRECTIVAS']),ArchivoSST.referencia_id==item.id,ArchivoSST.activo.is_(True)).scalar() or 0
    if float(item.avance or 0)<100: raise HTTPException(status_code=400, detail='No se puede cerrar. La medida debe tener avance del 100%')
    if seg==0: raise HTTPException(status_code=400, detail='No se puede cerrar. Debe registrar al menos un seguimiento')
    if evi==0: raise HTTPException(status_code=400, detail='No se puede cerrar. Debe adjuntar al menos una evidencia')

@router.get('/', response_model=list[MedidaCorrectivaResponse])
def listar_medidas_correctivas(empresa_id:int|None=Query(None),sede_id:int|None=Query(None),area_id:int|None=Query(None),estado:str|None=Query(None),prioridad:str|None=Query(None),tipo_accion:str|None=Query(None),origen:str|None=Query(None),q:str|None=Query(None),db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    _validar_empresa_usuario(usuario,empresa_id); return [_medida_to_response(db,i) for i in _query_medidas(db,empresa_id,sede_id,area_id,estado,prioridad,tipo_accion,origen,q,usuario).all()]

@router.get('/dashboard/resumen', response_model=MedidaCorrectivaDashboardResponse)
def dashboard_medidas_correctivas(empresa_id:int|None=Query(None),db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    _validar_empresa_usuario(usuario,empresa_id); medidas=_query_medidas(db,empresa_id=empresa_id,usuario=usuario).all(); total=len(medidas); hoy=date.today()
    abiertas=sum(1 for c in medidas if c.estado not in ['CERRADA','ANULADA']); cerradas=sum(1 for c in medidas if c.estado=='CERRADA'); en_ejecucion=sum(1 for c in medidas if c.estado in ['PLANIFICADA','EN_EJECUCION','VERIFICACION']); vencidas=sum(1 for c in medidas if c.fecha_compromiso and c.fecha_compromiso<hoy and c.estado not in ['CERRADA','ANULADA']); criticas=sum(1 for c in medidas if c.prioridad in ['CRITICA','CRÍTICA']); cumplimiento=round((cerradas/total)*100,1) if total else 0; avance_promedio=round(sum(float(c.avance or 0) for c in medidas)/total,1) if total else 0; costo_estimado=round(sum(float(getattr(c,'costo_estimado',0) or 0) for c in medidas),2); costo_real=round(sum(float(getattr(c,'costo_real',0) or 0) for c in medidas),2)
    def conteo(fn):
        d={}
        for item in medidas:
            k=fn(item) or 'Sin dato'; d[k]=d.get(k,0)+1
        return [{'name':k,'value':v} for k,v in sorted(d.items(), key=lambda x:x[1], reverse=True)[:8]]
    score=(45 if vencidas else 0)+(25 if criticas else 0)+(20 if abiertas else 0)+(10 if total and cumplimiento<60 else 0); semaforo='ROJO' if score>=60 else 'AMARILLO' if score>=25 else 'VERDE'
    rec=[]
    if vencidas: rec.append('Priorizar medidas vencidas y reasignar fechas compromiso realistas.')
    if criticas: rec.append('Escalar medidas críticas a revisión del responsable SST y gerencia.')
    if abiertas: rec.append('Registrar seguimientos y evidencias para las medidas abiertas.')
    if total and avance_promedio<60: rec.append('Revisar plan de trabajo: el avance promedio está por debajo del nivel esperado.')
    if not rec: rec.append('Centro de medidas correctivas estable. Mantener seguimiento preventivo.')
    return {'kpis':{'total':total,'abiertas':abiertas,'cerradas':cerradas,'en_ejecucion':en_ejecucion,'vencidas':vencidas,'criticas':criticas,'cumplimiento':cumplimiento,'avance_promedio':avance_promedio,'costo_estimado':costo_estimado,'costo_real':costo_real,'riesgo_score':score,'semaforo':semaforo},'charts':{'por_tipo':conteo(lambda x:x.tipo_accion),'por_estado':conteo(lambda x:x.estado),'por_prioridad':conteo(lambda x:x.prioridad),'por_origen':conteo(lambda x:x.origen),'por_area':conteo(lambda x:x.area.nombre if x.area else 'Sin área')},'alertas':{'vencidas':vencidas,'criticas':criticas,'abiertas':abiertas},'recomendaciones':rec}

@router.post('/', response_model=MedidaCorrectivaResponse)
def crear_medida_correctiva(data:MedidaCorrectivaCreate,db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    _validar_empresa_usuario(usuario,data.empresa_id); _validar_empresa(db,data.empresa_id); _validar_opcional(db,Sede,data.sede_id,'Sede'); _validar_opcional(db,Area,data.area_id,'Área'); _validar_opcional(db,Cargo,data.cargo_id,'Cargo'); _validar_opcional(db,Empleado,data.empleado_id,'Empleado'); _validar_opcional(db,InspeccionSST,data.inspeccion_id,'Inspección'); _validar_opcional(db,InspeccionHallazgoSST,data.hallazgo_id,'Hallazgo')
    if db.query(CapaSST).filter(CapaSST.empresa_id==data.empresa_id,func.upper(CapaSST.codigo)==data.codigo.upper()).first(): raise HTTPException(status_code=400, detail='Ya existe una medida con ese código para la empresa')
    payload=data.model_dump(); payload['usuario_id']=getattr(usuario,'id',None); payload['fecha_apertura']=payload.get('fecha_apertura') or date.today(); payload['tipo_accion']=_upper(payload.get('tipo_accion'),'CORRECTIVA'); payload['origen']=_upper(payload.get('origen'),'MANUAL'); payload['prioridad']=_upper(payload.get('prioridad'),'MEDIA'); payload['estado']=_upper(payload.get('estado'),'ABIERTA')
    item=CapaSST(**payload); _traza(item,f'Medida correctiva creada desde Centro Enterprise por usuario {getattr(usuario,"id","")}.'); db.add(item); db.commit(); db.refresh(item); return obtener_medida_correctiva(item.id,db,usuario)

@router.get('/{medida_id}', response_model=MedidaCorrectivaResponse)
def obtener_medida_correctiva(medida_id:int,db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    item=_query_medidas(db,usuario=usuario).filter(CapaSST.id==medida_id).first()
    if not item: raise HTTPException(status_code=404, detail='Medida correctiva no encontrada')
    return _medida_to_response(db,item)

@router.put('/{medida_id}', response_model=MedidaCorrectivaResponse)
def actualizar_medida_correctiva(medida_id:int,data:MedidaCorrectivaUpdate,db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    item=db.query(CapaSST).filter(CapaSST.id==medida_id,CapaSST.activo.is_(True)).first()
    if not item: raise HTTPException(status_code=404, detail='Medida correctiva no encontrada')
    _validar_empresa_usuario(usuario,item.empresa_id)
    if item.estado=='CERRADA': raise HTTPException(status_code=400, detail='La medida está cerrada y no permite modificaciones')
    for k,v in data.model_dump(exclude_unset=True).items():
        if k in {'tipo_accion','origen','prioridad','estado'} and v: v=_upper(v)
        setattr(item,k,v)
    _traza(item,f'Medida correctiva actualizada por usuario {getattr(usuario,"id","")}.'); db.commit(); db.refresh(item); return obtener_medida_correctiva(item.id,db,usuario)

@router.delete('/{medida_id}')
def eliminar_medida_correctiva(medida_id:int,db:Session=Depends(get_db),usuario=Depends(require_roles(['SUPER_ADMIN','ADMIN_EMPRESA']))):
    item=db.query(CapaSST).filter(CapaSST.id==medida_id).first()
    if not item: raise HTTPException(status_code=404, detail='Medida correctiva no encontrada')
    _validar_empresa_usuario(usuario,item.empresa_id); item.activo=False; item.estado='ANULADA'; _traza(item,f'Medida correctiva anulada por usuario {getattr(usuario,"id","")}.'); db.commit(); return {'ok':True,'message':'Medida correctiva anulada'}

@router.post('/{medida_id}/aprobar', response_model=MedidaCorrectivaResponse)
def aprobar_medida_correctiva(medida_id:int,data:MedidaCorrectivaAprobacionRequest,db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    item=db.query(CapaSST).filter(CapaSST.id==medida_id,CapaSST.activo.is_(True)).first()
    if not item: raise HTTPException(status_code=404, detail='Medida correctiva no encontrada')
    _validar_empresa_usuario(usuario,item.empresa_id); item.aprobada_por=getattr(usuario,'id',None); item.fecha_aprobacion=datetime.utcnow(); item.estado='PLANIFICADA' if item.estado=='ABIERTA' else item.estado; _traza(item,f'Medida aprobada por usuario {getattr(usuario,"id","")}. {data.observacion or ""}'); db.commit(); return obtener_medida_correctiva(item.id,db,usuario)

@router.post('/{medida_id}/cerrar', response_model=MedidaCorrectivaResponse)
def cerrar_medida_correctiva(medida_id:int,data:MedidaCorrectivaCierreRequest,db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    item=db.query(CapaSST).filter(CapaSST.id==medida_id,CapaSST.activo.is_(True)).first()
    if not item: raise HTTPException(status_code=404, detail='Medida correctiva no encontrada')
    _validar_empresa_usuario(usuario,item.empresa_id); _validar_cierre(db,item); item.estado='CERRADA'; item.fecha_cierre=date.today(); item.efectiva=data.efectiva; item.verificacion_eficacia=data.verificacion_eficacia; _traza(item,f'Medida cerrada digitalmente por usuario {getattr(usuario,"id","")}. {data.observacion or ""}'); db.commit(); return obtener_medida_correctiva(item.id,db,usuario)

@router.get('/{medida_id}/seguimientos', response_model=list[MedidaCorrectivaSeguimientoResponse])
def listar_seguimientos_medida(medida_id:int,db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    medida=db.query(CapaSST).filter(CapaSST.id==medida_id,CapaSST.activo.is_(True)).first()
    if not medida: raise HTTPException(status_code=404, detail='Medida correctiva no encontrada')
    _validar_empresa_usuario(usuario,medida.empresa_id); return db.query(CapaSeguimientoSST).filter(CapaSeguimientoSST.capa_id==medida_id,CapaSeguimientoSST.activo.is_(True)).order_by(CapaSeguimientoSST.fecha_seguimiento.desc(),CapaSeguimientoSST.id.desc()).all()

@router.post('/{medida_id}/seguimientos', response_model=MedidaCorrectivaSeguimientoResponse)
def crear_seguimiento_medida(medida_id:int,data:MedidaCorrectivaSeguimientoCreate,db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    medida=db.query(CapaSST).filter(CapaSST.id==medida_id,CapaSST.activo.is_(True)).first()
    if not medida: raise HTTPException(status_code=404, detail='Medida correctiva no encontrada')
    _validar_empresa_usuario(usuario,medida.empresa_id)
    if medida.estado=='CERRADA': raise HTTPException(status_code=400, detail='La medida está cerrada y no permite seguimientos')
    payload=data.model_dump(); payload['capa_id']=medida_id; payload['empresa_id']=medida.empresa_id; payload['usuario_id']=getattr(usuario,'id',None); payload['fecha_seguimiento']=payload.get('fecha_seguimiento') or date.today(); item=CapaSeguimientoSST(**payload); medida.avance=max(float(medida.avance or 0),float(item.avance or 0)); medida.estado='VERIFICACION' if float(medida.avance or 0)>=100 else ('EN_EJECUCION' if medida.estado in ['ABIERTA','PLANIFICADA'] else medida.estado); _traza(medida,f'Seguimiento registrado. Avance {item.avance}%.'); db.add(item); db.commit(); db.refresh(item); return item

@router.get('/{medida_id}/evidencias')
def listar_evidencias_medida(medida_id:int,db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    medida=db.query(CapaSST).filter(CapaSST.id==medida_id,CapaSST.activo.is_(True)).first()
    if not medida: raise HTTPException(status_code=404, detail='Medida correctiva no encontrada')
    _validar_empresa_usuario(usuario,medida.empresa_id); archivos=db.query(ArchivoSST).filter(ArchivoSST.modulo.in_(['CAPA','MEDIDAS_CORRECTIVAS']),ArchivoSST.referencia_id==medida_id,ArchivoSST.activo.is_(True)).order_by(ArchivoSST.fecha_creacion.desc()).all(); return [_archivo_to_dict(a) for a in archivos]

@router.post('/{medida_id}/evidencias')
def subir_evidencia_medida(medida_id:int,tipo_evidencia:str=Form(default='EVIDENCIA_MEDIDA'),descripcion:str=Form(default=''),archivo:UploadFile=File(...),db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    medida=db.query(CapaSST).filter(CapaSST.id==medida_id,CapaSST.activo.is_(True)).first()
    if not medida: raise HTTPException(status_code=404, detail='Medida correctiva no encontrada')
    _validar_empresa_usuario(usuario,medida.empresa_id)
    if medida.estado=='CERRADA': raise HTTPException(status_code=400, detail='La medida está cerrada y no permite subir evidencias')
    path,original,filename,mime_type,size=_guardar_upload(archivo); reg=ArchivoSST(empresa_id=medida.empresa_id,usuario_id=getattr(usuario,'id',None),tipo=(tipo_evidencia or 'EVIDENCIA_MEDIDA').upper().strip(),nombre_original=original,nombre_archivo=filename,ruta=str(path),url=_public_upload_url(path),extension=filename.rsplit('.',1)[-1].lower(),mime_type=mime_type,tamano_bytes=size,modulo='MEDIDAS_CORRECTIVAS',referencia_id=medida.id,descripcion=descripcion or 'Evidencia de medida correctiva',activo=True); _traza(medida,f'Evidencia adjuntada: {original}.'); db.add(reg); db.commit(); db.refresh(reg); return _archivo_to_dict(reg)

@router.get('/exportaciones/excel-general')
def exportar_excel_medidas(empresa_id:int|None=Query(None),db:Session=Depends(get_db),usuario=Depends(require_roles(ROLES_SST))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    _validar_empresa_usuario(usuario,empresa_id); wb=Workbook(); ws=wb.active; ws.title='Medidas Correctivas'; headers=['Código','Título','Empresa','Tipo','Origen','Prioridad','Estado','Responsable','Compromiso','Avance','Costo Estimado','Costo Real','Efectiva']; ws.append(headers)
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='173A8A')
    for c in _query_medidas(db,empresa_id=empresa_id,usuario=usuario).all(): ws.append([c.codigo,c.titulo,c.empresa.nombre if c.empresa else '',c.tipo_accion,c.origen,c.prioridad,c.estado,c.responsable or '',str(c.fecha_compromiso or ''),float(c.avance or 0),float(getattr(c,'costo_estimado',0) or 0),float(getattr(c,'costo_real',0) or 0),'SI' if c.efectiva else 'NO' if c.efectiva is False else ''])
    output=io.BytesIO(); wb.save(output); output.seek(0); return StreamingResponse(output,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename=medidas_correctivas_sst.xlsx'})
