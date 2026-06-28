# ============================================================
# SERVICE EXCEL EJECUTIVO MEDIDAS CORRECTIVAS
# ERP SST PRO
# FASE 1.1.8.7.6.1 — Exportaciones Enterprise PDF + Excel
# ============================================================

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.archivo_sst import ArchivoSST
from app.models.capa import CapaSST, CapaSeguimientoSST


AZUL = "1D4ED8"
AZUL_OSCURO = "0F2D6B"
GRIS = "64748B"
BLANCO = "FFFFFF"
BORDE = "CBD5E1"


def _safe(value: Any, default: str = "—") -> str:
    if value is None:
        return default
    text = str(value).strip()
    return text if text else default


def _fmt_date(value: Any) -> str:
    if not value:
        return "—"
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    return str(value)


def _num(value: Any) -> float:
    try:
        return float(value or 0)
    except Exception:
        return 0.0


def _eficacia_estado(item: CapaSST) -> str:
    porcentaje = getattr(item, "porcentaje_eficacia", None)
    pct = _num(porcentaje) if porcentaje is not None else None

    if item.efectiva is True or (pct is not None and pct >= 80):
        return "EFICAZ"
    if pct is not None and pct >= 50:
        return "PARCIAL"
    if pct is not None and pct < 50:
        return "NO EFICAZ"
    if item.efectiva is False and item.verificacion_eficacia:
        return "NO EFICAZ"
    return "PENDIENTE"


def _style(ws, title: str):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=17, color=BLANCO)
    ws["A1"].fill = PatternFill("solid", fgColor=AZUL_OSCURO)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:H2")
    ws["A2"] = f"ERP SST PRO · Generado {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = Font(size=10, color=GRIS)
    ws["A2"].alignment = Alignment(horizontal="center")

    thin = Side(style="thin", color=BORDE)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _header(ws, row: int, columns: list[str]):
    for col, name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.font = Font(bold=True, color=BLANCO)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws):
    for col in range(1, ws.max_column + 1):
        size = 12
        for row in range(1, ws.max_row + 1):
            value = ws.cell(row=row, column=col).value
            if value is not None:
                size = max(size, min(len(str(value)) + 2, 55))
        ws.column_dimensions[get_column_letter(col)].width = size


def generar_excel_medida_correctiva(db: Session, medida_id: int) -> bytes:
    medida = db.query(CapaSST).filter(CapaSST.id == medida_id, CapaSST.activo.is_(True)).first()
    if not medida:
        raise ValueError("Medida correctiva no encontrada")

    seguimientos = (
        db.query(CapaSeguimientoSST)
        .filter(CapaSeguimientoSST.capa_id == medida.id, CapaSeguimientoSST.activo.is_(True))
        .order_by(CapaSeguimientoSST.fecha_seguimiento.asc(), CapaSeguimientoSST.id.asc())
        .all()
    )

    evidencias = (
        db.query(ArchivoSST)
        .filter(
            ArchivoSST.modulo.in_(["CAPA", "MEDIDAS_CORRECTIVAS"]),
            ArchivoSST.referencia_id == medida.id,
            ArchivoSST.activo.is_(True),
        )
        .order_by(ArchivoSST.fecha_creacion.desc())
        .all()
    )

    wb = Workbook()

    ws = wb.active
    ws.title = "Resumen Ejecutivo"
    _style(ws, f"Reporte Ejecutivo Medida Correctiva · {_safe(medida.codigo)}")
    _header(ws, 4, ["Campo", "Valor"])

    resumen = [
        ("Código", _safe(medida.codigo)),
        ("Título", _safe(medida.titulo)),
        ("Estado", _safe(medida.estado)),
        ("Prioridad", _safe(medida.prioridad)),
        ("Origen", _safe(medida.origen)),
        ("Responsable", _safe(medida.responsable)),
        ("Fecha apertura", _fmt_date(medida.fecha_apertura)),
        ("Fecha compromiso", _fmt_date(medida.fecha_compromiso)),
        ("Fecha cierre", _fmt_date(medida.fecha_cierre)),
        ("Avance", _num(medida.avance)),
        ("Eficacia", _eficacia_estado(medida)),
        ("% eficacia", _num(getattr(medida, "porcentaje_eficacia", 0))),
        ("Costo estimado", _num(getattr(medida, "costo_estimado", 0))),
        ("Costo real", _num(getattr(medida, "costo_real", 0))),
    ]
    for r, (campo, valor) in enumerate(resumen, start=5):
        ws.cell(r, 1, campo)
        ws.cell(r, 2, valor)
        ws.cell(r, 1).font = Font(bold=True, color=AZUL_OSCURO)

    ws2 = wb.create_sheet("Acciones y Causa")
    _style(ws2, "Acciones y causa raíz")
    _header(ws2, 4, ["Elemento", "Detalle"])
    data = [
        ("Descripción", _safe(medida.descripcion)),
        ("Causa raíz", _safe(medida.causa_raiz)),
        ("Por qué 1", _safe(medida.porque_1)),
        ("Por qué 2", _safe(medida.porque_2)),
        ("Por qué 3", _safe(medida.porque_3)),
        ("Por qué 4", _safe(medida.porque_4)),
        ("Por qué 5", _safe(medida.porque_5)),
        ("Acción inmediata", _safe(medida.accion_inmediata)),
        ("Acción correctiva", _safe(medida.accion_correctiva)),
        ("Acción preventiva", _safe(medida.accion_preventiva)),
        ("Verificación eficacia", _safe(medida.verificacion_eficacia)),
        ("Observaciones", _safe(medida.observaciones)),
    ]
    for r, (campo, valor) in enumerate(data, start=5):
        ws2.cell(r, 1, campo)
        ws2.cell(r, 2, valor)
        ws2.cell(r, 1).font = Font(bold=True, color=AZUL_OSCURO)

    ws3 = wb.create_sheet("Seguimientos")
    _style(ws3, "Seguimientos")
    _header(ws3, 4, ["ID", "Fecha", "Responsable", "Avance", "Resultado", "Comentario"])
    if seguimientos:
        for r, seg in enumerate(seguimientos, start=5):
            ws3.cell(r, 1, seg.id)
            ws3.cell(r, 2, _fmt_date(getattr(seg, "fecha_seguimiento", None)))
            ws3.cell(r, 3, _safe(getattr(seg, "responsable", None)))
            ws3.cell(r, 4, _num(getattr(seg, "avance", 0)))
            ws3.cell(r, 5, _safe(getattr(seg, "resultado", None)))
            ws3.cell(r, 6, _safe(getattr(seg, "comentario", None)))
    else:
        ws3.cell(5, 1, "No existen seguimientos registrados.")

    ws4 = wb.create_sheet("Evidencias")
    _style(ws4, "Evidencias")
    _header(ws4, 4, ["ID", "Nombre", "Módulo", "Tipo", "MIME", "URL", "Fecha"])
    if evidencias:
        for r, archivo in enumerate(evidencias, start=5):
            ws4.cell(r, 1, archivo.id)
            ws4.cell(r, 2, _safe(archivo.nombre_original))
            ws4.cell(r, 3, _safe(archivo.modulo))
            ws4.cell(r, 4, _safe(archivo.tipo))
            ws4.cell(r, 5, _safe(archivo.mime_type))
            ws4.cell(r, 6, _safe(archivo.url))
            ws4.cell(r, 7, _fmt_date(archivo.fecha_creacion))
    else:
        ws4.cell(5, 1, "No existen evidencias registradas.")

    ws5 = wb.create_sheet("Trazabilidad")
    _style(ws5, "Trazabilidad técnica")
    _header(ws5, 4, ["Trazabilidad"])
    ws5.cell(5, 1, _safe(medida.trazabilidad, "Sin trazabilidad registrada."))

    for sheet in wb.worksheets:
        _autosize(sheet)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
