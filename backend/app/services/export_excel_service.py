# ============================================================
# SERVICIO EXCEL CORPORATIVO SST
# FASE 2.2.1C - Motor de Exportación Corporativo SST
# ============================================================

from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generar_excel_corporativo(
    titulo: str,
    codigo: str,
    empresa,
    configuracion=None,
    columnas=None,
    filas=None,
):
    """
    Genera Excel corporativo reutilizable para SG-SST.
    """

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte SST"

    columnas = columnas or []
    filas = filas or []

    prefijo = configuracion.prefijo_documental if configuracion else "SGSST"
    version = configuracion.version_documental if configuracion else "1.0"
    pie = (
        configuracion.pie_documental
        if configuracion
        else "Documento controlado generado desde ERP SST PRO."
    )

    total_cols = max(len(columnas), 6)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    ws["A1"] = empresa.nombre
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    ws["A2"] = f"{titulo} | Código: {prefijo}-{codigo} | Versión: {version}"
    ws["A2"].font = Font(bold=True, size=12)
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=total_cols)
    ws["A3"] = f"NIT: {empresa.nit} | Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A3"].alignment = Alignment(horizontal="center")

    start_row = 5

    header_fill = PatternFill("solid", fgColor="1D4ED8")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(border_style="thin", color="D1D5DB")

    for col_idx, columna in enumerate(columnas, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=columna)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

    for row_idx, fila in enumerate(filas, start=start_row + 1):
        for col_idx, valor in enumerate(fila, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

    for col_idx in range(1, total_cols + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 24

    footer_row = start_row + len(filas) + 3
    ws.merge_cells(
        start_row=footer_row,
        start_column=1,
        end_row=footer_row,
        end_column=total_cols,
    )
    ws.cell(row=footer_row, column=1, value=pie)
    ws.cell(row=footer_row, column=1).font = Font(italic=True)
    ws.cell(row=footer_row, column=1).alignment = Alignment(horizontal="center")

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer