import asyncio
from io import BytesIO
from datetime import date
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from app.models.empleado import Empleado
from app.routers.empleados_perfil import (
    _construir_dashboard,
    dashboard_perfil_sociodemografico,
    exportar_dashboard_excel,
    exportar_dashboard_pdf,
    router,
)
from app.schemas.empleado_schema import EmpleadoCreate, EmpleadoUpdate


def _empleado(
    empleado_id,
    fecha_ingreso,
    *,
    fecha_retiro=None,
    estado_laboral="ACTIVO",
    activo=True,
    sede_id=10,
    sede_nombre="Principal",
):
    return SimpleNamespace(
        id=empleado_id,
        fecha_ingreso=fecha_ingreso,
        fecha_retiro=fecha_retiro,
        estado_laboral=estado_laboral,
        activo=activo,
        sede_id=sede_id,
        sede=SimpleNamespace(nombre=sede_nombre) if sede_id else None,
    )


def test_dashboard_aplica_corte_retiro_y_legado_incompleto():
    empleados = [
        _empleado(1, date(2026, 1, 10)),
        _empleado(2, date(2026, 2, 1), fecha_retiro=date(2026, 3, 15)),
        _empleado(3, date(2026, 4, 1)),
        _empleado(4, None),
        _empleado(5, date(2026, 1, 5), estado_laboral="RETIRADO", activo=False),
        _empleado(
            6,
            date(2026, 1, 20),
            fecha_retiro=date(2026, 4, 10),
            estado_laboral="INACTIVO",
            activo=False,
            sede_id=None,
        ),
    ]
    perfiles = [
        SimpleNamespace(empleado_id=1, completado=True),
        SimpleNamespace(empleado_id=2, completado=False),
        SimpleNamespace(empleado_id=3, completado=True),
    ]

    resultado = _construir_dashboard(empleados, perfiles, 2026, 3)

    assert resultado["periodo"] == {
        "anio": 2026,
        "mes": 3,
        "fecha_corte": "2026-03-31",
        "etiqueta": "Marzo 2026",
    }
    assert resultado["kpis"] == {
        "total_periodo": 4,
        "activos": 2,
        "inactivos": 2,
        "perfiles_registrados": 2,
        "perfiles_completados": 1,
        "cobertura_perfil": 50.0,
        "sin_fecha_ingreso": 1,
        "datos_historicos_incompletos": 1,
    }
    assert resultado["tendencia_mensual"] == [
        {"mes": 1, "nombre": "Ene", "activos": 2, "inactivos": 1},
        {"mes": 2, "nombre": "Feb", "activos": 3, "inactivos": 1},
        {"mes": 3, "nombre": "Mar", "activos": 2, "inactivos": 2},
    ]
    assert sum(item["total"] for item in resultado["por_sede"]) == 4


def test_dashboard_rechaza_empresa_ajena_antes_de_consultar():
    db = MagicMock()
    usuario = SimpleNamespace(rol="RESPONSABLE_SST", empresa_id=7)

    with pytest.raises(HTTPException) as exc:
        dashboard_perfil_sociodemografico(
            empresa_id=8,
            sede_id=None,
            anio=2026,
            mes=None,
            db=db,
            usuario=usuario,
        )

    assert exc.value.status_code == 403
    db.query.assert_not_called()


def test_todos_los_endpoints_requieren_rol_sst():
    for route in router.routes:
        dependencies = route.dependant.dependencies
        assert any(
            getattr(dependency.call, "__name__", "") == "role_checker"
            for dependency in dependencies
        ), f"Endpoint sin RBAC ROLES_SST: {route.path}"


def test_exportaciones_generan_xlsx_y_pdf_validos():
    empleado = _empleado(1, date(2026, 1, 10))
    empleado.nombres = "Ana"
    empleado.apellidos = "Pérez"
    empleado.documento = "123"
    empleado.empresa_id = 7
    empleado.empresa = SimpleNamespace(nombre="Empresa Uno")
    perfil = SimpleNamespace(empleado_id=1, completado=True)

    class Query:
        def __init__(self, items):
            self.items = items

        def options(self, *args):
            return self

        def filter(self, *args):
            return self

        def order_by(self, *args):
            return self

        def all(self):
            return self.items

    class DB:
        def query(self, model):
            return Query([empleado] if model is Empleado else [perfil])

    async def contenido(response):
        return b"".join([chunk async for chunk in response.body_iterator])

    parametros = {
        "empresa_id": 7,
        "sede_id": None,
        "anio": 2026,
        "mes": 3,
        "db": DB(),
        "usuario": SimpleNamespace(rol="RESPONSABLE_SST", empresa_id=7),
    }
    excel = asyncio.run(contenido(exportar_dashboard_excel(**parametros)))
    pdf = asyncio.run(contenido(exportar_dashboard_pdf(**parametros)))

    libro = load_workbook(BytesIO(excel), read_only=True)
    assert libro.sheetnames == ["Resumen", "Detalle empleados"]
    libro.close()
    assert pdf.startswith(b"%PDF-")


@pytest.mark.parametrize("schema", [EmpleadoCreate, EmpleadoUpdate])
def test_schema_rechaza_retiro_anterior_al_ingreso(schema):
    payload = {
        "fecha_ingreso": date(2026, 2, 1),
        "fecha_retiro": date(2026, 1, 31),
    }
    if schema is EmpleadoCreate:
        payload.update(
            nombres="Ana",
            apellidos="Pérez",
            documento="123",
            empresa_id=1,
        )

    with pytest.raises(ValidationError, match="fecha_retiro"):
        schema(**payload)
